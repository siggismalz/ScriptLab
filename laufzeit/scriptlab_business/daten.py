"""Verzögerte lokale Tabellenpläne mit fest begrenzter Ausdruckssprache.

SQL-Fragmente kommen ausschließlich aus diesem Modul. Benutzerwerte werden
gebunden; Bezeichner werden gegen das jeweilige Schema geprüft und zitiert.
"""

from __future__ import annotations

import csv
import json
import math
import re
import threading
import time
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

import duckdb

from .typen import (AGGREGATE, FORMATE, KONFLIKTE, PRUEFUNGEN, VERGLEICHE,
                    DatenFehler, Spalte, Tabelle, auswahl, objekt, text)


@dataclass(frozen=True)
class _Quelle:
    pfad: Path
    groesse: int
    hash: str
    stat: tuple


@dataclass(frozen=True)
class _Plan:
    sql: str
    parameter: tuple[Any, ...]
    quellen: tuple[_Quelle, ...]


def _bezeichner(name: str) -> str:
    return '"' + text(name, "Spaltenname").replace('"', '""') + '"'


def _typ(wert: Any) -> str:
    fest = {"TEXT": "VARCHAR", "GANZZAHL": "BIGINT", "ZAHL": "DOUBLE",
            "WAHRHEIT": "BOOLEAN", "DATUM": "DATE", "ZEITSTEMPEL": "TIMESTAMP", "DEZIMAL": "DECIMAL(18,2)"}
    if isinstance(wert, str) and wert in fest:
        return fest[wert]
    if isinstance(wert, dict) and wert.get("art") == "DEZIMAL":
        p, s = wert.get("praezision", 18), wert.get("skala", 2)
        if type(p) is int and type(s) is int and 1 <= p <= 38 and 0 <= s <= p:
            return f"DECIMAL({p},{s})"
    raise DatenFehler("Wähle einen unterstützten Datentyp und eine gültige Dezimalgenauigkeit.")


def _literal(wert: dict) -> tuple[str, list[Any]]:
    objekt(wert, "Wert")
    typ = _typ(wert.get("typ"))
    roh = wert.get("wert")
    if roh is None:
        return f"CAST(? AS {typ})", [None]
    try:
        if typ == "VARCHAR":
            text(roh, "Textwert", leer=True)
        elif typ == "BIGINT":
            if type(roh) is not int or not -(2**63) <= roh < 2**63:
                raise ValueError()
        elif typ.startswith("DECIMAL"):
            if not isinstance(roh, str) or not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", roh):
                raise ValueError()
            zahl = Decimal(roh)
            p, s = map(int, re.findall(r"\d+", typ))
            if -zahl.as_tuple().exponent > s or max(zahl.adjusted()+1, 0) > p-s:
                raise ValueError()
            roh = zahl
        elif typ == "DOUBLE":
            if type(roh) not in (int, float) or not math.isfinite(roh):
                raise ValueError()
        elif typ == "BOOLEAN":
            if type(roh) is not bool:
                raise ValueError()
        elif typ == "DATE":
            if not isinstance(roh, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", roh):
                raise ValueError()
            roh = date.fromisoformat(roh)
        elif typ == "TIMESTAMP":
            roh = datetime.fromisoformat(text(roh, "Zeitstempel"))
            if roh.tzinfo is not None:
                raise ValueError()
    except (ValueError, TypeError, InvalidOperation):
        raise DatenFehler("Der Wert passt nicht zum gewählten Datentyp.") from None
    return f"CAST(? AS {typ})", [roh]


def _jsonwert(wert: Any) -> Any:
    if isinstance(wert, Decimal):
        return str(wert)
    if isinstance(wert, (date, datetime)):
        return wert.isoformat()
    if isinstance(wert, float) and not math.isfinite(wert):
        raise DatenFehler("Das Ergebnis enthält eine nicht endliche Zahl.")
    if isinstance(wert, (bytes, list, tuple, dict)):
        raise DatenFehler("Verschachtelte oder binäre Daten benötigen eine ausdrückliche Aufbereitung.")
    return wert


class DatenEngine:
    """Eine lauflokale Datenbank; Handles sind außerhalb dieser Engine wertlos."""

    def __init__(self, kontext):
        self.kontext = kontext
        self._plaene: dict[str, _Plan] = {}
        self._tabellen: dict[str, Tabelle] = {}
        ordner = Path(kontext.arbeitsordner)
        ordner.mkdir(parents=True, exist_ok=True)
        self._ordner = ordner
        speicher = getattr(getattr(kontext, "grenzen", None), "duckdb_speicher", "256MB")
        if not isinstance(speicher, str) or not re.fullmatch(r"[1-9]\d*(?:MB|GB)", speicher):
            raise DatenFehler("Die Speichergrenze der Datenengine ist ungültig.")
        self._verbindung = duckdb.connect(str(ordner / f"daten-{uuid.uuid4().hex}.duckdb"), config={
            "autoinstall_known_extensions": "false", "autoload_known_extensions": "false",
            "allow_unsigned_extensions": "false", "threads": "2", "memory_limit": speicher,
        })
        self._verbindung.execute("SET temp_directory = ?", [str(ordner / "duckdb-temp")])
        temp_limit = self._grenze("max_temp_bytes",4*1024**3)
        if type(temp_limit) is not int or temp_limit <= 0:
            raise DatenFehler("Die temporäre Speichergrenze ist ungültig.")
        self._verbindung.execute("SET max_temp_directory_size = ?", [f"{temp_limit}B"])
        self._geschlossen = False

    def _grenze(self, name, vorgabe):
        return getattr(getattr(self.kontext, "grenzen", None), name, vorgabe)

    def _pruefen(self):
        self.kontext.pruefen()

    def schliessen(self) -> None:
        if not self._geschlossen:
            self._verbindung.close()
            self._geschlossen = True

    def tabelle(self, referenz: Any) -> Tabelle:
        kennung = referenz.id if isinstance(referenz, Tabelle) else objekt(referenz, "Tabelle").get("tabelle")
        if not isinstance(kennung, str) or kennung not in self._tabellen:
            raise DatenFehler("Die Tabelle gehört nicht zu diesem Lauf oder ist nicht mehr verfügbar.")
        return self._tabellen[kennung]

    def _plan(self, tabelle: Tabelle) -> _Plan:
        return self._plaene[self.tabelle(tabelle).id]

    def quellpfade(self, tabelle: Tabelle) -> tuple[Path, ...]:
        plan = self._plan(tabelle)
        self._pruefe_quellen(plan)
        return tuple(q.pfad for q in plan.quellen)

    def _neu(self, sql: str, parameter=(), quellen=(), *, max_spalten=16384) -> Tabelle:
        try:
            relation = self._verbindung.sql(sql, params=list(parameter))
            namen, typen = relation.columns, relation.types
            if len({n.casefold() for n in namen}) != len(namen):
                raise DatenFehler("Spaltennamen müssen auch bei unterschiedlicher Groß-/Kleinschreibung eindeutig sein.")
            if len(namen) > max_spalten:
                raise DatenFehler("Die Tabelle enthält mehr als 16384 Spalten.")
            tabelle = Tabelle(uuid.uuid4().hex, tuple(Spalte(n, str(t)) for n, t in zip(namen, typen)))
            self._plaene[tabelle.id] = _Plan(sql, tuple(parameter), tuple(dict.fromkeys(quellen)))
            self._tabellen[tabelle.id] = tabelle
            return tabelle
        except duckdb.Error as fehler:
            raise DatenFehler("Die Tabelle konnte nicht aufgebaut werden. Prüfe Spaltennamen, Datentypen und Dateiformat.") from fehler

    def _spalte(self, tabelle: Tabelle, name: Any) -> str:
        name = text(name, "Spalte")
        if name not in [s.name for s in tabelle.spalten]:
            raise DatenFehler(f"Die Spalte „{name}“ ist in dieser Tabelle nicht vorhanden.")
        return _bezeichner(name)

    def _namen(self, tabelle: Tabelle, namen: Any, *, leer=False) -> list[str]:
        if not isinstance(namen, list) or (not leer and not namen) or len(namen) != len(set(n for n in namen if isinstance(n, str))):
            raise DatenFehler("Wähle unterschiedliche vorhandene Spalten aus.")
        return [self._spalte(tabelle, n) for n in namen]

    def _pruefe_quellen(self, plan: _Plan) -> None:
        from .dateien import datei_fingerabdruck
        self._pruefen()
        for quelle in plan.quellen:
            aktuell = datei_fingerabdruck(self.kontext, quelle.pfad)
            if aktuell != {"pfad":str(quelle.pfad),"bytes":quelle.groesse,"sha256":quelle.hash,"stat":list(quelle.stat)}:
                raise DatenFehler("Eine Quelldatei wurde seit dem Import verändert. Lies sie erneut ein.")

    def _skalar(self, tabelle: Tabelle, ausdruck: str, parameter=()) -> Any:
        plan = self._plan(tabelle)
        self._pruefe_quellen(plan)
        return self._verbindung.execute(f"SELECT {ausdruck} FROM ({plan.sql}) AS t", [*parameter, *plan.parameter]).fetchone()[0]

    def zeilen(self, tabelle: Tabelle, paket: int = 1000, *, zeitlimit: float | None = None) -> Iterator[tuple]:
        if type(paket) is not int or not 1 <= paket <= 10000:
            raise DatenFehler("Die Paketgröße ist ungültig.")
        plan = self._plan(tabelle)
        self._pruefe_quellen(plan)
        # Ein eigener Cursor verhindert, dass ein weiteres Metadatenlesen die
        # laufende Ausgabe ersetzt; die Daten bleiben auf DB-Seite gestreamt.
        cursor = self._verbindung.cursor()
        abgelaufen = threading.Event()
        def abbrechen():
            abgelaufen.set()
            cursor.interrupt()
        timer = threading.Timer(zeitlimit, abbrechen) if zeitlimit is not None else None
        if timer:
            timer.daemon = True
            timer.start()
        try:
            cursor.execute(plan.sql, list(plan.parameter))
            while zeilen := cursor.fetchmany(paket):
                self._pruefen()
                yield from zeilen
        except duckdb.Error as fehler:
            if abgelaufen.is_set():
                raise DatenFehler("Die Vorschau hat ihr Zeitlimit überschritten. Vereinfache den Ablauf oder begrenze die Quelle.") from fehler
            raise DatenFehler("Die Tabellendaten konnten nicht verarbeitet werden. Prüfe Werte und Datentypen.") from fehler
        finally:
            if timer:
                timer.cancel()
                timer.join()
            cursor.close()
        self._pruefe_quellen(plan)

    def importieren(self, format: str, pfad: str, optionen: dict | None = None) -> Tabelle:
        auswahl(format, FORMATE, "Dateiformat")
        optionen = objekt(optionen or {}, "Importoptionen")
        quelle = Path(self.kontext.pfad(pfad))
        if not quelle.is_file():
            raise DatenFehler("Die Quelldatei ist nicht vorhanden.")
        if quelle.stat().st_size > self._grenze("max_datei_bytes", 64*1024*1024):
            raise DatenFehler("Die Quelldatei überschreitet die eingestellte Dateigröße.")
        from .dateien import datei_fingerabdruck
        fp = datei_fingerabdruck(self.kontext, quelle)
        finger = _Quelle(quelle,fp["bytes"],fp["sha256"],tuple(fp["stat"]))
        try:
            if format == "XLSX":
                ergebnis = self._xlsx_lesen(quelle, optionen, finger)
            elif format == "CSV":
                ergebnis = self._csv_lesen(quelle, optionen, finger)
            elif format == "PARQUET":
                ergebnis = self._neu("SELECT * FROM read_parquet(?)", [str(quelle)], [finger])
            else:
                art = "newline_delimited" if optionen.get("zeilenformat", False) else "array"
                ergebnis = self._neu("SELECT * FROM read_json(?,format=?,sample_size=-1,ignore_errors=false)", [str(quelle), art], [finger])
                if any(any(s in sp.typ for s in ("STRUCT", "MAP", "[]", "BLOB")) for sp in ergebnis.spalten):
                    raise DatenFehler("JSON muss eine Tabelle mit einfachen Werten enthalten.")
            self._pruefe_quellen(self._plan(ergebnis))
            return ergebnis
        except (OSError, UnicodeError, csv.Error, duckdb.Error) as fehler:
            raise DatenFehler("Die Datei konnte nicht eingelesen werden. Prüfe Format, Kodierung und Zugriffsrechte.") from fehler

    def _csv_lesen(self, quelle: Path, optionen: dict, finger: _Quelle) -> Tabelle:
        trenner = optionen.get("trennzeichen", ";")
        if not isinstance(trenner, str) or len(trenner) != 1 or trenner in ('\r', '\n', '"', '\x00'):
            raise DatenFehler("Wähle genau ein zulässiges Trennzeichen.")
        kodierung = auswahl(optionen.get("kodierung", "UTF8"), ("UTF8", "UTF8_BOM", "UTF16", "CP1252"), "Kodierung")
        policy = auswahl(optionen.get("fehlzeilen", "STOPP"), ("STOPP", "AUSLASSEN_MIT_BERICHT"), "Fehlzeilen")
        eingang = quelle
        if kodierung not in ("UTF8", "UTF8_BOM"):
            eingang = self._ordner / f"csv-{uuid.uuid4().hex}.csv"
            with quelle.open("r", encoding={"UTF16": "utf-16", "CP1252": "cp1252"}[kodierung], errors="strict", newline="") as von, eingang.open("w", encoding="utf-8", newline="") as nach:
                for teil in iter(lambda: von.read(131072), ""):
                    self._pruefen()
                    nach.write(teil)
        nullwert = text(optionen.get("nullwert", ""), "Leerwert", leer=True)
        header = optionen.get("kopfzeile", True)
        if type(header) is not bool:
            raise DatenFehler("Die Einstellung Kopfzeile muss Ja oder Nein sein.")
        namen = optionen.get("spalten")
        with eingang.open("r", encoding="utf-8-sig", errors="strict", newline="") as datei:
            erste = next(csv.reader(datei, delimiter=trenner), None)
        if erste is None:
            raise DatenFehler("Die CSV-Datei ist leer.")
        if namen is not None:
            if not isinstance(namen, list) or not all(isinstance(n, str) and n.strip() for n in namen):
                raise DatenFehler("Gib gültige Spaltennamen an.")
            if len(namen) != len(erste):
                raise DatenFehler("Die Zahl der Spaltennamen passt nicht zur CSV-Datei.")
        elif header:
            namen = erste
        else:
            namen = [f"Spalte{i+1}" for i in range(len(erste))]
        if namen is not None and (len({n.casefold() for n in namen}) != len(namen) or any(not n.strip() for n in namen)):
            raise DatenFehler("CSV-Spaltennamen müssen gefüllt und eindeutig sein.")
        parameter = [str(eingang), trenner, header, nullwert, policy != "STOPP", {n:"VARCHAR" for n in namen}]
        sql = "SELECT * FROM read_csv(?,delim=?,header=?,nullstr=?,ignore_errors=?,strict_mode=true,auto_detect=false,quote='\"',escape='\"',columns=?)"
        if policy == "STOPP":
            return self._neu(sql, parameter, [finger])
        # Fehlerberichte müssen zur tatsächlichen einmaligen Einlesung passen.
        # Deshalb wird nur die ausdrücklich tolerante Einlesung DB-seitig fixiert.
        reject = f"sl_reject_{uuid.uuid4().hex}"
        tname = f"sl_csv_{uuid.uuid4().hex}"
        sql = sql[:-1] + ",store_rejects=true,rejects_table=?)"
        self._verbindung.execute(f"CREATE TABLE {_bezeichner(tname)} AS {sql}", [*parameter, reject])
        zahl = self._verbindung.execute(f"SELECT count(DISTINCT (scan_id,file_id,line)) FROM {_bezeichner(reject)}").fetchone()[0]
        self.kontext.ereignis("daten.fehlzeilen", anzahl=zahl, meldung=f"{zahl} fehlerhafte CSV-Zeilen wurden ausgelassen.")
        return self._neu(f"SELECT * FROM {_bezeichner(tname)}", [], [finger])

    def _xlsx_lesen(self, quelle: Path, optionen: dict, finger: _Quelle) -> Tabelle:
        from openpyxl import load_workbook
        from .dateien import zip_pruefen
        zip_pruefen(self.kontext, quelle)
        auswahl(optionen.get("formeln", "STOPP"), ("STOPP", "WERTE"), "Formeln")
        wb = load_workbook(quelle, read_only=True, data_only=optionen.get("formeln", "STOPP") == "WERTE", keep_links=False)
        try:
            blatt = optionen.get("blatt") or wb.sheetnames[0]
            if blatt not in wb.sheetnames:
                raise DatenFehler("Das gewählte Tabellenblatt ist nicht vorhanden.")
            ws = wb[blatt]
            ws.reset_dimensions()
            zeilen = ws.iter_rows()
            erste = next(zeilen, None)
            if not erste:
                raise DatenFehler("Das Tabellenblatt ist leer.")
            if optionen.get("kopfzeile", True):
                namen = [text(z.value, "XLSX-Spaltenname") for z in erste]
            else:
                namen = [f"Spalte{i+1}" for i in range(len(erste))]
                from itertools import chain
                zeilen = chain([erste], zeilen)
            if len({n.casefold() for n in namen}) != len(namen):
                raise DatenFehler("XLSX-Spaltennamen müssen eindeutig sein.")
            # Import als Text ist ausdrücklich verlustfrei für Kennungen. Die
            # fachliche Typisierung erfolgt danach im sichtbaren Typbaustein.
            tname = f"sl_xlsx_{uuid.uuid4().hex}"
            self._verbindung.execute(f"CREATE TABLE {_bezeichner(tname)} ({','.join(_bezeichner(n)+' VARCHAR' for n in namen)})")
            insert = f"INSERT INTO {_bezeichner(tname)} VALUES ({','.join('?' for _ in namen)})"
            paket = []
            for nr, zeile in enumerate(zeilen, 1):
                self._pruefen()
                if nr > self._grenze("max_zeilen", 100000):
                    raise DatenFehler("Das Tabellenblatt überschreitet die eingestellte Zeilengrenze.")
                if len(zeile) != len(namen):
                    raise DatenFehler("Eine XLSX-Zeile hat eine abweichende Spaltenzahl.")
                if optionen.get("formeln", "STOPP") == "STOPP" and any(z.data_type == "f" for z in zeile):
                    raise DatenFehler("Das Tabellenblatt enthält Formeln. Wähle bewusst gespeicherte Werte oder eine Quelle ohne Formeln.")
                paket.append(tuple(None if z.value is None else _xlsx_text(z.value) for z in zeile))
                if len(paket) == 1000:
                    self._verbindung.executemany(insert, paket)
                    paket.clear()
            if paket:
                self._verbindung.executemany(insert, paket)
            return self._neu(f"SELECT * FROM {_bezeichner(tname)}", [], [finger])
        finally:
            wb.close()

    def _ausdruck(self, tabelle: Tabelle, ast: Any, tiefe=0, budget=None) -> tuple[str, list[Any]]:
        budget = budget if budget is not None else [0]
        budget[0] += 1
        if tiefe > 20 or budget[0] > 500:
            raise DatenFehler("Der Ausdruck ist zu stark verschachtelt oder zu groß.")
        ast = objekt(ast, "Ausdruck")
        art = ast.get("art")
        if art == "spalte":
            return self._spalte(tabelle, ast.get("name")), []
        if art == "wert":
            return _literal(ast)
        if art == "rechnen":
            operator = auswahl(ast.get("operator"), ("ADD", "SUB", "MUL", "DIV"), "Rechenart")
            links, lp = self._ausdruck(tabelle, ast.get("links"), tiefe+1, budget)
            rechts, rp = self._ausdruck(tabelle, ast.get("rechts"), tiefe+1, budget)
            if operator == "DIV":
                # Kein Infinity und kein stilles NULL bei Division durch Null.
                return f"(({links}) / CASE WHEN ({rechts})=0 THEN error('Division durch 0 ist nicht erlaubt.') ELSE ({rechts}) END)", lp+rp+rp
            return f"(({links}) {dict(ADD='+',SUB='-',MUL='*')[operator]} ({rechts}))", lp+rp
        if art == "funktion":
            funktion = auswahl(ast.get("funktion"), ("JAHR", "MONAT", "TAG", "TRIM", "GROSS", "KLEIN", "LAENGE", "REGEX"), "Funktion")
            argument, parameter = self._ausdruck(tabelle, ast.get("argument"), tiefe+1, budget)
            if funktion == "REGEX":
                muster = text(ast.get("muster"), "Suchmuster")
                gruppe = ast.get("gruppe", 0)
                if len(muster) > 1000 or type(gruppe) is not int or not 0 <= gruppe <= 99:
                    raise DatenFehler("Das Suchmuster oder die Gruppennummer ist ungültig.")
                return f"regexp_extract(CAST({argument} AS VARCHAR),?,?)", parameter+[muster, gruppe]
            funktion_sql = dict(JAHR="year", MONAT="month", TAG="day", TRIM="trim", GROSS="upper", KLEIN="lower", LAENGE="length")[funktion]
            return f"{funktion_sql}({argument})", parameter
        raise DatenFehler("Dieser Ausdruck ist nicht unterstützt. Freie Formeln sind nicht erlaubt.")

    def _bedingung(self, tabelle: Tabelle, ast: Any, tiefe=0, budget=None) -> tuple[str, list[Any]]:
        budget = budget if budget is not None else [0]
        budget[0] += 1
        if tiefe > 20 or budget[0] > 500:
            raise DatenFehler("Die Bedingung ist zu stark verschachtelt oder zu groß.")
        ast = objekt(ast, "Bedingung")
        art = ast.get("art")
        if art in ("UND", "ODER"):
            kinder = ast.get("kinder")
            if not isinstance(kinder, list) or not 1 <= len(kinder) <= 100:
                raise DatenFehler("Verbinde mindestens eine und höchstens 100 Bedingungen.")
            teile = [self._bedingung(tabelle, k, tiefe+1, budget) for k in kinder]
            return '(' + (" AND " if art == "UND" else " OR ").join(f"({s})" for s, _ in teile) + ')', [p for _, ps in teile for p in ps]
        if art == "NICHT":
            sql, parameter = self._bedingung(tabelle, ast.get("kind"), tiefe+1, budget)
            return f"NOT ({sql})", parameter
        if art != "vergleich":
            raise DatenFehler("Wähle einen unterstützten Vergleich.")
        operator = auswahl(ast.get("operator"), VERGLEICHE, "Vergleich")
        links, lp = self._ausdruck(tabelle, ast.get("links", {"art": "spalte", "name": ast.get("spalte")}), tiefe+1, budget)
        if operator in ("IST_LEER", "IST_NICHT_LEER"):
            return f"({links}) IS {'NOT ' if operator == 'IST_NICHT_LEER' else ''}NULL", lp
        if operator == "IN":
            werte = ast.get("werte")
            if not isinstance(werte, list) or not 1 <= len(werte) <= 1000:
                raise DatenFehler("Gib zwischen 1 und 1000 erlaubte Werte an.")
            teile = [_literal(w) for w in werte]
            if any(w.get("wert") is None for w in werte):
                raise DatenFehler("Prüfe leere Werte mit „ist leer“.")
            return f"({links}) IN ({','.join(s for s,_ in teile)})", lp+[p for _, ps in teile for p in ps]
        rechts_ast = ast.get("rechts", ast.get("wert"))
        if isinstance(rechts_ast, dict) and rechts_ast.get("art") == "wert" and rechts_ast.get("wert") is None:
            raise DatenFehler("Prüfe leere Werte mit „ist leer“.")
        rechts, rp = self._ausdruck(tabelle, rechts_ast, tiefe+1, budget)
        if operator in ("ENTHAELT", "BEGINNT", "ENDET"):
            funktion = dict(ENTHAELT="contains", BEGINNT="starts_with", ENDET="ends_with")[operator]
            return f"{funktion}(CAST({links} AS VARCHAR),CAST({rechts} AS VARCHAR))", lp+rp
        return f"(({links}) {dict(EQ='=',NE='<>',GT='>',GE='>=',LT='<',LE='<=')[operator]} ({rechts}))", lp+rp

    def filtern(self, tabelle: Tabelle, bedingung: dict) -> Tabelle:
        sql, parameter = self._bedingung(tabelle, bedingung)
        plan = self._plan(tabelle)
        return self._neu(f"SELECT * FROM ({plan.sql}) AS t WHERE {sql}", [*plan.parameter, *parameter], plan.quellen)

    def spalten(self, tabelle: Tabelle, konfig: dict) -> Tabelle:
        aktion = auswahl(konfig.get("aktion"), ("AUSWAEHLEN", "UMBENENNEN", "TYP", "BERECHNEN"), "Spaltenaktion")
        plan = self._plan(tabelle)
        parameter = []
        if aktion == "AUSWAEHLEN":
            teile = self._namen(tabelle, konfig.get("spalten"))
        elif aktion == "UMBENENNEN":
            namen = objekt(konfig.get("namen"), "Neue Spaltennamen")
            for name in namen:
                self._spalte(tabelle, name)
            neu = [text(namen.get(s.name, s.name), "Neuer Spaltenname") for s in tabelle.spalten]
            if len({n.casefold() for n in neu}) != len(neu):
                raise DatenFehler("Die neuen Spaltennamen sind nicht eindeutig.")
            teile = [f"{_bezeichner(s.name)} AS {_bezeichner(n)}" for s,n in zip(tabelle.spalten, neu)]
        else:
            name = text(konfig.get("spalte"), "Ergebnisspalte")
            if aktion == "TYP":
                spalte = self._spalte(tabelle, name)
                zieltyp = _typ(konfig.get('typ'))
                fehlwerte = auswahl(konfig.get("fehlwerte","STOPP"), ("STOPP","NULL","AUSLASSEN_MIT_BERICHT"), "Ungültige Werte")
                ausdruck = spalte
                if zieltyp.startswith("DECIMAL") or zieltyp in ("DOUBLE", "BIGINT"):
                    dezimal = auswahl(konfig.get("dezimaltrennzeichen", "."), (".",","), "Dezimaltrennzeichen")
                    tausender = auswahl(konfig.get("tausendertrennzeichen", ""), ("", ".", ",", " "), "Tausendertrennzeichen")
                    if tausender and dezimal == tausender:
                        raise DatenFehler("Dezimal- und Tausendertrennzeichen müssen verschieden sein.")
                    if tausender:
                        ausdruck = f"replace(CAST({ausdruck} AS VARCHAR),?,'')"
                        parameter.append(tausender)
                    if dezimal != ".":
                        ausdruck = f"replace(CAST({ausdruck} AS VARCHAR),?,'.')"
                        parameter.append(dezimal)
                    if next(s.typ for s in tabelle.spalten if s.name == name) == "VARCHAR":
                        ganz = r"\d+" if not tausender else rf"(?:\d+|\d{{1,3}}(?:{re.escape(tausender)}\d{{3}})+)"
                        muster = rf"[+-]?{ganz}(?:{re.escape(dezimal)}\d+)?"
                        # Die gewählten Trennzeichen sind verbindlich; kaputte
                        # Gruppierung wird nicht durch replace scheinbar gültig.
                        fehlerwert = "error('Der Wert hat ein ungueltiges Zahlenformat.')" if fehlwerte == "STOPP" else "NULL"
                        ausdruck = f"CASE WHEN {spalte} IS NULL THEN NULL WHEN regexp_full_match({spalte},?) THEN {ausdruck} ELSE {fehlerwert} END"
                        parameter.insert(0,muster)
                if zieltyp in ("DATE","TIMESTAMP"):
                    formate = {"ISO":"%Y-%m-%d", "DMY":"%d.%m.%Y", "MDY":"%m/%d/%Y", "YYYY-MM-DD":"%Y-%m-%d", "DD.MM.YYYY":"%d.%m.%Y", "MM/DD/YYYY":"%m/%d/%Y"}
                    formatwert = konfig.get("datumsformat","ISO")
                    if formatwert not in formate:
                        raise DatenFehler("Wähle ein unterstütztes Datumsformat.")
                    if zieltyp == "DATE" or formatwert != "ISO":
                        ausdruck = f"strptime(CAST({ausdruck} AS VARCHAR),?)"
                        parameter.append(formate[formatwert])
                ausdruck = f"CAST({ausdruck} AS {zieltyp})"
                if fehlwerte != "STOPP":
                    ausdruck = f"TRY({ausdruck})"
                    invalid = f"{spalte} IS NOT NULL AND ({ausdruck}) IS NULL"
                    anzahl = self._skalar(tabelle,f"count(*) FILTER (WHERE {invalid})",parameter)
                    self.kontext.ereignis("daten.typfehler",anzahl=anzahl,spalte=name,regel=fehlwerte)
                    if fehlwerte == "AUSLASSEN_MIT_BERICHT":
                        gefiltert = self._neu(f"SELECT * FROM ({plan.sql}) AS t WHERE NOT ({invalid})",[*plan.parameter,*parameter],plan.quellen)
                        neue_konfig = dict(konfig,fehlwerte="STOPP")
                        return self.spalten(gefiltert,neue_konfig)
            else:
                if name not in [s.name for s in tabelle.spalten] and name.casefold() in [s.name.casefold() for s in tabelle.spalten]:
                    raise DatenFehler("Der neue Spaltenname ist bereits in anderer Schreibweise vorhanden.")
                ausdruck, parameter = self._ausdruck(tabelle, konfig.get("ausdruck"))
            teile = [f"{ausdruck} AS {_bezeichner(name)}" if s.name == name else _bezeichner(s.name) for s in tabelle.spalten]
            if name not in [s.name for s in tabelle.spalten]:
                teile.append(f"{ausdruck} AS {_bezeichner(name)}")
        return self._neu(f"SELECT {','.join(teile)} FROM ({plan.sql}) AS t", [*parameter,*plan.parameter], plan.quellen)

    def werte(self, tabelle: Tabelle, konfig: dict) -> Tabelle:
        aktion = auswahl(konfig.get("aktion"), ("ERSETZEN", "BEREINIGEN", "LEERWERTE"), "Wertaktion")
        namen = konfig.get("spalten")
        self._namen(tabelle, namen)
        plan = self._plan(tabelle)
        teile, parameter = [], []
        if aktion == "LEERWERTE":
            regel = auswahl(konfig.get("leere", "STOPP"), ("STOPP", "AUSLASSEN", "ERSETZEN"), "Leerwertregel")
            pruefung = " OR ".join(f"{_bezeichner(n)} IS NULL" for n in namen)
            if regel == "STOPP":
                if self._skalar(tabelle, f"count(*) FILTER (WHERE {pruefung})"):
                    raise DatenFehler("Die ausgewählten Spalten enthalten leere Werte.")
                return tabelle
            if regel == "AUSLASSEN":
                return self._neu(f"SELECT * FROM ({plan.sql}) AS t WHERE NOT ({pruefung})", plan.parameter, plan.quellen)
        for s in tabelle.spalten:
            sp = _bezeichner(s.name)
            ausdruck = sp
            if s.name in namen:
                if aktion == "ERSETZEN":
                    suchen, ps = _literal(konfig.get("suchen"))
                    ersatz, pe = _literal(konfig.get("ersatz"))
                    modus = auswahl(konfig.get("modus","GANZER_WERT"),("GANZER_WERT","TEXTTEIL"),"Ersetzen")
                    if modus == "TEXTTEIL":
                        if s.typ != "VARCHAR":
                            raise DatenFehler("Das Ersetzen von Textteilen benötigt Textspalten.")
                        ausdruck = f"replace({sp},CAST({suchen} AS VARCHAR),CAST({ersatz} AS VARCHAR))"
                    else:
                        ausdruck = f"CASE WHEN {sp} IS NOT DISTINCT FROM {suchen} THEN {ersatz} ELSE {sp} END"
                    parameter.extend(ps+pe)
                elif aktion == "LEERWERTE":
                    ersatz, pe = _literal(konfig.get("ersatz"))
                    ausdruck = f"coalesce({sp},{ersatz})"
                    parameter.extend(pe)
                else:
                    if s.typ != "VARCHAR":
                        raise DatenFehler("Textbereinigung benötigt Textspalten.")
                    if konfig.get("trim", True):
                        ausdruck = f"trim({ausdruck})"
                    art = auswahl(konfig.get("grossklein", "UNVERAENDERT"), ("UNVERAENDERT", "GROSS", "KLEIN"), "Schreibweise")
                    if art != "UNVERAENDERT":
                        ausdruck = f"{'upper' if art == 'GROSS' else 'lower'}({ausdruck})"
            teile.append(f"{ausdruck} AS {sp}")
        return self._neu(f"SELECT {','.join(teile)} FROM ({plan.sql}) AS t", [*parameter,*plan.parameter], plan.quellen)

    def sortieren(self, tabelle: Tabelle, schluessel: list) -> Tabelle:
        if not isinstance(schluessel, list) or not schluessel:
            raise DatenFehler("Wähle mindestens eine Sortierspalte.")
        teile = []
        for s in schluessel:
            objekt(s, "Sortierung")
            sp = self._spalte(tabelle, s.get("spalte"))
            richtung = auswahl(s.get("richtung", "ASC"), ("ASC", "DESC"), "Sortierrichtung")
            nullen = auswahl(s.get("nullen", "LETZTE"), ("ERSTE", "LETZTE"), "Leere Werte")
            teile.append(f"{sp} {richtung} NULLS {'FIRST' if nullen=='ERSTE' else 'LAST'}")
        plan = self._plan(tabelle)
        return self._neu(f"SELECT * FROM ({plan.sql}) AS t ORDER BY {','.join(teile)}", plan.parameter, plan.quellen)

    def dubletten(self, tabelle: Tabelle, schluessel: list | None, behalten="ERSTE") -> Tabelle:
        behalten = auswahl(behalten, ("ERSTE", "LETZTE", "KEINE"), "Dublettenregel")
        namen = self._namen(tabelle, schluessel or [s.name for s in tabelle.spalten])
        plan = self._plan(tabelle)
        intern = "sl_rang_" + uuid.uuid4().hex
        nummer = "sl_nr_" + uuid.uuid4().hex
        funktion = f"count(*) OVER (PARTITION BY {','.join(namen)})" if behalten == "KEINE" else f"row_number() OVER (PARTITION BY {','.join(namen)} ORDER BY {_bezeichner(nummer)} {'DESC' if behalten=='LETZTE' else 'ASC'})"
        sql = f"SELECT {','.join(_bezeichner(s.name) for s in tabelle.spalten)} FROM (SELECT *,{funktion} AS {_bezeichner(intern)} FROM (SELECT *,row_number() OVER () AS {_bezeichner(nummer)} FROM ({plan.sql}) AS t) AS n) AS r WHERE {_bezeichner(intern)}=1"
        return self._neu(sql, plan.parameter, plan.quellen)

    def vereinigen(self, links: Tabelle, rechts: Tabelle, *, nach_name=True, dubletten=False) -> Tabelle:
        if type(nach_name) is not bool or type(dubletten) is not bool:
            raise DatenFehler("Wähle gültige Optionen zum Zusammenführen.")
        ln, rn = [s.name for s in links.spalten], [s.name for s in rechts.spalten]
        if (set(ln) != set(rn)) if nach_name else (len(ln) != len(rn)):
            raise DatenFehler("Die Tabellen haben keine passenden Spalten.")
        rechte_namen = ln if nach_name else rn
        for ls, rs in zip(links.spalten, [next(s for s in rechts.spalten if s.name == n) for n in rechte_namen]):
            if ls.typ != rs.typ:
                raise DatenFehler("Die Spaltentypen der Tabellen müssen übereinstimmen.")
        lp, rp = self._plan(links), self._plan(rechts)
        sql = f"SELECT * FROM ({lp.sql}) AS l UNION {'ALL' if not dubletten else ''} SELECT {','.join(_bezeichner(n) for n in rechte_namen)} FROM ({rp.sql}) AS r"
        return self._neu(sql, [*lp.parameter,*rp.parameter], [*lp.quellen,*rp.quellen])

    def verknuepfen(self, links: Tabelle, rechts: Tabelle, konfig: dict) -> Tabelle:
        art = auswahl(konfig.get("art", "LEFT"), ("INNER", "LEFT", "FULL"), "Verknüpfungsart")
        kard = auswahl(konfig.get("kardinalitaet", "n:1"), ("1:1", "1:n", "n:1", "n:m"), "Zuordnung")
        fehlend = auswahl(konfig.get("fehlende_zuordnung", "STOPP"), ("STOPP", "BEHALTEN"), "Fehlende Zuordnung")
        schluessel = konfig.get("schluessel")
        if not isinstance(schluessel, list) or not schluessel:
            raise DatenFehler("Wähle mindestens ein Spaltenpaar zur Verknüpfung.")
        ls, rs = [], []
        for s in schluessel:
            objekt(s, "Spaltenpaar")
            ls.append(self._spalte(links, s.get("links")))
            rs.append(self._spalte(rechts, s.get("rechts")))
            if next(x.typ for x in links.spalten if x.name == s['links']) != next(x.typ for x in rechts.spalten if x.name == s['rechts']):
                raise DatenFehler("Verknüpfte Schlüsselspalten benötigen denselben Datentyp.")
        lp, rp = self._plan(links), self._plan(rechts)
        for tab, p, namen, noetig in ((links, lp, ls, kard in ("1:1", "1:n")), (rechts, rp, rs, kard in ("1:1", "n:1"))):
            if noetig:
                self._pruefe_quellen(p)
                sql = f"SELECT EXISTS(SELECT 1 FROM ({p.sql}) AS t GROUP BY {','.join(namen)} HAVING count(*)>1)"
                if self._verbindung.execute(sql, list(p.parameter)).fetchone()[0]:
                    raise DatenFehler("Die gewählte Zuordnung ist nicht eindeutig. Prüfe doppelte Schlüssel oder ändere die Kardinalität.")
        on = ' AND '.join(f"l.{l}=r.{r}" for l,r in zip(ls,rs))
        if fehlend == "STOPP":
            self._pruefe_quellen(lp)
            self._pruefe_quellen(rp)
            sql = f"SELECT count(*) FROM ({lp.sql}) AS l WHERE NOT EXISTS(SELECT 1 FROM ({rp.sql}) AS r WHERE {on})"
            anzahl = self._verbindung.execute(sql, [*lp.parameter,*rp.parameter]).fetchone()[0]
            if anzahl:
                raise DatenFehler(f"Für {anzahl} Zeilen fehlt eine passende Zuordnung in der rechten Tabelle.")
        verwendet = {s.name.casefold() for s in links.spalten}
        teile = [f"l.{_bezeichner(s.name)}" for s in links.spalten]
        for s in rechts.spalten:
            name, nr = s.name, 1
            while name.casefold() in verwendet:
                name = f"{s.name}_rechts" + (str(nr) if nr > 1 else "")
                nr += 1
            verwendet.add(name.casefold())
            teile.append(f"r.{_bezeichner(s.name)} AS {_bezeichner(name)}")
        return self._neu(f"SELECT {','.join(teile)} FROM ({lp.sql}) AS l {art} JOIN ({rp.sql}) AS r ON {on}", [*lp.parameter,*rp.parameter], [*lp.quellen,*rp.quellen])

    def gruppieren(self, tabelle: Tabelle, schluessel: list, aggregate: list) -> Tabelle:
        namen = self._namen(tabelle, schluessel, leer=True)
        if not isinstance(aggregate, list) or not aggregate:
            raise DatenFehler("Wähle mindestens eine Auswertung.")
        teile = list(namen)
        ergebnisse = list(schluessel)
        for a in aggregate:
            objekt(a, "Auswertung")
            funktion = auswahl(a.get("funktion"), AGGREGATE, "Auswertung")
            name = text(a.get("name"), "Ergebnisspalte")
            if name.casefold() in [n.casefold() for n in ergebnisse]:
                raise DatenFehler("Ergebnisspalten müssen eindeutig sein.")
            ergebnisse.append(name)
            spalte = "*" if funktion == "ANZAHL" else self._spalte(tabelle, a.get("spalte"))
            agg = dict(ANZAHL="count", ANZAHL_WERTE="count", ANZAHL_VERSCHIEDEN="count", SUMME="sum", MITTELWERT="avg", MIN="min", MAX="max")[funktion]
            teile.append(f"{agg}({'DISTINCT ' if funktion=='ANZAHL_VERSCHIEDEN' else ''}{spalte}) AS {_bezeichner(name)}")
        plan = self._plan(tabelle)
        return self._neu(f"SELECT {','.join(teile)} FROM ({plan.sql}) AS t" + (f" GROUP BY {','.join(namen)}" if namen else ""), plan.parameter, plan.quellen)

    def pruefen(self, tabelle: Tabelle, regeln: list, bei_fehler="BERICHT") -> dict:
        auswahl(bei_fehler, ("BERICHT", "STOPP"), "Prüfergebnis")
        if not isinstance(regeln, list) or not 1 <= len(regeln) <= 100:
            raise DatenFehler("Wähle zwischen 1 und 100 Prüfregeln.")
        ergebnisse = []
        gesamt = self._skalar(tabelle, "count(*)")
        for regel in regeln:
            objekt(regel, "Prüfregel")
            art = auswahl(regel.get("art"), PRUEFUNGEN, "Prüfregel")
            parameter = []
            if art == "ZEILENZAHL":
                minimum, maximum = regel.get("minimum", 0), regel.get("maximum")
                if type(minimum) is not int or minimum < 0 or (maximum is not None and (type(maximum) is not int or maximum < minimum)):
                    raise DatenFehler("Der erlaubte Zeilenbereich ist ungültig.")
                fehler = int(gesamt < minimum or maximum is not None and gesamt > maximum)
            elif art == "EINDEUTIG":
                namen = self._namen(tabelle, regel.get("spalten", [regel.get("spalte")]))
                plan = self._plan(tabelle)
                fehler = self._verbindung.execute(f"SELECT coalesce(sum(n-1),0) FROM (SELECT count(*) AS n FROM ({plan.sql}) AS t GROUP BY {','.join(namen)} HAVING count(*)>1) AS d", list(plan.parameter)).fetchone()[0]
            else:
                spalte = self._spalte(tabelle, regel.get("spalte"))
                if art == "NICHT_LEER":
                    wo = f"{spalte} IS NULL OR trim(CAST({spalte} AS VARCHAR))=''"
                elif art == "TYP":
                    wo = f"{spalte} IS NOT NULL AND TRY_CAST({spalte} AS {_typ(regel.get('typ'))}) IS NULL"
                elif art == "ERLAUBTE_WERTE":
                    ast = {"art":"vergleich", "spalte":regel.get("spalte"), "operator":"IN", "werte":regel.get("werte")}
                    sql, parameter = self._bedingung(tabelle, ast)
                    wo = f"{spalte} IS NOT NULL AND NOT ({sql})"
                else:
                    teile = []
                    for grenze, operator in (("minimum", "<"), ("maximum", ">")):
                        if grenze in regel:
                            sql, ps = _literal(regel[grenze])
                            teile.append(f"{spalte}{operator}{sql}")
                            parameter.extend(ps)
                    if not teile:
                        raise DatenFehler("Gib mindestens eine Bereichsgrenze an.")
                    wo = ' OR '.join(teile)
                # Parameter erscheinen vor dem Quellplan im SELECT-Ausdruck.
                fehler = self._skalar(tabelle, f"count(*) FILTER (WHERE {wo})", parameter)
            ergebnisse.append({"art": art, "fehler": int(fehler), "gueltig": fehler == 0})
        bericht = {"gueltig": all(e["gueltig"] for e in ergebnisse), "zeilen": gesamt, "regeln": ergebnisse}
        self.kontext.ereignis("daten.pruefung", **bericht)
        if bei_fehler == "STOPP" and not bericht["gueltig"]:
            raise DatenFehler("Die Daten erfüllen die gewählten Qualitätsregeln nicht. Prüfe den Prüfbericht.")
        return bericht

    def vorschau(self, tabelle: Tabelle, limit=200, byte_limit=262144) -> dict:
        if type(limit) is not int or not 1 <= limit <= min(200,self._grenze("max_vorschau_zeilen",200)):
            raise DatenFehler("Die Vorschau zeigt zwischen 1 und 200 Zeilen.")
        plan = self._plan(tabelle)
        if type(byte_limit) is not int or byte_limit <= 0:
            raise DatenFehler("Das Bytebudget der Vorschau ist ungültig.")
        byte_limit = min(byte_limit,self._grenze("max_vorschau_bytes",262144))
        # Eine einzelne riesige Textzelle darf bereits beim Fetch nicht das
        # Ausgabe-Bytebudget um Größenordnungen übersteigen. Die SQL-Projektion
        # begrenzt Text vor der Übergabe an Python; Kürzung wird offen gemeldet.
        zellgrenze = max(1,min(4096,byte_limit//max(6*len(tabelle.spalten),1)))
        texte = [s for s in tabelle.spalten if s.typ == "VARCHAR"]
        teile = [f"CASE WHEN length({_bezeichner(s.name)})>{zellgrenze} THEN substr({_bezeichner(s.name)},1,{zellgrenze}) || '…' ELSE {_bezeichner(s.name)} END AS {_bezeichner(s.name)}" if s.typ == "VARCHAR" else _bezeichner(s.name) for s in tabelle.spalten]
        marker = "sl_gekuerzt_"+uuid.uuid4().hex
        warnung = " OR ".join(f"coalesce(length({_bezeichner(s.name)})>{zellgrenze},false)" for s in texte) or "false"
        klein = self._neu(f"SELECT {','.join(teile)},({warnung}) AS {_bezeichner(marker)} FROM ({plan.sql}) AS t LIMIT {limit+1}",plan.parameter,plan.quellen,max_spalten=16385)
        zeilen, groesse, gekuerzt = [], 0, False
        zellen_gekuerzt = False
        start = time.monotonic()
        iterator = self.zeilen(klein, 1, zeitlimit=min(5,self._grenze("zeitlimit_sekunden",120)))
        try:
            for zeile in iterator:
                werte = [_jsonwert(w) for w in zeile[:-1]]
                laenge = len(json.dumps(werte, ensure_ascii=True).encode("utf-8"))
                if len(zeilen) == limit or groesse + laenge > byte_limit:
                    gekuerzt = True
                    break
                zeilen.append(werte)
                zellen_gekuerzt = zellen_gekuerzt or zeile[-1]
                groesse += laenge
        finally:
            iterator.close()
            self._plaene.pop(klein.id, None)
            self._tabellen.pop(klein.id, None)
        self._pruefe_quellen(plan)
        return {"spalten": tabelle.json()["spalten"], "zeilen": zeilen, "gekuerzt": gekuerzt or zellen_gekuerzt, "zellen_gekuerzt": zellen_gekuerzt, "bytes": groesse, "dauer_ms": round((time.monotonic()-start)*1000)}

    def exportieren(self, tabelle: Tabelle, format: str, pfad: str, konflikt="FEHLER", optionen: dict | None = None) -> dict:
        from .dateien import publizieren
        auswahl(format, FORMATE, "Dateiformat")
        auswahl(konflikt, KONFLIKTE, "Vorhandenes Ziel")
        optionen = objekt(optionen or {}, "Exportoptionen")
        plan = self._plan(tabelle)
        self._pruefe_quellen(plan)
        ziel = self.kontext.pfad(pfad, schreiben=True)
        return publizieren(self.kontext, ziel, lambda staging: self._schreiben(tabelle, format, Path(staging), optionen), konflikt=konflikt, quellen=[q.pfad for q in plan.quellen])

    def _schreiben(self, tabelle: Tabelle, format: str, pfad: Path, optionen: dict) -> None:
        if self._skalar(tabelle,"count(*)") > self._grenze("max_zeilen",100000):
            raise DatenFehler("Die Ausgabe überschreitet die eingestellte Zeilengrenze.")
        namen = [s.name for s in tabelle.spalten]
        if format == "CSV":
            kodierung = auswahl(optionen.get("kodierung", "UTF8_BOM"), ("UTF8", "UTF8_BOM", "UTF16", "CP1252"), "Kodierung")
            trenner = optionen.get("trennzeichen", ";")
            if not isinstance(trenner, str) or len(trenner) != 1 or trenner in ('\r','\n','"','\x00'):
                raise DatenFehler("Wähle genau ein zulässiges Trennzeichen.")
            with pfad.open("w", encoding=dict(UTF8="utf-8",UTF8_BOM="utf-8-sig",UTF16="utf-16",CP1252="cp1252")[kodierung], errors="strict", newline="") as datei:
                writer = csv.writer(datei, delimiter=trenner, lineterminator="\r\n")
                if optionen.get("kopfzeile", True):
                    writer.writerow(namen)
                nullwert = text(optionen.get("nullwert", ""), "Leerwert", leer=True)
                for zeile in self.zeilen(tabelle):
                    writer.writerow([nullwert if w is None else _jsonwert(w) for w in zeile])
        elif format == "JSON":
            newline = optionen.get("zeilenformat", False)
            with pfad.open("w", encoding="utf-8", errors="strict", newline="") as datei:
                if not newline:
                    datei.write('[')
                erste = True
                for zeile in self.zeilen(tabelle):
                    if not erste and not newline:
                        datei.write(',\r\n')
                    json.dump(dict(zip(namen, map(_jsonwert, zeile))), datei, ensure_ascii=False, allow_nan=False)
                    if newline:
                        datei.write('\r\n')
                    erste = False
                if not newline:
                    datei.write(']\r\n')
        elif format == "XLSX":
            self._xlsx_schreiben(tabelle, pfad, optionen)
        else:
            plan = self._plan(tabelle)
            self._pruefe_quellen(plan)
            # DuckDB erhält den Pfad über seine Parameterbindung; kein SQL-Literal
            # aus einem Benutzerfeld wird in einen COPY-Befehl eingesetzt.
            self._verbindung.execute(f"COPY ({plan.sql}) TO ? (FORMAT PARQUET, COMPRESSION ZSTD)", [*plan.parameter, str(pfad)])

    def _xlsx_schreiben(self, tabelle: Tabelle, pfad: Path, optionen: dict) -> None:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        blatt = text(optionen.get("blatt", "Daten"), "Tabellenblatt")
        if len(blatt) > 31 or any(z in blatt for z in "[]:*?/\\"):
            raise DatenFehler("Der Tabellenblattname ist ungültig oder länger als 31 Zeichen.")
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(blatt)
        def zellen(werte):
            for wert in werte:
                # XLSX kann Dezimalwerte >15 Stellen nicht verlustfrei numerisch
                # darstellen; Dezimalzahlen werden daher bewusst Textzellen.
                if isinstance(wert, Decimal):
                    wert = str(wert)
                if isinstance(wert, str):
                    if len(wert) > 32767:
                        raise DatenFehler("Ein Text überschreitet die Excel-Zellgrenze von 32767 Zeichen.")
                    zelle = WriteOnlyCell(ws, value=wert)
                    zelle.data_type = "s"
                    yield zelle
                else:
                    yield wert
        try:
            ws.append(list(zellen(s.name for s in tabelle.spalten)))
            for nr, zeile in enumerate(self.zeilen(tabelle), 2):
                if nr > 1048576:
                    raise DatenFehler("Die Tabelle überschreitet die Excel-Zeilengrenze. Verwende CSV oder Parquet.")
                ws.append(list(zellen(zeile)))
            wb.save(pfad)
        except BaseException:
            # Write-only-Blätter halten einen XML-Stream und technische Tempdateien.
            # Auch bei einem fachlichen Abbruch werden sie über öffentliche APIs
            # geschlossen; die zentrale Publikation verwirft diese Staging-Datei.
            try:
                ws.close()
                wb.save(pfad)
            except Exception:
                pass
            raise
        finally:
            wb.close()


def _xlsx_text(wert: Any) -> str:
    if isinstance(wert, (date, datetime)):
        return wert.isoformat()
    if isinstance(wert, bool):
        return "true" if wert else "false"
    return str(wert)


def engine(kontext) -> DatenEngine:
    if "daten" not in kontext.dienste:
        kontext.dienste["daten"] = DatenEngine(kontext)
    return kontext.dienste["daten"]


def _handler(operation):
    def ausfuehren(kontext, konfiguration, eingaben):
        k, e = objekt(konfiguration, "Konfiguration"), objekt(eingaben, "Eingaben")
        d = engine(kontext)
        try:
            if operation == "importieren":
                ergebnis = d.importieren(k.get("format"), k.get("pfad"), k.get("optionen"))
            elif operation in ("vereinigen", "verknuepfen"):
                links, rechts = d.tabelle(e.get("links")), d.tabelle(e.get("rechts"))
                ergebnis = d.vereinigen(links,rechts,nach_name=k.get("nach_name",True),dubletten=k.get("dubletten",False)) if operation == "vereinigen" else d.verknuepfen(links,rechts,k)
            else:
                t = d.tabelle(e.get("tabelle"))
                if operation == "exportieren": ergebnis = d.exportieren(t,k.get("format"),k.get("pfad"),k.get("konflikt","FEHLER"),k.get("optionen"))
                elif operation == "spalten": ergebnis = d.spalten(t,k)
                elif operation == "werte": ergebnis = d.werte(t,k)
                elif operation == "filtern": ergebnis = d.filtern(t,k.get("bedingung"))
                elif operation == "sortieren": ergebnis = d.sortieren(t,k.get("schluessel"))
                elif operation == "dubletten": ergebnis = d.dubletten(t,k.get("schluessel"),k.get("behalten","ERSTE"))
                elif operation == "gruppieren": ergebnis = d.gruppieren(t,k.get("schluessel",[]),k.get("aggregate"))
                elif operation == "pruefen": ergebnis = d.pruefen(t,k.get("regeln"),k.get("bei_fehler","BERICHT"))
                else: ergebnis = d.vorschau(t,k.get("limit",200))
            return ergebnis.json() if isinstance(ergebnis,Tabelle) else ergebnis
        except duckdb.Error as fehler:
            raise DatenFehler("Die Datenverarbeitung ist fehlgeschlagen. Prüfe Werte, Spalten und Datentypen.") from fehler
    return ausfuehren


HANDLER = {f"daten.{op}": _handler(op) for op in ("importieren","exportieren","spalten","werte","filtern","sortieren","dubletten","vereinigen","verknuepfen","gruppieren","pruefen","vorschau")}
