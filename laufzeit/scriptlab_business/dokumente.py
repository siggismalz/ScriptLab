"""Begrenzte Dokumentoperationen; Bibliotheken starten weder Office noch einen Viewer."""
from __future__ import annotations

from copy import copy, deepcopy
from datetime import date, datetime
from decimal import Decimal
import hashlib
import re
from pathlib import Path
from xml.sax.saxutils import escape

from .kontext import BusinessFehler, Kontext
from .dateien import publizieren, quellkopie, zip_pruefen, datei_fingerabdruck, quellen_pruefen


def _ziel(k, endung):
    wert = k.get("ziel")
    if not isinstance(wert, str) or Path(wert).suffix.lower() != endung:
        raise BusinessFehler(f"Wähle eine Ausgabedatei mit der Endung {endung}.")
    return wert


def _tabelle(kontext, eingaben):
    dienst = kontext.dienste.get("daten")
    if not dienst or "tabelle" not in eingaben:
        raise BusinessFehler("Verbinde eine vorherige Datentabelle mit diesem Dokumentbaustein.")
    tabelle = dienst.tabelle(eingaben["tabelle"])
    return dienst, tabelle


def _zeilen(kontext, dienst, tabelle):
    namen = [spalte.name for spalte in tabelle.spalten]
    for nummer, zeile in enumerate(dienst.zeilen(tabelle), 1):
        kontext.pruefen()
        if nummer > kontext.grenzen.max_zeilen:
            raise BusinessFehler("Die Dokumentdaten überschreiten die gewählte Zeilenbegrenzung.")
        yield dict(zip(namen, zeile))


def _text(wert):
    if wert is None:
        return ""
    if isinstance(wert, (datetime, date)):
        return wert.isoformat()
    if isinstance(wert, (dict, list)):
        raise BusinessFehler("Ein Dokumentfeld benötigt einen einzelnen Text- oder Zahlenwert.")
    return str(wert)


def _office_pruefer(kontext):
    return lambda pfad: zip_pruefen(kontext, pfad)


def xlsx_vorlage(kontext: Kontext, k: dict, e: dict):
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import PatternFill
    from openpyxl.utils.cell import column_index_from_string
    from openpyxl.workbook.properties import CalcProperties
    ziel = _ziel(k, ".xlsx")
    quelle = kontext.pfad(k["vorlage"])
    if quelle.suffix.lower() != ".xlsx":
        raise BusinessFehler("Excel-Vorlagen müssen normale .xlsx-Dateien sein; Makrodateien werden nicht geöffnet.")
    def schreiben(temp):
        with quellkopie(kontext, quelle, office=True) as kopie:
            buch = load_workbook(kopie, data_only=False, keep_links=False)
            try:
                blattname = k.get("blatt") or buch.sheetnames[0]
                if blattname not in buch.sheetnames:
                    raise BusinessFehler("Das gewählte Tabellenblatt fehlt in der Vorlage.")
                blatt = buch[blattname]
                zellen, formeln, formate = k.get("zellen", {}), k.get("formeln", {}), k.get("formate", [])
                if not isinstance(zellen, dict) or not isinstance(formeln, dict) or set(zellen) & set(formeln):
                    raise BusinessFehler("Zellwerte und ausdrücklich gewählte Formeln müssen getrennt zugeordnet sein.")
                if not isinstance(formate, list) or len(zellen) + len(formeln) + len(formate) > kontext.grenzen.max_dokument_zellen:
                    raise BusinessFehler("Die Excel-Zellzuweisungen überschreiten die Dokumentzellengrenze.")
                def zelle_lesen(adresse):
                    treffer = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", adresse) if isinstance(adresse, str) else None
                    if not treffer:
                        raise BusinessFehler("Eine Excel-Zelladresse ist ungültig, beispielsweise A1.")
                    if int(treffer[2]) > 1048576 or column_index_from_string(treffer[1]) > 16384:
                        raise BusinessFehler("Die Zielzelle liegt außerhalb des Excel-Bereichs.")
                    zelle = blatt[adresse]
                    if isinstance(zelle, MergedCell):
                        raise BusinessFehler("Die Zielzelle liegt außerhalb des Excel-Bereichs oder innerhalb verbundener Zellen.")
                    return zelle
                for adresse, wert in [*zellen.items(), *formeln.items()]:
                    kontext.pruefen()
                    zelle = zelle_lesen(adresse)
                    if adresse in formeln:
                        if not isinstance(wert, str) or not wert.startswith("="):
                            raise BusinessFehler("Eine ausdrücklich gesetzte Excel-Formel muss mit = beginnen.")
                        zelle.value = wert
                    else:
                        if isinstance(wert, (dict, list)):
                            raise BusinessFehler("Eine Excel-Zelle benötigt einen einzelnen Wert.")
                        zelle.value = str(wert) if isinstance(wert, Decimal) else wert
                        if isinstance(wert, str):
                            zelle.data_type = "s"  # Fremde Texte werden niemals zu Formeln.
                for format in formate:
                    kontext.pruefen()
                    if not isinstance(format, dict) or "adresse" not in format or set(format) - {"adresse", "zahlformat", "fett", "fuellung", "ausrichtung"}:
                        raise BusinessFehler("Ein Excel-Format benötigt eine Zelladresse und ausdrücklich erlaubte Formatangaben.")
                    zelle = zelle_lesen(format["adresse"])
                    if "zahlformat" in format:
                        wert = format["zahlformat"]
                        if not isinstance(wert, str) or not 1 <= len(wert) <= 256 or re.search(r"[\x00-\x1f]", wert):
                            raise BusinessFehler("Das Excel-Zahlenformat muss ein kurzer gültiger Formattext sein.")
                        zelle.number_format = wert
                    if "fett" in format:
                        if type(format["fett"]) is not bool:
                            raise BusinessFehler("Die Excel-Einstellung Fett benötigt wahr oder falsch.")
                        schrift = copy(zelle.font)
                        schrift.bold = format["fett"]
                        zelle.font = schrift
                    if "fuellung" in format:
                        farbe = format["fuellung"]
                        if not isinstance(farbe, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}", farbe):
                            raise BusinessFehler("Die Excel-Füllfarbe benötigt genau sechs Hexzeichen, beispielsweise E8EBEF.")
                        zelle.fill = PatternFill(fill_type="solid", fgColor="FF" + farbe.upper())
                    if "ausrichtung" in format:
                        wert = format["ausrichtung"]
                        if not isinstance(wert, str) or wert not in {"links", "mitte", "rechts"}:
                            raise BusinessFehler("Wähle links, mitte oder rechts als Excel-Ausrichtung.")
                        ausrichtung = copy(zelle.alignment)
                        ausrichtung.horizontal = {"links": "left", "mitte": "center", "rechts": "right"}[wert]
                        zelle.alignment = ausrichtung
                buch.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False)
                buch.save(temp)
            finally:
                buch.close()
    return {**publizieren(kontext, ziel, schreiben, k.get("konflikt", "FEHLER"), [quelle], _office_pruefer(kontext)), "dokumenttyp": "XLSX", "formeln_berechnet": False}


PLATZHALTER = re.compile(r"\{\{([^{}\r\n]{1,160})\}\}")


def _absatz_ersetzen(absatz, werte):
    laeufe = list(absatz.runs)
    gesamt = "".join(lauf.text for lauf in laeufe)
    rest = PLATZHALTER.sub("", gesamt)
    if "{{" in rest or "}}" in rest:
        raise BusinessFehler("Die Word-Vorlage enthält einen unvollständigen Platzhalter.")
    for treffer in reversed(list(PLATZHALTER.finditer(gesamt))):
        schluessel = treffer.group(1).strip()
        if schluessel not in werte:
            raise BusinessFehler(f"Der Word-Platzhalter hat keinen Wert: {schluessel}")
        anfang, ende = treffer.span()
        position, erster, letzter = 0, None, None
        for i, lauf in enumerate(laeufe):
            grenze = position + len(lauf.text)
            if erster is None and anfang < grenze:
                erster = (i, anfang - position)
            if ende <= grenze:
                letzter = (i, ende - position)
                break
            position = grenze
        if erster is None or letzter is None:
            raise BusinessFehler("Der Word-Platzhalter kann nicht eindeutig ersetzt werden.")
        a, offset_a = erster
        b, offset_b = letzter
        ersatz = _text(werte[schluessel])
        if a == b:
            laeufe[a].text = laeufe[a].text[:offset_a] + ersatz + laeufe[a].text[offset_b:]
        else:
            suffix = laeufe[b].text[offset_b:]
            laeufe[a].text = laeufe[a].text[:offset_a] + ersatz
            for i in range(a + 1, b):
                laeufe[i].text = ""
            laeufe[b].text = suffix


def _absatzliste(teil):
    yield from teil.paragraphs
    for tabelle in teil.tables:
        gesehen = set()
        for zeile in tabelle.rows:
            for zelle in zeile.cells:
                if zelle._tc not in gesehen:
                    gesehen.add(zelle._tc)
                    yield from _absatzliste(zelle)


def _word_ersetzen(dokument, werte):
    for absatz in _absatzliste(dokument):
        _absatz_ersetzen(absatz, werte)
    for abschnitt in dokument.sections:
        for teil in (abschnitt.header, abschnitt.footer, abschnitt.first_page_header,
                     abschnitt.first_page_footer, abschnitt.even_page_header, abschnitt.even_page_footer):
            for absatz in _absatzliste(teil):
                _absatz_ersetzen(absatz, werte)


def _positionen(kontext, dokument, konfiguration, zeilen, gemeinsame_werte):
    from docx.table import _Row
    index_tabelle, index_zeile = konfiguration.get("tabelle", 0), konfiguration.get("zeile", 1)
    if type(index_tabelle) is not int or type(index_zeile) is not int or index_tabelle < 0 or index_zeile < 0:
        raise BusinessFehler("Wähle eine gültige Word-Positionstabelle und Musterzeile.")
    try:
        tabelle = dokument.tables[index_tabelle]
        muster = tabelle.rows[index_zeile]
    except IndexError as fehler:
        raise BusinessFehler("Die Word-Positionstabelle oder ihre Musterzeile fehlt.") from fehler
    vorlage = deepcopy(muster._tr)
    anker = muster._tr
    praefix = konfiguration.get("praefix", "Position")
    for datensatz in zeilen:
        kontext.pruefen()
        neu = deepcopy(vorlage)
        anker.addprevious(neu)
        zeile = _Row(neu, tabelle)
        werte = {**gemeinsame_werte, **{f"{praefix}.{name}": wert for name, wert in datensatz.items()}}
        for zelle in zeile.cells:
            for absatz in _absatzliste(zelle):
                _absatz_ersetzen(absatz, werte)
    anker.getparent().remove(anker)


def docx_vorlage(kontext: Kontext, k: dict, e: dict):
    from docx import Document
    ziel = _ziel(k, ".docx")
    quelle = kontext.pfad(k["vorlage"])
    if quelle.suffix.lower() != ".docx" or not isinstance(k.get("werte", {}), dict):
        raise BusinessFehler("Wähle eine .docx-Vorlage und benannte Platzhalterwerte.")
    dienst, tabelle = _tabelle(kontext, e) if "positionen" in k and "positions_zeilen" not in e else (None, None)
    quellen = [quelle, *e.get("quellen", ()), *(dienst.quellpfade(tabelle) if dienst else ())]
    def schreiben(temp):
        with quellkopie(kontext, quelle, office=True) as kopie:
            dokument = Document(kopie)
            if "positionen" in k:
                zeilen = _zeilen(kontext, dienst, tabelle) if dienst else e.get("positions_zeilen", [])
                if not dienst and (not isinstance(zeilen, list) or len(zeilen) > kontext.grenzen.max_zeilen or any(not isinstance(z, dict) for z in zeilen)):
                    raise BusinessFehler("Die Word-Positionsdaten sind ungültig oder zu groß.")
                _positionen(kontext, dokument, k["positionen"], zeilen, k.get("werte", {}))
            _word_ersetzen(dokument, k.get("werte", {}))
            dokument.save(temp)
    return {**publizieren(kontext, ziel, schreiben, k.get("konflikt", "FEHLER"), quellen, _office_pruefer(kontext)), "dokumenttyp": "DOCX"}


def _dateiname(muster, werte, nummer, endung):
    if not isinstance(muster, str) or "/" in muster or "\\" in muster:
        raise BusinessFehler("Das Dateinamensmuster darf keinen Ordnerpfad enthalten.")
    def ersetzen(treffer):
        name = treffer.group(1)
        if name != "nummer" and name not in werte:
            raise BusinessFehler(f"Dem Dateinamensmuster fehlt der Wert {name}.")
        text = str(nummer if name == "nummer" else werte[name])
        return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text).strip(" .") or "leer"
    name = re.sub(r"\{([^{}]+)\}", ersetzen, muster)
    if not name or "{" in name or "}" in name or re.search(r'[<>:"/\\|?*\x00-\x1f]', name) or name.endswith((".", " ")) or len(name) > 200 or Path(name).suffix.lower() != endung:
        raise BusinessFehler("Das Dateinamensmuster ergibt keinen gültigen Windows-Dateinamen.")
    if re.fullmatch(r"(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])", Path(name).stem):
        name = "_" + name
    return name


def docx_serie(kontext: Kontext, k: dict, e: dict):
    dienst, tabelle = _tabelle(kontext, e)
    ausgaben = []
    if "gruppe_nach" in k:
        schluessel = k["gruppe_nach"]
        schluessel = [schluessel] if isinstance(schluessel, str) else schluessel
        namen = {s.name for s in tabelle.spalten}
        if not isinstance(schluessel, list) or not schluessel or any(s not in namen for s in schluessel) or "positionen" not in k:
            raise BusinessFehler("Wähle vorhandene Gruppenspalten und eine Word-Positionstabelle.")
        gruppen = {}
        zellen = 0
        for datensatz in _zeilen(kontext, dienst, tabelle):
            zellen += len(datensatz)
            if zellen > kontext.grenzen.max_dokument_zellen:
                raise BusinessFehler("Die Word-Seriendaten überschreiten die Dokumentzellengrenze.")
            key = tuple(datensatz[s] for s in schluessel)
            gruppen.setdefault(key, []).append(datensatz)
        for nummer, positionen in enumerate(gruppen.values(), 1):
            if nummer > kontext.grenzen.max_ausgaben:
                raise BusinessFehler("Die Word-Serie enthält zu viele Gruppen.")
            werte = {name: positionen[0][name] for name in namen if all(z[name] == positionen[0][name] for z in positionen)}
            for name, spalte in k.get("summen", {}).items():
                if spalte not in namen:
                    raise BusinessFehler("Eine Summenspalte fehlt in den Word-Seriendaten.")
                try:
                    zahlen = [Decimal(str(z[spalte])) for z in positionen]
                    if not all(zahl.is_finite() for zahl in zahlen):
                        raise ValueError()
                    werte[name] = sum(zahlen, Decimal(0))
                except (ValueError, ArithmeticError) as fehler:
                    raise BusinessFehler("Eine Word-Gruppensumme benötigt gültige Zahlen.") from fehler
            werte["AnzahlPositionen"] = len(positionen)
            name = _dateiname(k.get("dateiname", "Dokument-{nummer}.docx"), werte, nummer, ".docx")
            config = {"vorlage": k["vorlage"], "ziel": str(kontext.pfad(k["zielordner"]) / name), "werte": {**k.get("werte", {}), **werte}, "positionen": k["positionen"], "konflikt": k.get("konflikt", "FEHLER")}
            ausgaben.append(docx_vorlage(kontext, config, {"positions_zeilen": positionen, "quellen": dienst.quellpfade(tabelle)}))
        return {"dokumente": ausgaben, "anzahl": len(ausgaben)}
    for nummer, datensatz in enumerate(_zeilen(kontext, dienst, tabelle), 1):
        if nummer > kontext.grenzen.max_ausgaben:
            raise BusinessFehler("Die Dokumentserie enthält mehr Ausgaben als erlaubt.")
        name = _dateiname(k.get("dateiname", "Dokument-{nummer}.docx"), datensatz, nummer, ".docx")
        konfiguration = {"vorlage": k["vorlage"], "ziel": str(kontext.pfad(k["zielordner"]) / name), "werte": {**k.get("werte", {}), **datensatz}, "konflikt": k.get("konflikt", "FEHLER")}
        ausgaben.append(docx_vorlage(kontext, konfiguration, {"quellen": dienst.quellpfade(tabelle)}))
    return {"dokumente": ausgaben, "anzahl": len(ausgaben)}


def _pdf_lesen(kontext, pfad):
    from pypdf import PdfReader
    datei_fingerabdruck(kontext, pfad)
    try:
        leser = PdfReader(pfad, strict=True)
        if leser.is_encrypted:
            raise BusinessFehler("Verschlüsselte PDF-Dateien werden nicht geöffnet.")
        if len(leser.pages) > kontext.grenzen.max_seiten:
            raise BusinessFehler("Das PDF überschreitet die Seitenbegrenzung.")
        return leser
    except BusinessFehler:
        raise
    except Exception as fehler:
        raise BusinessFehler("Die PDF-Datei ist beschädigt oder nicht unterstützt.") from fehler


def seiten_lesen(text, anzahl):
    if text is None or text == "":
        return list(range(anzahl))
    if not isinstance(text, str) or len(text) > 20000:
        raise BusinessFehler("Gib PDF-Seiten beispielsweise als 1-3,5 an.")
    ergebnis = []
    for teil in text.split(","):
        if not re.fullmatch(r"\s*[1-9][0-9]*(\s*-\s*[1-9][0-9]*)?\s*", teil):
            raise BusinessFehler("Die PDF-Seitenauswahl ist ungültig.")
        zahlen = [int(z) for z in teil.replace(" ", "").split("-")]
        a, b = zahlen[0], zahlen[-1]
        if a > b or b > anzahl:
            raise BusinessFehler("Eine PDF-Seite fehlt oder der Bereich ist umgekehrt.")
        for seite in range(a - 1, b):
            if seite in ergebnis:
                raise BusinessFehler("Eine PDF-Seite wurde doppelt ausgewählt.")
            ergebnis.append(seite)
    return ergebnis


def _pdf_schreiben(kontext, ziel, quellen, auswahl, konflikt="FEHLER", ausgabe_id=None):
    from pypdf import PdfWriter
    def schreiben(temp):
        schreiber = PdfWriter()
        try:
            for quelle, seiten in zip(quellen, auswahl):
                with quellkopie(kontext, quelle) as kopie:
                    leser = _pdf_lesen(kontext, kopie)
                    for seite in seiten_lesen(seiten, len(leser.pages)):
                        kontext.pruefen()
                        schreiber.add_page(leser.pages[seite], excluded_keys=("/AA",))
                        if len(schreiber.pages) > kontext.grenzen.max_seiten:
                            raise BusinessFehler("Die PDF-Ausgabe enthält zu viele Seiten.")
            schreiber.write(temp)
        finally:
            schreiber.close()
    return {**publizieren(kontext, ziel, schreiben, konflikt, quellen, lambda pfad: _pdf_lesen(kontext, pfad), ausgabe_id), "dokumenttyp": "PDF"}


def _pfade(k, e):
    quellen = k.get("quellen", e.get("dateien"))
    if isinstance(quellen, dict):
        quellen = quellen.get("dateien", quellen.get("dokumente"))
    if not isinstance(quellen, list) or not quellen:
        raise BusinessFehler("Wähle mindestens eine PDF-Quelldatei.")
    return [q if isinstance(q, str) else q["pfad"] for q in quellen]


def pdf_verbinden(kontext: Kontext, k: dict, e: dict):
    return _pdf_schreiben(kontext, _ziel(k, ".pdf"), [kontext.pfad(p) for p in _pfade(k, e)], [None] * len(_pfade(k, e)), k.get("konflikt", "FEHLER"))


def pdf_seiten(kontext: Kontext, k: dict, e: dict):
    return _pdf_schreiben(kontext, _ziel(k, ".pdf"), [kontext.pfad(k["quelle"])], [k["seiten"]], k.get("konflikt", "FEHLER"))


def pdf_teilen(kontext: Kontext, k: dict, e: dict):
    bereiche = k.get("bereiche")
    if not isinstance(bereiche, list) or not bereiche or len(bereiche) > kontext.grenzen.max_ausgaben:
        raise BusinessFehler("Lege eine begrenzte Liste von PDF-Seitenbereichen fest.")
    ausgaben = []
    for nummer, bereich in enumerate(bereiche, 1):
        name = _dateiname(k.get("dateiname", "Teil-{nummer}.pdf"), {}, nummer, ".pdf")
        ausgaben.append(_pdf_schreiben(kontext, kontext.pfad(k["zielordner"]) / name, [kontext.pfad(k["quelle"])], [bereich], k.get("konflikt", "FEHLER"), f"pdf-{nummer}"))
    return {"dokumente": ausgaben, "anzahl": len(ausgaben)}


def pdf_text(kontext: Kontext, k: dict, e: dict):
    bei_scan = k.get("bei_scan", "STOPP")
    if bei_scan not in {"STOPP", "BERICHT"}:
        raise BusinessFehler("Wähle STOPP oder BERICHT für PDF-Dateien ohne verlässliche Textschicht.")
    quelle = kontext.pfad(k["quelle"])
    with quellkopie(kontext, quelle) as kopie:
        leser = _pdf_lesen(kontext, kopie)
        seiten, gesamt = [], 0
        for index in seiten_lesen(k.get("seiten"), len(leser.pages)):
            kontext.pruefen()
            seite = leser.pages[index]
            text = seite.extract_text() or ""
            gesamt += len(text.encode("utf-8"))
            if gesamt > kontext.grenzen.max_vorschau_bytes:
                raise BusinessFehler("Der PDF-Text überschreitet 256 KiB. Wähle weniger Seiten.")
            seiten.append({"seite": index + 1, "text": text, "scan_verdacht": not text.strip() and bool(list(seite.images))})
        ergebnis = {"seiten": seiten, "text": "\n".join(s["text"] for s in seiten), "scan_verdacht": any(s["scan_verdacht"] for s in seiten), "ocr": False, "status": "gelesen"}
        if ergebnis["scan_verdacht"] or not ergebnis["text"].strip():
            meldung = ("Mindestens eine ausgewählte PDF-Seite hat keine verlässliche Textschicht. Der Textauszug ist unvollständig."
                       if ergebnis["text"].strip() else "Die PDF-Datei enthält keine auslesbare Textschicht.")
            meldung += " OCR ist nicht eingerichtet; die Quelle bleibt unverändert."
            if bei_scan == "STOPP":
                raise BusinessFehler(meldung)
            status = "unvollstaendig" if ergebnis["text"].strip() else "kein_text"
            hinweis = {"status": status, "quelle": str(quelle), "scan_verdacht": ergebnis["scan_verdacht"], "meldung": meldung}
            kontext.hinweise.append(hinweis)
            kontext.ereignis("hinweis", **hinweis)
            ergebnis.update(status=status, hinweis=hinweis)
        return ergebnis


def pdf_bericht(kontext: Kontext, k: dict, e: dict):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont, TTFError
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from reportlab.platypus.doctemplate import LayoutError
    dienst, tabelle = _tabelle(kontext, e)
    spalten = k.get("spalten") or [s.name for s in tabelle.spalten]
    if not isinstance(spalten, list) or not 1 <= len(spalten) <= 20 or any(s not in {s.name for s in tabelle.spalten} for s in spalten):
        raise BusinessFehler("Wähle zwischen einer und zwanzig vorhandenen Berichtsspalten.")
    schriftwahl = k.get("schrift", "INTERN:VERA")
    schrift = (Path(__file__).resolve().parent / "schriften" / "Vera.ttf"
               if schriftwahl == "INTERN:VERA" else kontext.pfad(schriftwahl))
    if schrift.suffix.lower() != ".ttf":
        raise BusinessFehler("Wähle eine lokale TrueType-Schriftdatei (.ttf).")
    fingerprint = datei_fingerabdruck(kontext, schrift)
    fontname = "ScriptLab-" + fingerprint["sha256"][:16]
    try:
        with quellkopie(kontext, schrift) as kopie:
            font = TTFont(fontname, str(kopie))
    except (TTFError, OSError, ValueError) as fehler:
        raise BusinessFehler("Die gewählte TrueType-Schrift konnte nicht gelesen werden.") from fehler
    pdfmetrics.registerFont(font)
    def schrift_pruefen(text):
        fehlend = sorted({ord(zeichen) for zeichen in text
                          if zeichen not in "\r\n\t" and not font.face.charToGlyph.get(ord(zeichen))})
        if fehlend:
            codes = ", ".join(f"U+{code:04X}" for code in fehlend[:12])
            raise BusinessFehler(f"Die gewählte Schrift enthält Zeichen des Berichts nicht ({codes}). Wähle eine geeignete TTF-Schrift.")
        return text
    stil = ParagraphStyle("ScriptLab", fontName=fontname, fontSize=9, leading=12)
    def absatz(text):
        # Erst Benutzerdaten maskieren, danach ausschließlich eigene Umbruch-Tags einfügen.
        maskiert = escape(schrift_pruefen(text)).replace("\r\n", "\n").replace("\r", "\n")
        return Paragraph(maskiert.replace("\n", "<br/>"), stil)
    daten = [[absatz(s) for s in spalten]]
    for zeile in _zeilen(kontext, dienst, tabelle):
        if (len(daten) + 1) * len(spalten) > kontext.grenzen.max_dokument_zellen:
            raise BusinessFehler("Der PDF-Bericht überschreitet die Dokumentzellengrenze. Teile den Bericht oder erhöhe die Grenze bewusst.")
        daten.append([absatz(_text(zeile[s])) for s in spalten])
    titel = schrift_pruefen(_text(k.get("titel", "Bericht")))
    schrift_pruefen("Seite 0123456789")
    def schreiben(temp):
        raender = k.get("rand_mm", 12.7)
        if type(raender) not in {int, float} or not 5 <= raender <= 40:
            raise BusinessFehler("Der Berichtsrand muss zwischen 5 und 40 mm liegen.")
        from reportlab.lib.pagesizes import A3, LETTER, landscape
        from reportlab.lib.units import mm
        formatname = k.get("seitenformat", "A4")
        if formatname not in {"A4", "A3", "LETTER"}:
            raise BusinessFehler("Wähle A4, A3 oder LETTER als PDF-Seitenformat.")
        format = {"A4": A4, "A3": A3, "LETTER": LETTER}[formatname]
        if k.get("querformat", False):
            format = landscape(format)
        tabelle_pdf = LongTable(daten, colWidths=[(format[0] - 2 * raender * mm) / len(spalten)] * len(spalten), repeatRows=1, splitInRow=1)
        tabelle_pdf.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), fontname), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EBEF")), ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
        def seitenkopf(canvas, dokument):
            kontext.pruefen()
            if dokument.page > kontext.grenzen.max_seiten:
                raise BusinessFehler("Der PDF-Bericht enthält zu viele Seiten.")
            canvas.setFont(fontname, 9)
            kopftitel = titel.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")
            canvas.drawString(raender * mm, format[1] - 24, kopftitel[:140])
            canvas.drawRightString(format[0] - raender * mm, 24, f"Seite {dokument.page}")
        try:
            SimpleDocTemplate(str(temp), pagesize=format, topMargin=max(42, raender * mm), bottomMargin=max(42, raender * mm), leftMargin=raender * mm, rightMargin=raender * mm).build([absatz(titel), Spacer(1, 12), tabelle_pdf], onFirstPage=seitenkopf, onLaterPages=seitenkopf)
        except LayoutError as fehler:
            raise BusinessFehler("Der PDF-Inhalt passt nicht in das Seitenlayout. Kürze sehr große Zellen oder wähle ein größeres Seitenformat.") from fehler
    quellen_pruefen(kontext, [fingerprint])
    return {**publizieren(kontext, _ziel(k, ".pdf"), schreiben, k.get("konflikt", "FEHLER"), [schrift, *dienst.quellpfade(tabelle)], lambda pfad: _pdf_lesen(kontext, pfad)), "dokumenttyp": "PDF"}


def docx_erstellen(kontext: Kontext, k: dict, e: dict):
    from docx import Document
    from docx.shared import Inches
    abschnitte = k.get("abschnitte", [])
    if not isinstance(abschnitte, list) or len(abschnitte) > kontext.grenzen.max_zeilen:
        raise BusinessFehler("Lege eine begrenzte Liste von Word-Absätzen, Überschriften und Bildern fest.")
    bilder = [kontext.pfad(a["pfad"]) for a in abschnitte if a.get("art") == "BILD"]
    def schreiben(temp):
        dokument = Document()
        if k.get("titel"):
            dokument.add_heading(_text(k["titel"]), level=0)
        for abschnitt in abschnitte:
            kontext.pruefen()
            art = abschnitt.get("art")
            if art == "ABSATZ":
                stil = abschnitt.get("stil") or None
                if stil is not None and stil not in dokument.styles:
                    raise BusinessFehler("Der gewählte Absatzstil fehlt im neuen Word-Dokument.")
                dokument.add_paragraph(_text(abschnitt.get("text")), style=stil)
            elif art == "UEBERSCHRIFT":
                stufe = abschnitt.get("stufe", 1)
                if type(stufe) is not int or not 0 <= stufe <= 9:
                    raise BusinessFehler("Die Word-Überschrift benötigt eine Stufe von 0 bis 9.")
                dokument.add_heading(_text(abschnitt.get("text")), level=stufe)
            elif art == "BILD":
                breite = abschnitt.get("breite_cm", 12)
                if type(breite) not in {int, float} or not 0 < breite <= 20:
                    raise BusinessFehler("Die Bildbreite muss größer als 0 und höchstens 20 cm sein.")
                with quellkopie(kontext, abschnitt["pfad"]) as kopie:
                    dokument.add_picture(str(kopie), width=Inches(breite / 2.54))
            else:
                raise BusinessFehler("Dieser Word-Abschnitt wird nicht unterstützt.")
        if "tabelle" in e:
            dienst, tabelle = _tabelle(kontext, e)
            namen = [s.name for s in tabelle.spalten]
            t = dokument.add_table(rows=1, cols=len(namen))
            for zelle, name in zip(t.rows[0].cells, namen):
                zelle.text = name
            for nummer, zeile in enumerate(_zeilen(kontext, dienst, tabelle), 1):
                if (nummer + 1) * len(namen) > kontext.grenzen.max_dokument_zellen:
                    raise BusinessFehler("Die Word-Tabelle überschreitet die Dokumentzellengrenze.")
                for zelle, name in zip(t.add_row().cells, namen):
                    zelle.text = _text(zeile[name])
        dokument.save(temp)
    if "tabelle" in e:
        dienst, tabelle = _tabelle(kontext, e)
        bilder.extend(dienst.quellpfade(tabelle))
    return {**publizieren(kontext, _ziel(k, ".docx"), schreiben, k.get("konflikt", "FEHLER"), bilder, _office_pruefer(kontext)), "dokumenttyp": "DOCX"}


HANDLER = {
    "dokument.xlsx_vorlage": xlsx_vorlage, "dokument.docx_vorlage": docx_vorlage,
    "dokument.docx_serie": docx_serie, "dokument.docx_erstellen": docx_erstellen,
    "dokument.pdf_verbinden": pdf_verbinden, "dokument.pdf_seiten": pdf_seiten,
    "dokument.pdf_teilen": pdf_teilen, "dokument.pdf_text": pdf_text,
    "dokument.pdf_bericht": pdf_bericht,
}
