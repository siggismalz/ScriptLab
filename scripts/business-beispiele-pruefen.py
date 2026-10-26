"""Reproduzierbare C/D/E-Prüfung mit eigenen, ausdrücklich synthetischen Dateien.

Der äußere Aufruf bindet diesen festen Prüfer an den vorhandenen Windows-Jobhost.
Die fachlichen Läufe arbeiten nur in neuen Kopien unter artifacts, nie im Beispiel.
"""
from __future__ import annotations

import argparse
import ctypes
from ctypes import wintypes
from datetime import datetime
import hashlib
from importlib.metadata import version
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile

WURZEL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WURZEL / "laufzeit"))
BEISPIELE = WURZEL / "beispiele" / "business"
NAMEN = {"C": "C-Bestellvorschlag", "D": "D-Dokumenteingang", "E": "E-Dateistapel"}


def prozessspeicher():
    """Nur der eigene Prüfprozess; keine Suche nach fremden Prozesskennungen."""
    if os.name != "nt":
        raise RuntimeError("Der feste Speicher-Nachweis benötigt Windows.")
    class Speicher(ctypes.Structure):
        _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD)] + [
            (name, ctypes.c_size_t) for name in (
                "PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage", "QuotaPagedPoolUsage",
                "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage", "PrivateUsage")]
    kernel = ctypes.WinDLL("kernel32.dll", winmode=0x800)
    api = ctypes.WinDLL("psapi.dll", winmode=0x800)
    kernel.GetCurrentProcess.restype = wintypes.HANDLE
    api.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Speicher), wintypes.DWORD]
    api.GetProcessMemoryInfo.restype = wintypes.BOOL
    daten = Speicher()
    daten.cb = ctypes.sizeof(daten)
    if not api.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(daten), daten.cb):
        raise ctypes.WinError()
    return {"working_set_bytes": daten.WorkingSetSize, "os_peak_working_set_bytes": daten.PeakWorkingSetSize,
            "private_bytes": daten.PrivateUsage}


def rechnerdaten():
    import winreg
    class Speicherstatus(ctypes.Structure):
        _fields_ = [("dwLength", wintypes.DWORD), ("dwMemoryLoad", wintypes.DWORD)] + [
            (name, ctypes.c_ulonglong) for name in ("ullTotalPhys", "ullAvailPhys", "ullTotalPageFile",
                "ullAvailPageFile", "ullTotalVirtual", "ullAvailVirtual", "ullAvailExtendedVirtual")]
    kernel = ctypes.WinDLL("kernel32.dll", winmode=0x800)
    kernel.GlobalMemoryStatusEx.argtypes = [ctypes.POINTER(Speicherstatus)]
    kernel.GlobalMemoryStatusEx.restype = wintypes.BOOL
    status = Speicherstatus()
    status.dwLength = ctypes.sizeof(status)
    if not kernel.GlobalMemoryStatusEx(ctypes.byref(status)):
        raise ctypes.WinError()
    with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"HARDWARE\DESCRIPTION\System\CentralProcessor\0") as key:
        cpu = winreg.QueryValueEx(key, "ProcessorNameString")[0].strip()
    windows = sys.getwindowsversion()
    return {"cpu": cpu, "logische_prozessoren": os.cpu_count(), "physischer_ram_bytes": status.ullTotalPhys,
            "windows": f"{windows.major}.{windows.minor}.{windows.build}", "architektur": "x64" if ctypes.sizeof(ctypes.c_void_p) == 8 else "x86",
            "python": sys.version.split()[0],
            "pakete": {name: version(name) for name in ("duckdb", "openpyxl", "python-docx", "reportlab", "pypdf",
                "et-xmlfile", "lxml", "typing-extensions", "pillow", "charset-normalizer", "defusedxml")}}


class Stapelmessung:
    """25-ms-Sollabstand, ausschließlich eigener Prozess-TEMP und eigenes Zielstaging."""
    def __init__(self, kontext):
        self.kontext = kontext
        self.temp = Path(tempfile.gettempdir()).resolve()
        erlaubt = (WURZEL / "artifacts" / "business-beispiele").resolve()
        if not self.temp.is_relative_to(erlaubt) or not self.temp.name.endswith("-prozess-temp"):
            raise RuntimeError("Die Lastmessung benötigt den privaten Prozess-TEMP des festen Controllers.")
        self.ende = threading.Event()
        self.fehler = None
        self.anzahl = 0
        self.temp_max = self.kontext_max = self.staging_max = self.privat_max = self.working_set_max = 0
        self.max_abstand = 0.0
        self.zuletzt = None
        self.thread = threading.Thread(target=self._messen, daemon=True)

    @staticmethod
    def _bytes(pfade):
        gesamt = 0
        for pfad in pfade:
            try:
                if pfad.is_file() and not pfad.is_symlink():
                    gesamt += pfad.stat().st_size
            except FileNotFoundError:
                pass  # Die eigene atomare Publikation kann die Datei inzwischen verschieben.
        return gesamt

    def _stichprobe(self):
        jetzt = time.perf_counter()
        if self.zuletzt is not None:
            self.max_abstand = max(self.max_abstand, jetzt - self.zuletzt)
        self.zuletzt = jetzt
        privat = self._bytes(self.temp.rglob("*"))
        staging = self._bytes(self.kontext.basis.rglob(".scriptlab-*"))
        self.temp_max = max(self.temp_max, privat + staging)
        self.kontext_max = max(self.kontext_max, self._bytes(self.kontext.arbeitsordner.rglob("*")))
        self.staging_max = max(self.staging_max, staging)
        speicher = prozessspeicher()
        self.working_set_max = max(self.working_set_max, speicher["working_set_bytes"])
        self.privat_max = max(self.privat_max, speicher["private_bytes"])
        self.anzahl += 1

    def _messen(self):
        try:
            while not self.ende.is_set():
                self._stichprobe()
                self.ende.wait(.025)
        except BaseException as fehler:
            self.fehler = fehler

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, *_):
        self.ende.set()
        self.thread.join(timeout=5)
        if self.thread.is_alive():
            raise RuntimeError("Die eigene Messung wurde nicht rechtzeitig beendet.")
        if self.fehler:
            raise self.fehler
        self._stichprobe()

    def ergebnis(self):
        return {**prozessspeicher(), "sampled_max_working_set_bytes": self.working_set_max,
            "sampled_max_private_bytes": self.privat_max, "sampled_max_temp_und_staging_bytes": self.temp_max,
            "sampled_max_context_temp_bytes": self.kontext_max, "sampled_max_staging_bytes": self.staging_max,
            "stichproben": self.anzahl, "soll_abstand_ms": 25, "max_beobachteter_abstand_ms": round(self.max_abstand * 1000, 3),
            "speicher_messumfang": "Eigener gebundener Prüfprozess; OS-Working-Set-Peak seit Prozessbeginn, Private Bytes als Stichprobe.",
            "temp_messumfang": "Privater TEMP/TMP/TMPDIR des eigenen Jobs inklusive Bibliotheken und Kontext plus .scriptlab-Staging unter der eigenen Stapelkopie; keine fremden Tempordner.",
            "grenze": "Stichproben können kurzlebige Spitzen übersehen; beobachtete Temp-/Private-Maxima sind keine garantierten absoluten Peaks."}


def sha(pfad):
    digest = hashlib.sha256()
    with Path(pfad).open("rb") as datei:
        for teil in iter(lambda: datei.read(131072), b""):
            digest.update(teil)
    return digest.hexdigest()


def dateistand(ordner):
    return {str(p.relative_to(ordner)).replace("\\", "/"): sha(p)
            for p in sorted(Path(ordner).rglob("*")) if p.is_file()}


def _zip_fest(pfad):
    """Auch Office-ZIP-Zeitstempel sind feste Fixture-Daten, keine Uhrzeitwerte."""
    with zipfile.ZipFile(pfad) as archiv:
        daten = [(eintrag.filename, archiv.read(eintrag)) for eintrag in archiv.infolist()]
    with zipfile.ZipFile(pfad, "w", compression=zipfile.ZIP_DEFLATED) as archiv:
        for name, inhalt in daten:
            if name == "docProps/core.xml":
                xml = inhalt.decode("utf-8")
                xml = re.sub(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                             r"\g<1>2026-10-26T12:00:00Z\g<2>", xml)
                inhalt = xml.encode("utf-8")
            info = zipfile.ZipInfo(name, (2026, 10, 26, 12, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archiv.writestr(info, inhalt)


def fixtures_erstellen(beispiele_wurzel=BEISPIELE):
    """Nur die hier bezeichneten neuen Testdateien werden deterministisch erzeugt."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = Path(beispiele_wurzel) / NAMEN["C"]
    (c / "eingaben").mkdir(parents=True, exist_ok=True)
    (c / "vorlagen").mkdir(parents=True, exist_ok=True)
    buch = Workbook()
    blatt = buch.active
    blatt.title = "Lager"
    for zeile in [("Artikel", "Lieferant", "Bestand", "Mindestbestand", "Zielbestand"),
                  ("Artikel A", "S1", 5, 10, 20), ("Artikel B", "S2", 30, 10, 40),
                  ("Artikel C", "S1", 0, 5, 12)]:
        blatt.append(zeile)
    for zelle in blatt[1]:
        zelle.font = Font(name="Calibri", bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="243F59")
    for spalte, breite in (("A", 22), ("B", 18), ("C", 14), ("D", 19), ("E", 17)):
        blatt.column_dimensions[spalte].width = breite
    blatt.freeze_panes = "A2"
    buch.properties.creator = "ScriptLab"
    buch.properties.created = buch.properties.modified = datetime(2026, 10, 26, 12)
    pfad = c / "eingaben" / "Lager.xlsx"
    buch.save(pfad)
    buch.close()
    _zip_fest(pfad)

    dokument = Document()
    abschnitt = dokument.sections[0]
    abschnitt.page_width, abschnitt.page_height = Inches(8.5), Inches(11)
    abschnitt.top_margin = abschnitt.bottom_margin = Inches(1)
    abschnitt.left_margin = abschnitt.right_margin = Inches(1)
    normal = dokument.styles["Normal"]
    normal.font.name, normal.font.size = "Segoe UI", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    titel = dokument.styles["Title"]
    # Die mitgelieferte python-docx-Vorlage hat einen blauen Titel-Unterrand.
    # Auch geerbte/verkettete Stile ausdrücklich bereinigen, nicht mit Weiß verdecken.
    for stilname in ("Title", "Title Char", "Normal"):
        stil = dokument.styles[stilname]
        while stil is not None:
            for rand in list(stil.element.xpath("./w:pPr/w:pBdr")):
                rand.getparent().remove(rand)
            stil.font.underline = False
            stil = stil.base_style
    titel.font.name, titel.font.size, titel.font.color.rgb = "Segoe UI", Pt(22), RGBColor(0, 0, 0)
    titel_absatz = dokument.add_paragraph("Bestellvorschlag", "Title")
    for rand in list(titel_absatz._p.xpath("./w:pPr/w:pBdr")):
        rand.getparent().remove(rand)
    for run in titel_absatz.runs:
        run.underline = False
        run.font.color.rgb = RGBColor(0, 0, 0)
    dokument.add_paragraph("Synthetisches Beispiel · keine Bestellung an einen echten Lieferanten")
    dokument.add_paragraph("Lieferant: {{Lieferant}}")
    dokument.add_paragraph("Bestelldatum: {{Datum}}")
    tabelle = dokument.add_table(rows=2, cols=4)
    tabelle.autofit = False
    for spalte, breite in zip(tabelle.columns, (2.6, 1.0, 1.4, 1.5)):
        spalte.width = Inches(breite)
    for zelle, text in zip(tabelle.rows[0].cells, ("Artikel", "Bestand", "Zielbestand", "Bestellmenge")):
        zelle.text = text
        farbe = OxmlElement("w:shd")
        farbe.set(qn("w:fill"), "243F59")
        zelle._tc.get_or_add_tcPr().append(farbe)
        for run in zelle.paragraphs[0].runs:
            run.font.bold, run.font.color.rgb = True, RGBColor(255, 255, 255)
    for zelle, text in zip(tabelle.rows[1].cells,
                          ("{{Position.Artikel}}", "{{Position.Bestand}}", "{{Position.Zielbestand}}", "{{Position.Bestellmenge}}")):
        zelle.text = text
    for zeile in tabelle.rows:
        for zelle in zeile.cells:
            rand = OxmlElement("w:tcBorders")
            for seite in ("top", "left", "bottom", "right"):
                linie = OxmlElement("w:" + seite)
                for attribut, wert in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    linie.set(qn("w:" + attribut), wert)
                rand.append(linie)
            zelle._tc.get_or_add_tcPr().append(rand)
    dokument.add_paragraph("Gesamtbestellmenge: {{Gesamtbestellmenge}}")
    dokument.add_paragraph("Bitte den Vorschlag fachlich prüfen. Dieses Beispiel sendet keine Nachricht und löst keine Bestellung aus.")
    dokument.core_properties.author = "ScriptLab"
    dokument.core_properties.created = dokument.core_properties.modified = datetime(2026, 10, 26, 12)
    pfad = c / "vorlagen" / "Bestellung.docx"
    dokument.save(pfad)
    _zip_fest(pfad)

    d = Path(beispiele_wurzel) / NAMEN["D"] / "eingaben"
    d.mkdir(parents=True, exist_ok=True)
    for name, nummer, datum in (("01-Rechnung.pdf", "Rechnung SYN-1001", "2026-10-03"),
                                 ("02-Beleg.pdf", "Beleg SYN-1002", "17.10.2026")):
        pdf = canvas.Canvas(str(d / name), pagesize=letter, invariant=1)
        pdf.setTitle("Synthetische Eingangsdatei")
        pdf.setAuthor("ScriptLab")
        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawString(72, 710, nummer)
        pdf.setFont("Helvetica", 12)
        pdf.drawString(72, 675, "Datum: " + datum)
        pdf.drawString(72, 642, "Nur synthetische Testdaten. Kein echter Geschaeftsvorgang.")
        pdf.drawString(72, 612, "Gesamtbetrag: 42,00 EUR")
        pdf.showPage()
        pdf.save()
    bild = Image.new("RGB", (900, 1165), "white")
    zeichnen = ImageDraw.Draw(bild)
    zeichnen.text((85, 130), "SYNTHETISCHER SCAN", fill="black", font=ImageFont.load_default(size=38))
    zeichnen.text((85, 235), "Keine Textebene", fill="black", font=ImageFont.load_default(size=28))
    zeichnen.text((85, 300), "Keine echten Geschaeftsdaten", fill="black", font=ImageFont.load_default(size=24))
    pdf = canvas.Canvas(str(d / "03-Scan.pdf"), pagesize=letter, invariant=1)
    pdf.setTitle("Synthetischer Scan ohne Textebene")
    pdf.setAuthor("ScriptLab")
    pdf.drawImage(ImageReader(bild), 0, 0, width=letter[0], height=letter[1])
    pdf.showPage()
    pdf.save()
    e = Path(beispiele_wurzel) / NAMEN["E"] / "eingaben"
    e.mkdir(parents=True, exist_ok=True)
    for name, inhalt in {
        "01-Nord.csv": "Bestellung;Kunde;Betrag\r\nSYN-001;Beispiel Nord;12,50\r\nSYN-002;Müller Test;7,50\r\n",
        "02-Sued.csv": "Bestellung;Kunde;Betrag\r\nSYN-003;Beispiel Süd;25,00\r\n",
        "03-Fehler.csv": "Bestellung;Kunde;Betrag\r\nSYN-FEHLER;NICHT_IM_JOURNAL\r\n",
    }.items():
        (e / name).write_bytes(b"\xef\xbb\xbf" + inhalt.encode("utf-8"))
    return {kennung: dateistand(Path(beispiele_wurzel) / name) for kennung, name in NAMEN.items()}


def beispiel_kopieren(kennung, ziel):
    basis = Path(ziel) / NAMEN[kennung]
    if basis.exists():
        raise RuntimeError("Der neue Prüfkopienordner ist bereits vorhanden.")
    shutil.copytree(BEISPIELE / NAMEN[kennung], basis)
    return basis


def laden(basis):
    return (json.loads((basis / "plan.json").read_text(encoding="utf-8")),
            json.loads((basis / "parameter.json").read_text(encoding="utf-8")))


def kontext(basis, modus="ausfuehren", **optionen):
    from scriptlab_business.kontext import Kontext, daten_hash
    plan, parameter = laden(basis)
    optionen.setdefault("workflow_hash", daten_hash(plan))
    return Kontext(basis, modus, parameter, **optionen)


def pruefe_ohne_ziele(basis):
    from scriptlab_business.ausfuehrung import ausfuehren
    plan, _ = laden(basis)
    vorher = dateistand(basis)
    for modus in ("vorschau", "prueflauf"):
        with kontext(basis, modus) as k:
            ausfuehren(plan, k)
        assert dateistand(basis) == vorher, f"{modus} hat Dateien im Beispiel verändert."
    return ["vorschau", "prueflauf"]


def pruefe_c(basis):
    from docx import Document
    from scriptlab_business.ausfuehrung import ausfuehren
    from scriptlab_business.kontext import BusinessFehler
    plan, _ = laden(basis)
    quellen = {p: sha(basis / p) for p in ("eingaben/Lager.xlsx", "vorlagen/Bestellung.docx")}
    with kontext(basis) as k:
        ergebnis = ausfuehren(plan, k)
        assert ergebnis["status"] == "erfolgreich"
        assert ergebnis["ergebnisse"]["c-bestellungen"]["anzahl"] == 1
    dateien = sorted((basis / "ausgaben").glob("*.docx"))
    assert [p.name for p in dateien] == ["Bestellung-S1.docx"]
    dokument = Document(dateien[0])
    assert not dokument.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert not dokument.styles["Normal"].element.xpath("./w:pPr/w:pBdr")
    assert not dokument.paragraphs[0]._p.xpath("./w:pPr/w:pBdr")
    assert all(run.underline is False and str(run.font.color.rgb) == "000000" for run in dokument.paragraphs[0].runs)
    zeilen = [[z.text for z in zeile.cells] for zeile in dokument.tables[0].rows]
    assert zeilen[1:] == [["Artikel A", "5", "20", "15"], ["Artikel C", "0", "12", "12"]], zeilen
    text = "\n".join([p.text for p in dokument.paragraphs] + [t for z in zeilen for t in z])
    assert "Gesamtbestellmenge: 27" in text and "Lieferant: S1" in text
    assert "Artikel B" not in text and "{{" not in text and "}}" not in text
    unveraendert = sha(dateien[0])
    try:
        with kontext(basis) as k:
            ausfuehren(plan, k)
    except BusinessFehler as fehler:
        assert "bereits vorhanden" in str(fehler)
    else:
        raise AssertionError("Das bestehende Word-Ziel wurde nicht geschützt.")
    assert sha(dateien[0]) == unveraendert
    assert all(sha(basis / p) == h for p, h in quellen.items())
    return {"dokumente": 1, "lieferant": "S1", "positionen": [15, 12], "gesamtbestellmenge": 27,
            "vorhandenes_ziel_geschuetzt": True, "quellen_unveraendert": True,
            "docx": str(dateien[0].relative_to(basis)), "docx_sha256": unveraendert}


def pruefe_d(basis):
    from pypdf import PdfReader
    from scriptlab_business.ausfuehrung import ausfuehren
    from scriptlab_business.kontext import BusinessFehler
    plan, _ = laden(basis)
    original = {p.name: sha(p) for p in (basis / "eingaben").glob("*.pdf")}
    scan = PdfReader(basis / "eingaben" / "03-Scan.pdf", strict=True)
    assert not scan.pages[0].extract_text().strip()
    objekte = scan.pages[0]["/Resources"]["/XObject"].get_object()
    assert any(o.get_object().get("/Subtype") == "/Image" for o in objekte.values())
    with kontext(basis) as k:
        ergebnis = ausfuehren(plan, k)
        hinweise = list(k.hinweise)
    ziele = sorted((basis / "ausgaben").rglob("*.pdf"))
    assert [p.name for p in ziele] == ["SYN-1001.pdf", "SYN-1002.pdf"]
    assert all(p.parent.name == "2026-10" for p in ziele)
    assert sha(ziele[0]) == original["01-Rechnung.pdf"]
    assert sha(ziele[1]) == original["02-Beleg.pdf"]
    assert ergebnis["ergebnisse"]["d-scanhinweis"]["scan_verdacht"] is True
    assert len(hinweise) == 1 and "Textschicht" in hinweise[0]["meldung"] and "OCR" in hinweise[0]["meldung"]
    try:
        with kontext(basis) as k:
            ausfuehren(plan, k)
    except BusinessFehler as fehler:
        assert "bereits vorhanden" in str(fehler)
    else:
        raise AssertionError("Das bestehende PDF-Ziel wurde nicht geschützt.")
    assert {p.name: sha(p) for p in (basis / "eingaben").glob("*.pdf")} == original
    assert sha(ziele[0]) == original["01-Rechnung.pdf"]
    return {"abgelegte_kopien": 2, "scan_hinweise": len(hinweise), "scan_ohne_textebene": True,
            "originale_unveraendert": True, "vorhandenes_ziel_geschuetzt": True,
            "zielnamen": [str(p.relative_to(basis)).replace("\\", "/") for p in ziele]}


def _e_zustaende(basis, original):
    for name in ("01-Nord.csv", "02-Sued.csv"):
        assert not (basis / "eingaben" / name).exists()
        assert sha(basis / "ausgaben" / name) == original[name]
        assert sha(basis / "archiv" / name) == original[name]
    assert sha(basis / "eingaben" / "03-Fehler.csv") == original["03-Fehler.csv"]
    assert not (basis / "ausgaben" / "03-Fehler.csv").exists()
    assert not (basis / "archiv" / "03-Fehler.csv").exists()


def pruefe_e(basis):
    from scriptlab_business.ausfuehrung import ausfuehren
    plan, _ = laden(basis)
    original = {p.name: sha(p) for p in (basis / "eingaben").glob("*.csv")}
    with kontext(basis) as k:
        ergebnis = ausfuehren(plan, k)
        assert ergebnis["status"] == "mit_fehlern" and len(k.fehler) == 1
        journal = k.journal.pfad
    _e_zustaende(basis, original)
    assert "NICHT_IM_JOURNAL" not in journal.read_text(encoding="utf-8")
    return {"ausgaben": 2, "archiviert": 2, "fehlerhafte_quelle_erhalten": True,
            "offene_fehler": 1, "keine_rohwerte_im_journal": True}


def pruefe_e_abbruch(basis):
    from scriptlab_business.ausfuehrung import ausfuehren
    from scriptlab_business.kontext import AbbruchFehler
    plan, _ = laden(basis)
    original = {p.name: sha(p) for p in (basis / "eingaben").glob("*.csv")}
    journal = basis / ".scriptlab" / "gezielter-abbruch.jsonl"
    with kontext(basis, journal_pfad=journal) as k:
        schreiben = k.journal.schreiben
        def nach_archiv_stoppen(zustand, *args, **kwargs):
            eintrag = schreiben(zustand, *args, **kwargs)
            if zustand == "archiviert":
                k.abgebrochen = True
            return eintrag
        k.journal.schreiben = nach_archiv_stoppen
        try:
            ausfuehren(plan, k)
        except AbbruchFehler:
            pass
        else:
            raise AssertionError("Der gezielte Abbruch wurde nicht beobachtet.")
    assert len(list((basis / "archiv").glob("*.csv"))) == 1
    erstes = basis / "ausgaben" / "01-Nord.csv"
    erster_stand = (sha(erstes), erstes.stat().st_mtime_ns)
    with kontext(basis, journal_pfad=journal, fortsetzen=True) as k:
        ergebnis = ausfuehren(plan, k)
        dateien = ergebnis["ergebnisse"]["e-csv-stapel"]["dateien"]
        assert sum(d["status"] == "wiederverwendet" for d in dateien) == 1
        assert len(k.fehler) == 1
    assert (sha(erstes), erstes.stat().st_mtime_ns) == erster_stand
    _e_zustaende(basis, original)
    assert len(list((basis / "ausgaben").glob("*.csv"))) == 2
    assert len(list((basis / "archiv").glob("*.csv"))) == 2
    return {"abbruch_nach_archivierung": 1, "wiederverwendet": 1,
            "keine_doppelten_ausgaben": True, "erste_ausgabe_nicht_neu_geschrieben": True}


def pruefe_e_unklar(basis):
    from scriptlab_business.ausfuehrung import ausfuehren
    from scriptlab_business.kontext import UnklarerZustand
    plan, _ = laden(basis)
    journal = basis / ".scriptlab" / "unklarer-zustand.jsonl"
    with kontext(basis, journal_pfad=journal) as k:
        schreiben = k.journal.schreiben
        def bestaetigung_unterbrechen(zustand, *args, **kwargs):
            if zustand == "publiziert":
                raise RuntimeError("Fester QA-Abbruch nach Rename vor Journalbestätigung")
            return schreiben(zustand, *args, **kwargs)
        k.journal.schreiben = bestaetigung_unterbrechen
        try:
            ausfuehren(plan, k)
        except RuntimeError as fehler:
            assert "Fester QA-Abbruch" in str(fehler)
        else:
            raise AssertionError("Die vorbereitete Unterbrechung wurde nicht getroffen.")
    assert (basis / "ausgaben" / "01-Nord.csv").exists()
    assert (basis / "eingaben" / "01-Nord.csv").exists()
    vorher = dateistand(basis)
    try:
        with kontext(basis, journal_pfad=journal, fortsetzen=True) as k:
            ausfuehren(plan, k)
    except UnklarerZustand as fehler:
        assert "nicht bestätigt" in str(fehler)
    else:
        raise AssertionError("Der unklare Zustand wurde automatisch wiederholt.")
    assert dateistand(basis) == vorher
    return {"unterbrechung_nach_atomic_rename": True, "quelle_erhalten": True,
            "automatische_wiederholung_verhindert": True}


def benchmark_stapel(basis, anzahl):
    from scriptlab_business.ausfuehrung import ausfuehren
    from scriptlab_business.kontext import Grenzen
    if anzahl != 1000:
        raise ValueError("Der feste optionale Dateistapel-Benchmark umfasst genau 1000 Dateien.")
    # Ausschließlich die neu angelegte Prüfkopie wird mit synthetischen Dateien erweitert.
    for pfad in (basis / "eingaben").glob("*.csv"):
        pfad.unlink()
    for nummer in range(anzahl):
        (basis / "eingaben" / f"SYN-{nummer:04}.csv").write_bytes(
            b"\xef\xbb\xbfBestellung;Kunde;Betrag\r\n" + f"SYN-{nummer:04};Testkunde;1,00\r\n".encode("ascii"))
    plan, _ = laden(basis)
    original = {p.name: sha(p) for p in (basis / "eingaben").glob("*.csv")}
    hardware = rechnerdaten()
    start = time.perf_counter()
    with kontext(basis, grenzen=Grenzen(max_ausgaben=3000, zeitlimit_sekunden=120)) as k, Stapelmessung(k) as messung:
        ergebnis = ausfuehren(plan, k)
        assert ergebnis["status"] == "erfolgreich"
    dauer = time.perf_counter() - start
    assert not list((basis / "eingaben").glob("*.csv"))
    for ordner in ("ausgaben", "archiv"):
        assert {p.name: sha(p) for p in (basis / ordner).glob("*.csv")} == original
    return {"dateien": anzahl, "publizierte_dateien": anzahl * 2, "dauer_sekunden": round(dauer, 3),
            "eingabebytes": sum((basis / "archiv" / n).stat().st_size for n in original),
            "alle_ausgaben_und_archive_hashgleich": True, "zeitlimit_sekunden": 120,
            "grenzen": {"max_ausgaben": k.grenzen.max_ausgaben, "max_temp_bytes": k.grenzen.max_temp_bytes,
                        "duckdb_speicher": k.grenzen.duckdb_speicher},
            "rechner": hardware, "messung": messung.ergebnis()}


def pruefung(ziel, benchmark=False, nur_benchmark=False):
    ziel.mkdir(parents=True, exist_ok=False)
    vorher = {k: dateistand(BEISPIELE / n) for k, n in NAMEN.items()}
    bericht = {"version": 1, "synthetische_daten": True, "python": sys.version.split()[0], "pruefungen": {}}
    quellpfade = ["scripts/business-beispiele-pruefen.py", "laufzeit/requirements-business.lock"] + [
        f"laufzeit/scriptlab_business/{name}.py" for name in ("ausfuehrung", "kontext", "dateien", "daten", "dokumente", "typen")]
    bericht["quell_sha256"] = {name: sha(WURZEL / name) for name in quellpfade}
    if not nur_benchmark:
        for kennung in NAMEN:
            basis = beispiel_kopieren(kennung, ziel)
            bericht["pruefungen"][kennung + "_ohne_zieldateien"] = pruefe_ohne_ziele(basis)
            bericht["pruefungen"][kennung] = {"C": pruefe_c, "D": pruefe_d, "E": pruefe_e}[kennung](basis)
        for name, funktion in (("E-Abbruch", pruefe_e_abbruch), ("E-Unklar", pruefe_e_unklar)):
            ordner = ziel / name
            ordner.mkdir()
            bericht["pruefungen"][name] = funktion(beispiel_kopieren("E", ordner))
    if benchmark:
        ordner = ziel / "E-1000-Dateien"
        ordner.mkdir()
        bericht["benchmark"] = benchmark_stapel(beispiel_kopieren("E", ordner), 1000)
        bericht["benchmark"]["isolierter_benchmark_ohne_cde_vorlauf"] = nur_benchmark
    assert vorher == {k: dateistand(BEISPIELE / n) for k, n in NAMEN.items()}
    bericht["beispielquellen_unveraendert"] = True
    bericht["visuelle_pruefung"] = "Separater Render- und Bilddurchlauf erforderlich; nicht aus Inhaltsassertions ableiten."
    (ziel / "pruefung.json").write_text(json.dumps(bericht, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return bericht


def _job(args, ziel):
    helper = WURZEL / "artifacts" / "native" / "ScriptLab.Testhost.exe"
    if os.name != "nt" or not helper.is_file():
        raise RuntimeError("Diese Prüfung benötigt den zuvor gebauten festen Windows-Jobhost.")
    ziel.parent.mkdir(parents=True, exist_ok=True)
    harness = ziel.parent / (ziel.name + "-aufruf.py")
    optionen = [str(Path(__file__).resolve()), "--intern", str(ziel)]
    if args.fixtures:
        optionen.append("--fixtures")
    if args.benchmark_dateien:
        optionen.extend(["--benchmark-dateien", "1000"])
    if args.tests:
        optionen.append("--tests")
    if args.nur_benchmark:
        optionen.append("--nur-benchmark")
    harness.write_text("import runpy,sys\nsys.argv=" + repr(optionen) + "\nrunpy.run_path(sys.argv[0],run_name='__main__')\n", encoding="utf-8")
    limit = 170 if args.benchmark_dateien else 45
    prozess_temp = ziel.parent / (ziel.name + "-prozess-temp")
    prozess_temp.mkdir(exist_ok=False)
    umgebung = {**os.environ, "TEMP": str(prozess_temp), "TMP": str(prozess_temp), "TMPDIR": str(prozess_temp)}
    prozess = subprocess.Popen([str(helper), "--script", str(harness), "--basis", str(ziel.parent),
        "--limit-ms", str(limit * 1000), "--token", uuid.uuid4().hex,
        "--sprache", "python", "--interpreter", sys.executable],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        creationflags=subprocess.CREATE_NO_WINDOW, cwd=ziel.parent, env=umgebung)
    ausgaben = [bytearray(), bytearray()]
    def lesen(strom, index):
        for teil in iter(lambda: strom.read(4096), b""):
            if len(ausgaben[index]) < 131072:
                ausgaben[index].extend(teil[:131072 - len(ausgaben[index])])
    threads = [threading.Thread(target=lesen, args=(strom, i), daemon=True)
               for i, strom in enumerate((prozess.stdout, prozess.stderr))]
    for thread in threads:
        thread.start()
    try:
        try:
            prozess.wait(timeout=limit + 5)
        except subprocess.TimeoutExpired:
            try:
                prozess.stdin.write(b"STOP\n")
                prozess.stdin.flush()
            except (OSError, BrokenPipeError):
                pass
            try:
                prozess.wait(timeout=5)
            except subprocess.TimeoutExpired:
                prozess.kill()  # Kill-on-close beendet ausschließlich diesen eigenen Job.
                prozess.wait(timeout=5)
            raise
    finally:
        prozess.stdin.close()
        for thread in threads:
            thread.join(timeout=2)
        prozess.stdout.close()
        prozess.stderr.close()
    (ziel.parent / (ziel.name + "-stdout.txt")).write_bytes(ausgaben[0])
    (ziel.parent / (ziel.name + "-stderr.txt")).write_bytes(ausgaben[1])
    sys.stdout.buffer.write(ausgaben[0])
    if prozess.returncode:
        sys.stderr.buffer.write(ausgaben[1])
    return prozess.returncode


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fixtures", action="store_true", help="Die bezeichneten synthetischen Eingaben und Vorlagen neu erzeugen.")
    parser.add_argument("--benchmark-dateien", type=int, choices=(1000,))
    parser.add_argument("--tests", action="store_true", help="Die eigene C/D/E-Fachtestsuite im begrenzten Job ausführen.")
    parser.add_argument("--nur-benchmark", action="store_true", help="Den festen 1000-Dateien-Lauf ohne CDE-Vorläufe in einem frischen Prüfprozess messen.")
    parser.add_argument("--intern", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.nur_benchmark and (args.benchmark_dateien != 1000 or args.tests):
        parser.error("--nur-benchmark benötigt --benchmark-dateien 1000 und kann nicht mit --tests kombiniert werden.")
    if args.intern:
        ziel = args.intern.resolve()
        erlaubt = (WURZEL / "artifacts" / "business-beispiele").resolve()
        if not ziel.is_relative_to(erlaubt) or ziel == erlaubt:
            raise RuntimeError("Prüfausgaben sind nur in einem neuen eigenen artifacts-Unterordner erlaubt.")
        if args.fixtures:
            fixtures_erstellen()
        if args.tests:
            import unittest
            suite = unittest.defaultTestLoader.discover(str(WURZEL / "tests" / "business"), "test_beispiele_dokumente.py")
            result = unittest.TextTestRunner(verbosity=2).run(suite)
            ziel.mkdir(parents=True, exist_ok=False)
            bericht = {"tests": result.testsRun, "fehler": len(result.errors),
                       "fehlgeschlagen": len(result.failures), "ausgelassen": len(result.skipped)}
            (ziel / "fachtests.json").write_text(json.dumps(bericht, indent=2) + "\n", encoding="utf-8")
            print(json.dumps(bericht))
            return 0 if result.wasSuccessful() and not result.skipped else 1
        bericht = pruefung(ziel, args.benchmark_dateien == 1000, args.nur_benchmark)
        print(json.dumps(bericht, ensure_ascii=True, indent=2))
        return 0
    ziel = WURZEL / "artifacts" / "business-beispiele" / (datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:8])
    return _job(args, ziel)


if __name__ == "__main__":
    raise SystemExit(main())
