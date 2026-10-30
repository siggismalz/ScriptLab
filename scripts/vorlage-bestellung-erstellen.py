"""Die vorhandene Bestellvorlage ohne neue Datenfelder reproduzierbar gestalten.

Aufruf: python -B scripts/vorlage-bestellung-erstellen.py [Ziel.docx]
Benötigt python-docx; der Generator gehört nicht zur Workflowausführung.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

WURZEL = Path(__file__).resolve().parents[1]
STANDARDZIEL = WURZEL / "beispiele/business/C-Bestellvorschlag/vorlagen/Bestellung.docx"


def schrift(stil, groesse, fett=False):
    stil.font.name = "Segoe UI"
    stil.font.size = Pt(groesse)
    stil.font.bold = fett
    stil.font.underline = False
    stil.font.color.rgb = RGBColor(0, 0, 0)
    eigenschaften = stil.element.get_or_add_rPr()
    sprache = OxmlElement("w:lang")
    sprache.set(qn("w:val"), "de-DE")
    eigenschaften.append(sprache)
    for rand in list(stil.element.xpath("./w:pPr/w:pBdr")):
        rand.getparent().remove(rand)


def zelle_gestalten(zelle, breite, kopf=False, zahl=False):
    zelle.width = Cm(breite)
    zelle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    eigenschaften = zelle._tc.get_or_add_tcPr()
    polster = OxmlElement("w:tcMar")
    for seite, wert in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):
        rand = OxmlElement("w:" + seite)
        rand.set(qn("w:w"), str(wert))
        rand.set(qn("w:type"), "dxa")
        polster.append(rand)
    eigenschaften.append(polster)
    fuellung = OxmlElement("w:shd")
    fuellung.set(qn("w:fill"), "EDF0F0" if kopf else "FFFFFF")
    eigenschaften.append(fuellung)
    raender = OxmlElement("w:tcBorders")
    for seite in ("top", "left", "bottom", "right"):
        linie = OxmlElement("w:" + seite)
        for attribut, wert in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            linie.set(qn("w:" + attribut), wert)
        raender.append(linie)
    eigenschaften.append(raender)
    for absatz in zelle.paragraphs:
        absatz.paragraph_format.space_before = Pt(0)
        absatz.paragraph_format.space_after = Pt(0)
        absatz.paragraph_format.line_spacing = 1.15
        absatz.alignment = WD_ALIGN_PARAGRAPH.RIGHT if zahl else WD_ALIGN_PARAGRAPH.LEFT
        for lauf in absatz.runs:
            lauf.font.name = "Segoe UI"
            lauf.font.size = Pt(11)
            lauf.font.bold = kopf
            lauf.font.color.rgb = RGBColor(0, 0, 0)


def erstellen(ziel: Path = STANDARDZIEL):
    if ziel.suffix.lower() != ".docx":
        raise ValueError("Die Bestellvorlage benötigt die Dateiendung .docx.")
    dokument = Document()
    abschnitt = dokument.sections[0]
    abschnitt.page_width, abschnitt.page_height = Cm(21), Cm(29.7)
    abschnitt.left_margin = abschnitt.right_margin = Cm(2)
    abschnitt.top_margin, abschnitt.bottom_margin = Cm(2), Cm(1.8)
    abschnitt.header_distance = abschnitt.footer_distance = Cm(0.8)

    normal = dokument.styles["Normal"]
    schrift(normal, 11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Title", "Title Char"):
        schrift(dokument.styles[name], 26)
    titel = dokument.styles["Title"]
    titel.paragraph_format.space_after = Pt(5)
    titel.paragraph_format.keep_with_next = True
    schrift(dokument.styles["Heading 2"], 12, True)
    dokument.styles["Heading 2"].paragraph_format.space_before = Pt(18)
    dokument.styles["Heading 2"].paragraph_format.space_after = Pt(6)

    angaben = dokument.styles.add_style("Bestellangaben", WD_STYLE_TYPE.PARAGRAPH)
    angaben.base_style = normal
    schrift(angaben, 11)
    angaben.paragraph_format.space_after = Pt(5)
    angaben.paragraph_format.keep_with_next = True

    titelabsatz = dokument.add_paragraph("Bestellvorschlag", "Title")
    for lauf in titelabsatz.runs:
        lauf.font.color.rgb = RGBColor(0, 0, 0)
        lauf.font.underline = False
    dokument.add_paragraph("Zur fachlichen Prüfung vor einer Bestellung")
    for beschriftung, platzhalter in (("Lieferant", "{{Lieferant}}"), ("Bestelldatum", "{{Datum}}")):
        absatz = dokument.add_paragraph(style="Bestellangaben")
        absatz.add_run(beschriftung + ": ").bold = True
        absatz.add_run(platzhalter)

    dokument.add_paragraph("Bestellpositionen", "Heading 2")
    einheit = dokument.add_paragraph("Alle Mengen in Stück. Die Bestellmenge ergänzt den Bestand auf den Zielbestand.")
    einheit.paragraph_format.keep_with_next = True
    einheit.paragraph_format.space_after = Pt(8)

    # Erste Tabelle und zweite Zeile sind der bestehende Positionsvertrag.
    tabelle = dokument.add_table(rows=2, cols=4)
    tabelle.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabelle.autofit = False
    breiten = (3.9, 3.8, 4.5, 4.8)
    for spalte, breite in zip(tabelle.columns, breiten):
        spalte.width = Cm(breite)
    for zelle, text in zip(tabelle.rows[0].cells, ("Artikel", "Bestand", "Zielbestand", "Bestellmenge")):
        zelle.text = text
    for zelle, text in zip(tabelle.rows[1].cells, (
        "{{Position.Artikel}}", "{{Position.Bestand}}", "{{Position.Zielbestand}}", "{{Position.Bestellmenge}}",
    )):
        zelle.text = text
    for nummer, zeile in enumerate(tabelle.rows):
        zeileneigenschaften = zeile._tr.get_or_add_trPr()
        zeileneigenschaften.append(OxmlElement("w:cantSplit"))
        if nummer == 0:
            zeileneigenschaften.append(OxmlElement("w:tblHeader"))
        for spaltennummer, (zelle, breite) in enumerate(zip(zeile.cells, breiten)):
            zelle_gestalten(zelle, breite, kopf=nummer == 0, zahl=spaltennummer > 0)
    raender = OxmlElement("w:tblBorders")
    for seite in ("top", "left", "bottom", "right", "insideH", "insideV"):
        linie = OxmlElement("w:" + seite)
        for attribut, wert in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            linie.set(qn("w:" + attribut), wert)
        raender.append(linie)
    tabelle._tbl.tblPr.append(raender)

    summe = dokument.add_paragraph()
    summe.paragraph_format.space_before = Pt(12)
    summe.paragraph_format.space_after = Pt(12)
    summe.paragraph_format.keep_with_next = True
    lauf = summe.add_run("Gesamtbestellmenge: {{Gesamtbestellmenge}} Stück")
    lauf.bold, lauf.font.size = True, Pt(12)
    dokument.add_paragraph("Synthetisches Beispiel. Bitte den Vorschlag fachlich prüfen. Dieser Ablauf löst keine Bestellung aus.")

    fuss = abschnitt.footer.paragraphs[0]
    fuss.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fuss.add_run("Seite ")
    feld = OxmlElement("w:fldSimple")
    feld.set(qn("w:instr"), "PAGE")
    fuss._p.append(feld)
    for lauf in fuss.runs:
        lauf.font.name, lauf.font.size = "Segoe UI", Pt(9)
    eigenschaften = dokument.core_properties
    eigenschaften.title = "Bestellvorschlag"
    eigenschaften.subject = "Synthetisches Beispiel für einen Bestellvorschlag"
    eigenschaften.author = eigenschaften.last_modified_by = ""
    eigenschaften.created = eigenschaften.modified = datetime(2026, 10, 26, 12)
    strom = BytesIO()
    dokument.save(strom)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    # Feste ZIP-Zeitpunkte vermeiden zufällige Binärunterschiede beim Neugenerieren.
    with ZipFile(strom) as quelle, ZipFile(ziel, "w", compression=ZIP_DEFLATED) as ausgabe:
        for name in sorted(quelle.namelist()):
            eintrag = ZipInfo(name, (2026, 10, 26, 12, 0, 0))
            eintrag.compress_type = ZIP_DEFLATED
            ausgabe.writestr(eintrag, quelle.read(name))
    return ziel


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("ziel", nargs="?", type=Path, default=STANDARDZIEL)
    print(erstellen(parser.parse_args().ziel))
