"""Begrenzte, reproduzierbare Datenlasten mit ausschließlich synthetischen Dateien.

Mit dem geprüften Business-Interpreter starten. Der Controller startet nur diese
Datei als festen Childprozess, setzt eine harte Frist und misst den Prozess.
Alle Ergebnisse bleiben im ignorierten artifacts/business-qa-Verzeichnis.
"""

from __future__ import annotations

import argparse
import ctypes
from ctypes import wintypes
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid

WURZEL = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(WURZEL/"laufzeit"))


def _speicher():
    if os.name != "nt":
        return {"peak_working_set_bytes":None,"private_bytes":None}
    class Speicher(ctypes.Structure):
        _fields_=[("cb",wintypes.DWORD),("PageFaultCount",wintypes.DWORD)]+[(n,ctypes.c_size_t) for n in ("PeakWorkingSetSize","WorkingSetSize","QuotaPeakPagedPoolUsage","QuotaPagedPoolUsage","QuotaPeakNonPagedPoolUsage","QuotaNonPagedPoolUsage","PagefileUsage","PeakPagefileUsage","PrivateUsage")]
    kernel=ctypes.WinDLL("kernel32.dll",winmode=0x800)
    api=ctypes.WinDLL("psapi.dll",winmode=0x800)
    kernel.GetCurrentProcess.restype=wintypes.HANDLE
    api.GetProcessMemoryInfo.argtypes=[wintypes.HANDLE,ctypes.POINTER(Speicher),wintypes.DWORD]
    api.GetProcessMemoryInfo.restype=wintypes.BOOL
    daten=Speicher();daten.cb=ctypes.sizeof(daten)
    if not api.GetProcessMemoryInfo(kernel.GetCurrentProcess(),ctypes.byref(daten),daten.cb):
        raise ctypes.WinError()
    return {"peak_working_set_bytes":daten.PeakWorkingSetSize,"private_bytes":daten.PrivateUsage}


def _bytes(ordner):
    gesamt=0
    for pfad in ordner.rglob("*"):
        try:
            if pfad.is_file(): gesamt+=pfad.stat().st_size
        except FileNotFoundError:
            pass
    return gesamt


class Messung:
    def __init__(self,kontext):
        self.kontext=kontext
        self.temp_max=0
        self.kontext_temp_max=0
        self.bibliothek_temp_max=0
        self.staging_max=0
        self.spill_max=0
        self.privat_max=0
        self.ende=threading.Event()
        self.phase="vorbereiten"
        self.thread=threading.Thread(target=self._messen,daemon=True)

    def _messen(self):
        while not self.ende.is_set():
            kontext_bytes=_bytes(self.kontext.arbeitsordner)
            bibliothek_bytes=_bytes(self.kontext.basis/"prozess-temp")
            staging_bytes=0
            for pfad in self.kontext.basis.rglob(".scriptlab-*"):
                try:
                    if pfad.is_file() and not pfad.is_relative_to(self.kontext.basis/"prozess-temp"): staging_bytes+=pfad.stat().st_size
                except FileNotFoundError: pass
            self.kontext_temp_max=max(self.kontext_temp_max,kontext_bytes)
            self.bibliothek_temp_max=max(self.bibliothek_temp_max,bibliothek_bytes)
            self.staging_max=max(self.staging_max,staging_bytes)
            self.temp_max=max(self.temp_max,bibliothek_bytes+staging_bytes)
            self.spill_max=max(self.spill_max,_bytes(self.kontext.arbeitsordner/"duckdb-temp"))
            self.privat_max=max(self.privat_max,_speicher()["private_bytes"] or 0)
            self.ende.wait(.025)

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self,*fehler):
        self.ende.set();self.thread.join()
        if fehler[1] is not None:
            fehler[1].business_messung={**self.json(),"phase":self.phase}

    def json(self):
        return {**_speicher(),"sampled_private_max_bytes":self.privat_max,"sampled_temp_max_bytes":self.temp_max,"sampled_context_temp_max_bytes":self.kontext_temp_max,"sampled_library_temp_max_bytes":self.bibliothek_temp_max,"sampled_staging_max_bytes":self.staging_max,"sampled_duckdb_spill_max_bytes":self.spill_max,"sample_interval_ms":25,"temp_messumfang":"Privater Prozess-TEMP inklusive Kontext/DuckDB/openpyxl plus Zielstaging; Enginephase, Stichproben alle 25 ms"}


def _sha(pfad):
    digest=hashlib.sha256()
    with pfad.open("rb") as datei:
        for teil in iter(lambda:datei.read(131072),b""): digest.update(teil)
    return digest.hexdigest()


def _job(args,ziel,art):
    """Vorhandener fester Job-Helper bindet auch den venv-Redirector-Nachkommen."""
    helper=WURZEL/"artifacts"/"native"/"ScriptLab.Testhost.exe"
    if os.name!="nt" or not helper.is_file():
        raise RuntimeError("Diese Windows-Prüfung benötigt den zuvor gebauten festen ScriptLab-Testhost.")
    auftrag={"art":art,"zeilen":args.zeilen,"xlsx_zeilen":args.xlsx_zeilen,"zeitlimit":args.zeitlimit,"profil":args.profil}
    (ziel/"auftrag.json").write_text(json.dumps(auftrag),encoding="utf-8")
    harness=ziel/"fester-pruefaufruf.py"
    harness.write_text("import json,pathlib,runpy,sys\nbasis=pathlib.Path(sys.argv[sys.argv.index('--scriptlab-basis')+1])\na=json.loads((basis/'auftrag.json').read_text(encoding='utf-8'))\nsys.argv=["+repr(str(Path(__file__).resolve()))+",'--art',a['art'],'--zeilen',str(a['zeilen']),'--xlsx-zeilen',str(a['xlsx_zeilen']),'--zeitlimit',str(a['zeitlimit']),'--profil',a['profil'],'--intern',str(basis)]\nrunpy.run_path(sys.argv[0],run_name='__main__')\n",encoding="utf-8")
    token=uuid.uuid4().hex
    befehl=[str(helper),"--script",str(harness),"--basis",str(ziel),"--limit-ms",str((args.zeitlimit+20)*1000),"--token",token,"--sprache","python","--interpreter",sys.executable]
    prozess=subprocess.Popen(befehl,stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,creationflags=subprocess.CREATE_NO_WINDOW,cwd=ziel)
    ausgaben=[bytearray(),bytearray()]
    def lesen(strom,index):
        for teil in iter(lambda:strom.read(4096),b""):
            if len(ausgaben[index])<65536: ausgaben[index].extend(teil[:65536-len(ausgaben[index])])
    threads=[threading.Thread(target=lesen,args=(strom,index),daemon=True) for index,strom in enumerate((prozess.stdout,prozess.stderr))]
    for thread in threads: thread.start()
    try:
        try:
            prozess.wait(timeout=args.zeitlimit+30)
        except subprocess.TimeoutExpired:
            try: prozess.stdin.write(b"STOP\n");prozess.stdin.flush()
            except (OSError,BrokenPipeError): pass
            try: prozess.wait(timeout=5)
            except subprocess.TimeoutExpired:
                # Ein beendeter fester Host schließt seinen Kill-on-close-Job.
                prozess.kill();prozess.wait(timeout=5)
            raise
    finally:
        prozess.stdin.close()
        for thread in threads: thread.join(timeout=2)
        prozess.stdout.close();prozess.stderr.close()
    return prozess.returncode,bytes(ausgaben[0]),bytes(ausgaben[1])


def _csv(basis,anzahl,zeitlimit,profil):
    from scriptlab_business.kontext import Kontext,Grenzen
    from scriptlab_business.daten import DatenEngine
    quelle=basis/"million.csv"
    with quelle.open("w",encoding="utf-8",newline="") as datei:
        datei.write("Bestellung;Kunde;Betrag;Aktiv;Beschreibung\r\n")
        for i in range(anzahl):
            datei.write(f"B{i:09};K{i%10000:05};1.25;{'ja' if i%2==0 else 'nein'};Synthetische Testzeile {i:09} – "+"Textdaten "*15+"\r\n")
    kunden=basis/"kunden.csv"
    with kunden.open("w",encoding="utf-8",newline="") as datei:
        datei.write("Kunde;Region\r\n")
        for i in range(10000): datei.write(f"K{i:05};Region{i%8}\r\n")
    quellhash=_sha(quelle)
    grenzen=Grenzen(max_datei_bytes=2*1024**3,max_zeilen=max(anzahl,100000),duckdb_speicher="64MB",max_temp_bytes=2*1024**3,zeitlimit_sekunden=zeitlimit)
    start=time.perf_counter()
    with Kontext(basis,"ausfuehren",grenzen=grenzen,workflow_hash="synthetische-million-csv") as k, Messung(k) as messung:
        d=DatenEngine(k);k.dienste["daten"]=d
        messung.phase="csv_import"
        t=d.importieren("CSV",quelle.name)
        t=d.spalten(t,{"aktion":"TYP","spalte":"Betrag","typ":"DEZIMAL"})
        t=d.filtern(t,{"art":"vergleich","spalte":"Aktiv","operator":"EQ","wert":{"art":"wert","typ":"TEXT","wert":"ja"}})
        kunden_t=d.importieren("CSV",kunden.name)
        messung.phase="join"
        t=d.verknuepfen(t,kunden_t,{"art":"LEFT","schluessel":[{"links":"Kunde","rechts":"Kunde"}],"kardinalitaet":"n:1","fehlende_zuordnung":"STOPP"})
        if profil=="schmal":
            t=d.spalten(t,{"aktion":"AUSWAEHLEN","spalten":["Bestellung","Kunde","Region","Betrag"]})
        t=d.sortieren(t,[{"spalte":"Bestellung","richtung":"DESC","nullen":"LETZTE"}])
        k.beginn_schritt("detail-export",{}, {})
        messung.phase="sortierter_export"
        output=d.exportieren(t,"PARQUET","sortierte-details.parquet")
        gruppen=d.gruppieren(t,["Region"],[{"funktion":"ANZAHL","name":"Datensaetze"},{"funktion":"SUMME","spalte":"Betrag","name":"Betrag"}])
        messung.phase="aggregation"
        ergebnis=list(d.zeilen(gruppen)) # Höchstens acht Ergebnisgruppen, keine Rohdatentabelle.
        erwartet=(anzahl+1)//2
        assert sum(z[1] for z in ergebnis)==erwartet
        assert sum(z[2] for z in ergebnis)==Decimal("1.25")*erwartet
        gruppensoll={f"Region{r}":max(0,(anzahl-1-r)//8+1) for r in (0,2,4,6)}
        assert {z[0]:z[1] for z in ergebnis}=={r:n for r,n in gruppensoll.items() if n}
        gelesen=d.importieren("PARQUET","sortierte-details.parquet")
        pruefung=d.pruefen(gelesen,[{"art":"ZEILENZAHL","minimum":erwartet,"maximum":erwartet}],"STOPP")
        vorherige=anzahl+1
        for zeile in d.zeilen(gelesen):
            kennung=int(zeile[0][1:])
            assert kennung%2==0 and kennung<vorherige
            vorherige=kennung
        vorschau=d.vorschau(gelesen)
        assert len(vorschau["zeilen"])==min(200,erwartet)
        assert _sha(quelle)==quellhash
        bericht={"art":"CSV","profil":profil,"status":"bestanden","eingabezeilen":anzahl,"gefilterte_zeilen":erwartet,"gruppenzahl":len(ergebnis),"gesamtbetrag":str(sum(z[2] for z in ergebnis)),"quelle_bytes":quelle.stat().st_size,"output_bytes":output["bytes"],"output_sha256":output["sha256"],"duckdb_speicher":"64MB","max_temp_bytes":grenzen.max_temp_bytes,"pruefung":pruefung,"dauer_sekunden":round(time.perf_counter()-start,3),**messung.json()}
    return bericht


def _xlsx(basis,anzahl,zeitlimit):
    from openpyxl import Workbook,load_workbook
    from openpyxl.cell import WriteOnlyCell
    from scriptlab_business.kontext import Kontext,Grenzen
    from scriptlab_business.daten import DatenEngine
    quelle=basis/"gross.xlsx"
    wb=Workbook(write_only=True);ws=wb.create_sheet("Daten");ws.append(["ID","Name","Betrag","Datum","Kategorie","Hinweis"])
    for i in range(anzahl):
        hinweis=WriteOnlyCell(ws,value="=bleibt Text" if i%100==0 else "synthetisch");hinweis.data_type="s"
        ws.append([f"{i:08}",f"Prüfung Ä {i}","1.25","2026-10-01",f"K{i%10}",hinweis])
    wb.save(quelle);wb.close()
    vorher=_sha(quelle)
    start=time.perf_counter()
    grenzen=Grenzen(max_zeilen=max(anzahl,100000),max_datei_bytes=512*1024**2,duckdb_speicher="64MB",max_temp_bytes=2*1024**3,zeitlimit_sekunden=zeitlimit)
    with Kontext(basis,"ausfuehren",grenzen=grenzen,workflow_hash="synthetische-grosse-xlsx") as k,Messung(k) as messung:
        d=DatenEngine(k);k.dienste["daten"]=d
        messung.phase="xlsx_import"
        t=d.importieren("XLSX",quelle.name)
        t=d.spalten(t,{"aktion":"TYP","spalte":"Betrag","typ":"DEZIMAL"})
        t=d.spalten(t,{"aktion":"TYP","spalte":"Datum","typ":"DATUM"})
        k.beginn_schritt("xlsx-export",{}, {})
        messung.phase="xlsx_export"
        output=d.exportieren(t,"XLSX","grosse-ausgabe.xlsx")
        gesamt=d.gruppieren(t,[],[{"funktion":"ANZAHL","name":"N"},{"funktion":"SUMME","spalte":"Betrag","name":"Summe"}])
        assert list(d.zeilen(gesamt))==[(anzahl,Decimal("1.25")*anzahl)]
        wb=load_workbook(basis/"grosse-ausgabe.xlsx",read_only=True,data_only=False)
        try:
            zeilen=wb.active.iter_rows(values_only=True);next(zeilen)
            erste=next(zeilen)
            assert erste[0]=="00000000" and erste[5]=="=bleibt Text"
            assert 1+sum(1 for _ in zeilen)==anzahl
        finally: wb.close()
        assert _sha(quelle)==vorher
        bericht={"art":"XLSX","status":"bestanden","eingabezeilen":anzahl,"outputzeilen":anzahl,"gesamtbetrag":str(Decimal("1.25")*anzahl),"quelle_bytes":quelle.stat().st_size,"output_bytes":output["bytes"],"output_sha256":output["sha256"],"duckdb_speicher":"64MB","dauer_sekunden":round(time.perf_counter()-start,3),**messung.json()}
    return bericht


def _beispiele(basis,zeitlimit):
    from scriptlab_business.kontext import Kontext
    from scriptlab_business.ausfuehrung import ausfuehren
    from pypdf import PdfReader
    from openpyxl import load_workbook
    schrift=Path(os.environ.get("WINDIR","C:/Windows"))/"Fonts"/"arial.ttf"
    if not schrift.is_file(): raise ValueError("Für die PDF-Probe fehlt eine lokale berechtigt vorhandene Arial-TTF.")
    berichte=[]
    for name in ("A-Monatsauswertung","B-Bestellungen"):
        ziel=basis/name;shutil.copytree(WURZEL/"beispiele"/"business"/name,ziel)
        plan=json.loads((ziel/"plan.json").read_text(encoding="utf-8"))
        parameter=json.loads((ziel/"parameter.json").read_text(encoding="utf-8"))
        parameter["Schriftdatei"]=str(schrift)
        with Kontext(ziel,"ausfuehren",parameter,grenzen={"zeitlimit_sekunden":zeitlimit}) as k:
            result=ausfuehren(plan,k)
            assert result["status"]==("erfolgreich" if name.startswith("A") else "mit_fehlern")
            if name.startswith("B"):
                assert k.ergebnisse["b-pruefen"]["regeln"][0]["fehler"]==1
                assert k.ergebnisse["b-fehlerhinweis"]["zeilen"][0][:2]==["B4","K9"]
        xlsx=ziel/parameter.get("XLSX_Ziel",parameter.get("Ausgabedatei"))
        wb=load_workbook(xlsx,read_only=True,data_only=False)
        try:
            zeilen=wb.active.iter_rows(values_only=True);next(zeilen)
            werte={z[0]:str(Decimal(str(z[2]))) for z in zeilen}
        finally: wb.close()
        assert {k:Decimal(v) for k,v in werte.items()}==({"Nord":Decimal(150),"Süd":Decimal(80),"Gesamt":Decimal(230)} if name.startswith("A") else {"Nord":Decimal(45),"Süd":Decimal(40),"Gesamt":Decimal(85)})
        bericht={"beispiel":name,"status":result["status"],"xlsx_sha256":_sha(xlsx),"summen":werte}
        if name.startswith("A"):
            pdf=ziel/parameter["PDF_Ziel"]
            with pdf.open("rb") as datei:
                reader=PdfReader(datei);text=reader.pages[0].extract_text()
                assert len(reader.pages)==1
            assert all(t in text for t in ("Nord","Süd","Gesamt","150.00","80.00","230.00"))
            bericht.update(pdf_sha256=_sha(pdf),pdf_seiten=1,pdf_text_geprueft=True)
        berichte.append(bericht)
    return {"art":"BEISPIELE","status":"bestanden","beispiele":berichte}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--art",choices=("csv","xlsx","beispiele","alle"),default="alle")
    parser.add_argument("--zeilen",type=int,default=1000000)
    parser.add_argument("--xlsx-zeilen",type=int,default=50000)
    parser.add_argument("--zeitlimit",type=int,default=120)
    parser.add_argument("--profil",choices=("breit","schmal"),default="breit")
    parser.add_argument("--intern",type=Path,help=argparse.SUPPRESS)
    args=parser.parse_args()
    if not 1<=args.zeilen<=5000000 or not 1<=args.xlsx_zeilen<=1048575 or not 1<=args.zeitlimit<=600:
        parser.error("Zeilenanzahl oder Zeitlimit liegen außerhalb der festen Prüfgrenzen.")
    archiv=(WURZEL/"artifacts"/"business-qa").resolve()
    if args.intern:
        basis=args.intern.resolve()
        if not basis.is_relative_to(archiv) or args.art=="alle": parser.error("Ungültiges technisches Prüfziel.")
        prozess_temp=basis/"prozess-temp";prozess_temp.mkdir()
        for name in ("TEMP","TMP","TMPDIR"): os.environ[name]=str(prozess_temp)
        tempfile.tempdir=str(prozess_temp)
        start=time.perf_counter()
        try:
            bericht={"csv":lambda:_csv(basis,args.zeilen,args.zeitlimit,args.profil),"xlsx":lambda:_xlsx(basis,args.xlsx_zeilen,args.zeitlimit),"beispiele":lambda:_beispiele(basis,args.zeitlimit)}[args.art]()
        except Exception as fehler:
            bericht={"art":args.art.upper(),"profil":args.profil,"status":"fehlgeschlagen","fehlerart":type(fehler).__name__,"meldung":str(fehler),**_speicher(),**getattr(fehler,"business_messung",{})}
        bericht["gesamt_dauer_sekunden"]=round(time.perf_counter()-start,3)
        (basis/"ergebnis.json").write_text(json.dumps(bericht,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print(json.dumps(bericht,ensure_ascii=True),flush=True)
        return 0 if bericht["status"]=="bestanden" else 1
    lauf=archiv/("daten-"+uuid.uuid4().hex);lauf.mkdir(parents=True)
    ergebnisse=[]
    for art in (("beispiele","csv","xlsx") if args.art=="alle" else (args.art,)):
        ziel=lauf/art;ziel.mkdir()
        start=time.perf_counter()
        try:
            exitcode,stdout,stderr=_job(args,ziel,art)
            if (ziel/"ergebnis.json").is_file():
                bericht=json.loads((ziel/"ergebnis.json").read_text(encoding="utf-8"))
            else:
                bericht={"art":art.upper(),"status":"prozessfehler","meldung":"Der Prüfprozess hat keinen fertigen Bericht geschrieben."}
            bericht["exitcode"]=exitcode
            if stderr: (ziel/"technische-fehler.txt").write_bytes(stderr[:65536])
        except subprocess.TimeoutExpired:
            bericht={"art":art.upper(),"status":"zeitlimit","harte_frist_sekunden":args.zeitlimit+30}
        bericht["controller_dauer_sekunden"]=round(time.perf_counter()-start,3)
        bericht["nachweisordner"]=str(ziel.relative_to(WURZEL)).replace("\\","/")
        ergebnisse.append(bericht)
        print(json.dumps(bericht,ensure_ascii=True),flush=True)
    zusammenfassung={"version":1,"python":sys.version.split()[0],"plattform":sys.platform,"ergebnisse":ergebnisse}
    (lauf/"zusammenfassung.json").write_text(json.dumps(zusammenfassung,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(str((lauf/"zusammenfassung.json").relative_to(WURZEL)),flush=True)
    return 0 if all(e["status"]=="bestanden" for e in ergebnisse) else 1


if __name__=="__main__":
    raise SystemExit(main())
