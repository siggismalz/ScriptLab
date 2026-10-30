"""Dokumentwerte bleiben Daten; Prüfungen lesen die fertigen Dateien erneut ein."""
from __future__ import annotations

from pathlib import Path
import random
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "laufzeit"))

from docx import Document
from docx.enum.section import WD_SECTION
from openpyxl import Workbook, load_workbook

from scriptlab_business.dokumente import docx_vorlage, xlsx_vorlage
from scriptlab_business.kontext import BusinessFehler, Kontext


class DokumentwerteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ScriptLab-Dokumentwerte-")
        self.addCleanup(self.temp.cleanup)
        self.basis = Path(self.temp.name).resolve()

    def kontext(self):
        kontext = Kontext(self.basis, "ausfuehren", workflow_hash="dokumentwerte")
        self.addCleanup(kontext.schliessen)
        return kontext

    def excelvorlage(self):
        pfad = self.basis / "Vorlage.xlsx"
        buch = Workbook()
        buch.active["A1"] = "Unveränderte Vorlage"
        buch.save(pfad)
        buch.close()
        return pfad

    def test_excel_text_an_der_zellgrenze_bleibt_vollstaendig_und_text(self):
        vorlage = self.excelvorlage()
        original = vorlage.read_bytes()
        text = "=" + "".join(random.Random(26).choices("abcdefghijklmnopqrstuvwxyz0123456789äöü", k=32766))
        bericht = xlsx_vorlage(self.kontext(), {
            "vorlage": vorlage.name, "ziel": "Ergebnis.xlsx", "zellen": {"A1": text},
        }, {})
        self.assertEqual(bericht["status"], "publiziert")
        buch = load_workbook(self.basis / "Ergebnis.xlsx", data_only=False)
        try:
            self.assertEqual(buch.active["A1"].value, text)
            self.assertEqual(buch.active["A1"].data_type, "s")
        finally:
            buch.close()
        self.assertEqual(vorlage.read_bytes(), original)

    def test_excel_zu_langer_text_verhindert_publikation_und_erhaelt_originale(self):
        vorlage = self.excelvorlage()
        original = vorlage.read_bytes()
        for vorhanden in (False, True):
            with self.subTest(vorhandenes_ziel=vorhanden):
                ziel = self.basis / ("Vorhanden.xlsx" if vorhanden else "Neu.xlsx")
                if vorhanden:
                    ziel.write_bytes(original)
                with self.kontext() as kontext:
                    with self.assertRaisesRegex(BusinessFehler, "32767"):
                        xlsx_vorlage(kontext, {
                            "vorlage": vorlage.name, "ziel": ziel.name,
                            "zellen": {"A1": "x" * 32768}, "konflikt": "ERSETZEN",
                        }, {})
                    self.assertEqual(kontext.ausgaben, [])
                    self.assertEqual(vorlage.read_bytes(), original)
                    if vorhanden:
                        self.assertEqual(ziel.read_bytes(), original)
                    else:
                        self.assertFalse(ziel.exists())
                    self.assertEqual(list(self.basis.glob(".scriptlab-*")), [])

    def test_word_positionswerte_werden_nicht_erneut_als_vorlage_gelesen(self):
        vorlage = self.basis / "Positionen.docx"
        dokument = Document()
        dokument.add_paragraph("Kunde {{Kunde}}")
        tabelle = dokument.add_table(rows=2, cols=1)
        tabelle.cell(0, 0).text = "Text"
        tabelle.cell(1, 0).paragraphs[0].add_run("{{Position.Text}}").bold = True
        dokument.save(vorlage)
        original = vorlage.read_bytes()
        texte = ["{{Kunde}}", "{{Unbekannt}}", "Unvollständig {{", "}} wörtlich"]
        bericht = docx_vorlage(self.kontext(), {
            "vorlage": vorlage.name, "ziel": "Positionen-fertig.docx",
            "werte": {"Kunde": "Lea"}, "positionen": {"tabelle": 0, "zeile": 1},
        }, {"positions_zeilen": [{"Text": text} for text in texte]})
        self.assertEqual(bericht["status"], "publiziert")
        ergebnis = Document(self.basis / "Positionen-fertig.docx")
        self.assertEqual(ergebnis.paragraphs[0].text, "Kunde Lea")
        self.assertEqual([zeile.cells[0].text for zeile in ergebnis.tables[0].rows[1:]], texte)
        self.assertTrue(ergebnis.tables[0].cell(1, 0).paragraphs[0].runs[0].bold)
        self.assertEqual(vorlage.read_bytes(), original)

    def test_excel_zu_lange_formel_verhindert_publikation_und_erhaelt_originale(self):
        vorlage = self.excelvorlage()
        original = vorlage.read_bytes()
        formel = "=" + "1+" * 16383 + "1"
        self.assertEqual(len(formel), 32768)
        for vorhanden in (False, True):
            with self.subTest(vorhandenes_ziel=vorhanden):
                ziel = self.basis / ("Formel-vorhanden.xlsx" if vorhanden else "Formel-neu.xlsx")
                if vorhanden:
                    ziel.write_bytes(original)
                with self.kontext() as kontext:
                    with self.assertRaisesRegex(BusinessFehler, "32767"):
                        xlsx_vorlage(kontext, {
                            "vorlage": vorlage.name, "ziel": ziel.name,
                            "formeln": {"A1": formel}, "konflikt": "ERSETZEN",
                        }, {})
                    self.assertEqual(kontext.ausgaben, [])
                    self.assertEqual(vorlage.read_bytes(), original)
                    if vorhanden:
                        self.assertEqual(ziel.read_bytes(), original)
                    else:
                        self.assertFalse(ziel.exists())
                    self.assertEqual(list(self.basis.glob(".scriptlab-*")), [])

    def test_word_verbundene_positionszelle_wird_nur_einmal_ersetzt(self):
        vorlage = self.basis / "Verbundene-Zellen.docx"
        dokument = Document()
        tabelle = dokument.add_table(rows=2, cols=2)
        tabelle.cell(1, 0).merge(tabelle.cell(1, 1)).text = "{{Position.Text}}"
        dokument.save(vorlage)
        bericht = docx_vorlage(self.kontext(), {
            "vorlage": vorlage.name, "ziel": "Verbundene-Zellen-fertig.docx",
            "werte": {"Kunde": "Lea"}, "positionen": {"tabelle": 0, "zeile": 1},
        }, {"positions_zeilen": [{"Text": "{{Kunde}}"}]})
        self.assertEqual(bericht["status"], "publiziert")
        ergebnis = Document(self.basis / "Verbundene-Zellen-fertig.docx")
        self.assertEqual(ergebnis.tables[0].cell(1, 0).text, "{{Kunde}}")
        self.assertEqual(ergebnis.tables[0].cell(1, 1).text, "{{Kunde}}")

    def test_word_geteilte_kopf_und_fusszeilen_werden_nur_einmal_ersetzt(self):
        vorlage = self.basis / "Abschnitte.docx"
        dokument = Document()
        dokument.sections[0].header.paragraphs[0].text = "Kopf {{Kunde}}"
        dokument.sections[0].footer.paragraphs[0].text = "Fuß {{Kunde}}"
        dokument.add_section(WD_SECTION.NEW_PAGE)
        dokument.save(vorlage)
        original = vorlage.read_bytes()
        bericht = docx_vorlage(self.kontext(), {
            "vorlage": vorlage.name, "ziel": "Abschnitte-fertig.docx",
            "werte": {"Kunde": "{{Unbekannt}}"},
        }, {})
        self.assertEqual(bericht["status"], "publiziert")
        ergebnis = Document(self.basis / "Abschnitte-fertig.docx")
        for abschnitt in ergebnis.sections:
            self.assertEqual(abschnitt.header.paragraphs[0].text, "Kopf {{Unbekannt}}")
            self.assertEqual(abschnitt.footer.paragraphs[0].text, "Fuß {{Unbekannt}}")
        self.assertEqual(vorlage.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
