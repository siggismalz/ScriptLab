"""Fachprüfungen der ausgelieferten synthetischen Beispiele A und B."""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from decimal import Decimal

from scriptlab_business.ausfuehrung import ausfuehren
from scriptlab_business.kontext import Kontext
from scriptlab_business.typen import DatenFehler

BEISPIELE = Path(__file__).resolve().parents[2] / "beispiele" / "business"


def fingerabdruecke(ordner):
    return {str(p.relative_to(ordner)):hashlib.sha256(p.read_bytes()).hexdigest() for p in (ordner/"eingaben").iterdir() if p.is_file()}


def xlsx_ergebnisse(pfad):
    from openpyxl import load_workbook
    wb=load_workbook(pfad,read_only=True,data_only=False)
    try:
        zeilen=wb.active.iter_rows(values_only=True)
        kopf=next(zeilen)
        return kopf,{z[0]:(int(z[1]),Decimal(str(z[2]))) for z in zeilen}
    finally:
        wb.close()


class BeispielDatenTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix="ScriptLab-Beispiel-Daten-")
        self.basis=Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def vorbereiten(self,name):
        quelle=BEISPIELE/name
        shutil.copytree(quelle,self.basis,dirs_exist_ok=True)
        plan=json.loads((self.basis/"plan.json").read_text(encoding="utf-8"))
        parameter=json.loads((self.basis/"parameter.json").read_text(encoding="utf-8"))
        return plan,parameter

    def test_a_monat_in_xlsx_und_pdf_mit_exakten_sollwerten(self):
        plan,parameter=self.vorbereiten("A-Monatsauswertung")
        schrift=Path(os.environ.get("WINDIR","C:/Windows"))/"Fonts"/"arial.ttf"
        if not schrift.is_file():
            self.skipTest("Für diesen realen PDF-Test fehlt eine berechtigt vorhandene lokale Arial-TTF.")
        parameter["Schriftdatei"]=str(schrift)
        vorher=fingerabdruecke(self.basis)
        with Kontext(self.basis,"ausfuehren",parameter) as kontext:
            ergebnis=ausfuehren(plan,kontext)
            self.assertEqual(ergebnis["status"],"erfolgreich")
            dienst=kontext.dienste["daten"]
            monat=dienst.tabelle(kontext.ergebnisse["a-monat"])
            self.assertEqual(len(list(dienst.zeilen(monat))),3)
        kopf,werte=xlsx_ergebnisse(self.basis/parameter["XLSX_Ziel"])
        self.assertEqual(kopf,("Region","Datensätze","Betrag"))
        self.assertEqual(werte,{"Nord":(2,Decimal("150.00")),"Süd":(1,Decimal("80.00")),"Gesamt":(3,Decimal("230.00"))})
        from pypdf import PdfReader
        with (self.basis/parameter["PDF_Ziel"]).open("rb") as datei:
            reader=PdfReader(datei)
            self.assertEqual(len(reader.pages),1)
            text=reader.pages[0].extract_text()
        for teil in ("Nord","Süd","150.00","80.00","Gesamt","230.00","Datensätze"):
            self.assertIn(teil,text)
        self.assertNotIn("70.00",text)
        self.assertEqual(fingerabdruecke(self.basis),vorher)

    def test_a_vorschau_und_prueflauf_schreiben_keine_berichte(self):
        plan,parameter=self.vorbereiten("A-Monatsauswertung")
        schrift=Path(os.environ.get("WINDIR","C:/Windows"))/"Fonts"/"arial.ttf"
        if not schrift.is_file(): self.skipTest("Lokale TTF für PDF-Prüflauf fehlt.")
        parameter["Schriftdatei"]=str(schrift)
        vorher=fingerabdruecke(self.basis)
        for modus in ("vorschau","prueflauf"):
            with self.subTest(modus=modus), Kontext(self.basis,modus,parameter) as kontext:
                ergebnis=ausfuehren(plan,kontext)
                self.assertEqual(ergebnis["status"],"erfolgreich")
                self.assertEqual(kontext.ergebnisse["a-xlsx"]["status"],"geplant")
                self.assertEqual(kontext.ergebnisse["a-pdf"]["status"],"geplant")
            self.assertFalse((self.basis/parameter["XLSX_Ziel"]).exists())
            self.assertFalse((self.basis/parameter["PDF_Ziel"]).exists())
        self.assertEqual(fingerabdruecke(self.basis),vorher)

    def test_b_k9_hinweis_drei_zuordnungen_und_85_gesamt(self):
        plan,parameter=self.vorbereiten("B-Bestellungen")
        vorher=fingerabdruecke(self.basis)
        with Kontext(self.basis,"ausfuehren",parameter) as kontext:
            ergebnis=ausfuehren(plan,kontext)
            self.assertEqual(ergebnis["status"],"mit_fehlern")
            self.assertEqual(kontext.ergebnisse["b-pruefen"]["regeln"][0]["fehler"],1)
            fehlend=kontext.ergebnisse["b-fehlerhinweis"]["zeilen"]
            self.assertEqual(len(fehlend),1)
            self.assertEqual(fehlend[0][:2],["B4","K9"])
            dienst=kontext.dienste["daten"]
            self.assertEqual(len(list(dienst.zeilen(dienst.tabelle(kontext.ergebnisse["b-gueltig"])))),3)
        kopf,werte=xlsx_ergebnisse(self.basis/parameter["Ausgabedatei"])
        self.assertEqual(kopf,("Region","Bestellungen","Betrag"))
        self.assertEqual(werte,{"Nord":(2,Decimal("45.00")),"Süd":(1,Decimal("40.00")),"Gesamt":(3,Decimal("85.00"))})
        self.assertEqual(fingerabdruecke(self.basis),vorher)

    def test_b_doppelter_k1_stoppt_vor_ausgabe(self):
        plan,parameter=self.vorbereiten("B-Bestellungen")
        parameter["Kundendatei"]="eingaben/kunden-doppelt.csv"
        vorher=fingerabdruecke(self.basis)
        with Kontext(self.basis,"ausfuehren",parameter) as kontext:
            with self.assertRaisesRegex(DatenFehler,"nicht eindeutig"):
                ausfuehren(plan,kontext)
            self.assertEqual(kontext.fehler[-1]["block_id"],"b-zuordnen")
        self.assertFalse((self.basis/parameter["Ausgabedatei"]).exists())
        self.assertEqual(fingerabdruecke(self.basis),vorher)


if __name__=="__main__":
    unittest.main()
