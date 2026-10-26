"""Reale Datenengine-Tests, ausschließlich mit selbst erzeugten Tempdateien.

Start: <Business-Python> -m unittest discover -s tests/business -p test_daten.py -v
PYTHONPATH muss auf den lokalen Ordner laufzeit zeigen.
"""

from __future__ import annotations

import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
import tempfile
import time
import unittest

from scriptlab_business.daten import DatenEngine, HANDLER
from scriptlab_business.kontext import BusinessFehler, Kontext, Grenzen
from scriptlab_business.typen import DatenFehler


def wert(wert, typ="TEXT"):
    return {"art": "wert", "typ": typ, "wert": wert}


def vergleich(spalte, operator, vergleichswert=None):
    ast = {"art": "vergleich", "spalte": spalte, "operator": operator}
    if vergleichswert is not None:
        ast["wert"] = vergleichswert
    return ast


class DatenTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ScriptLab-Daten-Test-")
        self.basis = Path(self.temp.name).resolve()
        self.meldungen = []
        self.kontext = Kontext(self.basis, "ausfuehren", workflow_hash="fester-datentest", ereignis_callback=self.meldungen.append)
        self.engine = DatenEngine(self.kontext)
        self.kontext.dienste["daten"] = self.engine

    def tearDown(self):
        self.kontext.schliessen()
        self.temp.cleanup()

    def csv(self, inhalt, name="quelle.csv", optionen=None):
        (self.basis/name).write_text(inhalt, encoding="utf-8", newline="")
        return self.engine.importieren("CSV", name, optionen)

    def zeilen(self, tabelle):
        return list(self.engine.zeilen(tabelle, 2))

    def typ(self, tabelle, spalte, typ):
        return self.engine.spalten(tabelle, {"aktion":"TYP", "spalte":spalte, "typ":typ})

    def test_csv_bewahrt_kennungen_unicode_zitate_und_leerwerte(self):
        t = self.csv('ID;Name;Wert\r\n001;"Müller; 😀";\r\n002;"Zitat ""gut""";7\r\n')
        self.assertEqual(self.zeilen(t), [("001", "Müller; 😀", None), ("002", 'Zitat "gut"', "7")])
        self.assertTrue(all(s.typ == "VARCHAR" for s in t.spalten))

    def test_csv_utf16_und_cp1252_streaming(self):
        for kodierung, codec in (("UTF16","utf-16"),("CP1252","cp1252")):
            with self.subTest(kodierung=kodierung):
                p = self.basis/f"{kodierung}.csv"
                p.write_text("ID;Name\r\n001;Ähre €\r\n", encoding=codec, newline="")
                t = self.engine.importieren("CSV",p.name,{"kodierung":kodierung})
                self.assertEqual(self.zeilen(t), [("001","Ähre €")])

    def test_csv_doppelte_kopfzeilen_werden_nicht_still_umbenannt(self):
        with self.assertRaises(DatenFehler): self.csv("ID;id\n1;2\n")

    def test_csv_explizite_spalten_ohne_header(self):
        t = self.csv("001;Äpfel\n",optionen={"kopfzeile":False,"spalten":["ID","Name"]})
        self.assertEqual(self.zeilen(t),[("001","Äpfel")])

    def test_csv_fehlzeilen_stoppen_oder_werden_gemeldet(self):
        for policy in ("STOPP","AUSLASSEN_MIT_BERICHT"):
            with self.subTest(policy=policy):
                if policy == "STOPP":
                    with self.assertRaises((DatenFehler, BusinessFehler)):
                        self.zeilen(self.csv("ID;Wert\n1;2\n3;4;5\n6;7\n",optionen={"fehlzeilen":policy}))
                else:
                    t = self.csv("ID;Wert\n1;2\n3;4;5\n6;7\n",optionen={"fehlzeilen":policy})
                    self.assertEqual(self.zeilen(t),[("1","2"),("6","7")])
                    self.assertEqual(self.meldungen[-1]["anzahl"],1)

    def test_typen_decimal_datum_boolean_null(self):
        t = self.csv("Betrag;Datum;Ja\n123456789.25;2024-02-29;true\n;;false\n")
        for sp,typ in (("Betrag",{"art":"DEZIMAL","praezision":18,"skala":2}),("Datum","DATUM"),("Ja","WAHRHEIT")):
            t = self.typ(t,sp,typ)
        self.assertEqual(self.zeilen(t),[(Decimal("123456789.25"),date(2024,2,29),True),(None,None,False)])

    def test_unpassender_typ_schlaegt_bei_auswertung_fehl(self):
        t = self.typ(self.csv("Wert\nnicht-zahl\n"),"Wert","GANZZAHL")
        with self.assertRaises(DatenFehler): self.zeilen(t)

    def test_identifier_und_wertinjektion_bleiben_daten(self):
        name = 'Spalte"; DROP TABLE fremd;--'
        datei = self.basis/"injektion.csv"
        with datei.open("w",encoding="utf-8",newline="") as f:
            writer=csv.writer(f,delimiter=";");writer.writerow([name]);writer.writerows([["x' OR 1=1 --"],["normal"]])
        t=self.engine.importieren("CSV",datei.name)
        t=self.engine.filtern(t,vergleich(name,"EQ",wert("x' OR 1=1 --")))
        self.assertEqual(self.zeilen(t),[("x' OR 1=1 --",)])
        with self.assertRaises(DatenFehler): self.engine.spalten(t,{"aktion":"TYP","spalte":name,"typ":"INT);DROP TABLE x;--"})

    def test_komplexe_filter_und_oder_nicht(self):
        t=self.typ(self.csv("Name;Zahl\nAnna;5\nBerta;10\nClaus;15\n"),"Zahl","GANZZAHL")
        t=self.engine.filtern(t,{"art":"UND","kinder":[{"art":"ODER","kinder":[vergleich("Name","BEGINNT",wert("A")),vergleich("Name","ENDET",wert("s"))]},{"art":"NICHT","kind":vergleich("Zahl","LT",wert(10,"GANZZAHL"))}]})
        self.assertEqual(self.zeilen(t),[("Claus",15)])

    def test_alle_vergleiche_und_in(self):
        t=self.typ(self.csv("Name;N\nAlpha;1\nBeta;2\nGamma;3\n"),"N","GANZZAHL")
        for op,n in (("EQ",1),("NE",2),("GT",1),("GE",2),("LT",1),("LE",2)):
            with self.subTest(op=op):
                self.assertEqual(len(self.zeilen(self.engine.filtern(t,vergleich("N",op,wert(2,"GANZZAHL"))))),n)
        self.assertEqual(len(self.zeilen(self.engine.filtern(t,vergleich("Name","ENTHAELT",wert("mm"))))),1)
        self.assertEqual(len(self.zeilen(self.engine.filtern(t,{"art":"vergleich","spalte":"N","operator":"IN","werte":[wert(1,"GANZZAHL"),wert(3,"GANZZAHL")]}))),2)

    def test_leerwerte_nur_ausdruecklich_vergleichen(self):
        t=self.csv("A;B\n;1\nx;2\n")
        self.assertEqual(len(self.zeilen(self.engine.filtern(t,vergleich("A","IST_LEER")))),1)
        self.assertEqual(len(self.zeilen(self.engine.filtern(t,vergleich("A","IST_NICHT_LEER")))),1)
        with self.assertRaises(DatenFehler): self.engine.filtern(t,vergleich("A","EQ",wert(None)))

    def test_ausdrucksgrenzen_und_freie_formel_abgewiesen(self):
        t=self.csv("A\nx\n")
        ast=vergleich("A","EQ",wert("x"))
        for _ in range(25): ast={"art":"NICHT","kind":ast}
        with self.assertRaises(DatenFehler): self.engine.filtern(t,ast)
        with self.assertRaises(DatenFehler): self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"X","ausdruck":{"art":"python","code":"print(1)"}})

    def test_decimal_subtraktion_ohne_floatverlust(self):
        t=self.csv("Brutto;Rabatt\n0.30;0.10\n")
        for name in ("Brutto","Rabatt"): t=self.typ(t,name,"DEZIMAL")
        t=self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"Netto","ausdruck":{"art":"rechnen","operator":"SUB","links":{"art":"spalte","name":"Brutto"},"rechts":{"art":"spalte","name":"Rabatt"}}})
        self.assertEqual(self.zeilen(t)[0][2],Decimal("0.20"))

    def test_rechenoperatoren_und_division_durch_null(self):
        t=self.csv("A\nx\n")
        for op, erwartet in (("ADD",8),("SUB",4),("MUL",12),("DIV",3)):
            with self.subTest(op=op):
                r=self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"R","ausdruck":{"art":"rechnen","operator":op,"links":wert(6,"GANZZAHL"),"rechts":wert(2,"GANZZAHL")}})
                self.assertEqual(self.zeilen(r)[0][1],erwartet)
        r=self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"R","ausdruck":{"art":"rechnen","operator":"DIV","links":wert(1,"GANZZAHL"),"rechts":wert(0,"GANZZAHL")}})
        with self.assertRaises(DatenFehler): self.zeilen(r)

    def test_datum_text_und_regex_funktionen(self):
        t=self.csv("Text;Datum\n  AB-123  ;2024-02-29\n")
        t=self.typ(t,"Datum","DATUM")
        for funktion,spalte,erwartet,extra in (("JAHR","Datum",2024,{}),("MONAT","Datum",2,{}),("TAG","Datum",29,{}),("TRIM","Text","AB-123",{}),("GROSS","Text","  AB-123  ",{}),("KLEIN","Text","  ab-123  ",{}),("LAENGE","Text",10,{}),("REGEX","Text","123",{"muster":r"AB-(\d+)","gruppe":1})):
            with self.subTest(funktion=funktion):
                r=self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"R","ausdruck":{"art":"funktion","funktion":funktion,"argument":{"art":"spalte","name":spalte},**extra}})
                self.assertEqual(self.zeilen(r)[0][-1],erwartet)

    def test_typisierte_literalfehler_werden_ohne_sql_ausfuehrung_erkannt(self):
        t=self.csv("A\nx\n")
        for literal in (wert("1,2","DEZIMAL"),wert("123.45",{"art":"DEZIMAL","praezision":4,"skala":2}),wert(float("nan"),"ZAHL"),wert("31.02.2024","DATUM"),wert(1,"WAHRHEIT"),wert(True,"GANZZAHL")):
            with self.subTest(literal=literal):
                with self.assertRaises(DatenFehler): self.engine.spalten(t,{"aktion":"BERECHNEN","spalte":"R","ausdruck":literal})

    def test_spalten_waehlen_umbenennen_und_kollisionen(self):
        t=self.csv("A;B\n1;2\n")
        t=self.engine.spalten(t,{"aktion":"AUSWAEHLEN","spalten":["B","A"]})
        t=self.engine.spalten(t,{"aktion":"UMBENENNEN","namen":{"A":"Neu"}})
        self.assertEqual([s.name for s in t.spalten],["B","Neu"])
        self.assertEqual(self.zeilen(t),[("2","1")])
        for k in ({"aktion":"AUSWAEHLEN","spalten":["SQL"]},{"aktion":"UMBENENNEN","namen":{"Neu":"b"}}):
            with self.assertRaises(DatenFehler): self.engine.spalten(t,k)

    def test_bereinigen_ersetzen_und_leerwerte(self):
        t=self.csv("A;B\n  hallo  ;1\n;2\n")
        t=self.engine.werte(t,{"aktion":"BEREINIGEN","spalten":["A"],"grossklein":"GROSS"})
        t=self.engine.werte(t,{"aktion":"ERSETZEN","spalten":["A"],"suchen":wert("HALLO"),"ersatz":wert("WELT")})
        with self.assertRaises(DatenFehler): self.engine.werte(t,{"aktion":"LEERWERTE","spalten":["A"]})
        self.assertEqual(self.zeilen(self.engine.werte(t,{"aktion":"LEERWERTE","spalten":["A"],"leere":"AUSLASSEN"})),[("WELT","1")])
        self.assertEqual(self.zeilen(self.engine.werte(t,{"aktion":"LEERWERTE","spalten":["A"],"leere":"ERSETZEN","ersatz":wert("FEHLT")})),[("WELT","1"),("FEHLT","2")])

    def test_sortierung_leerwerte_und_dubletten(self):
        t=self.csv("ID;Wert\na;erster\nb;einzig\na;letzter\n;leer\n")
        r=self.engine.sortieren(t,[{"spalte":"ID","richtung":"DESC","nullen":"ERSTE"}])
        self.assertEqual(self.zeilen(r)[0],(None,"leer"))
        for behalten, erwartet in (("ERSTE","erster"),("LETZTE","letzter")):
            r=self.engine.dubletten(t,["ID"],behalten)
            self.assertIn(("a",erwartet),self.zeilen(r))
            self.assertEqual(len(self.zeilen(r)),3)
        self.assertEqual(len(self.zeilen(self.engine.dubletten(t,["ID"],"KEINE"))),2)

    def test_vereinigen_nach_namen_und_typpruefung(self):
        l=self.csv("A;B\n1;2\n",name="links.csv")
        r=self.csv("B;A\n4;3\n",name="rechts.csv")
        self.assertEqual(self.zeilen(self.engine.vereinigen(l,r)),[("1","2"),("3","4")])
        self.assertEqual(len(self.zeilen(self.engine.vereinigen(l,l,dubletten=True))),1)
        with self.assertRaises(DatenFehler): self.engine.vereinigen(l,self.typ(r,"A","GANZZAHL"))

    def test_join_kardinalitaet_fehlende_zuordnung_und_namenskollision(self):
        l=self.csv("ID;Name\n1;A\n2;B\n",name="links.csv")
        r=self.csv("ID;Name\n1;X\n1;Y\n",name="rechts.csv")
        k={"art":"LEFT","schluessel":[{"links":"ID","rechts":"ID"}],"fehlende_zuordnung":"BEHALTEN"}
        with self.assertRaises(DatenFehler): self.engine.verknuepfen(l,r,{**k,"kardinalitaet":"n:1"})
        with self.assertRaises(DatenFehler): self.engine.verknuepfen(l,r,{**k,"kardinalitaet":"1:n","fehlende_zuordnung":"STOPP"})
        t=self.engine.verknuepfen(l,r,{**k,"kardinalitaet":"1:n"})
        self.assertEqual([s.name for s in t.spalten],["ID","Name","ID_rechts","Name_rechts"])
        self.assertEqual(set(self.zeilen(t)),{("1","A","1","X"),("1","A","1","Y"),("2","B",None,None)})

    def test_alle_joinarten(self):
        l=self.csv("ID\n1\n2\n",name="links.csv")
        r=self.csv("ID\n2\n3\n",name="rechts.csv")
        for art,anzahl in (("INNER",1),("LEFT",2),("FULL",3)):
            with self.subTest(art=art):
                t=self.engine.verknuepfen(l,r,{"art":art,"schluessel":[{"links":"ID","rechts":"ID"}],"kardinalitaet":"1:1","fehlende_zuordnung":"BEHALTEN"})
                self.assertEqual(len(self.zeilen(t)),anzahl)

    def test_gruppieren_alle_aggregate(self):
        t=self.typ(self.csv("Gruppe;N\na;1\na;3\na;3\nb;\n"),"N","GANZZAHL")
        agg=[{"funktion":f,"spalte":"N","name":f} for f in ("ANZAHL","ANZAHL_WERTE","ANZAHL_VERSCHIEDEN","SUMME","MITTELWERT","MIN","MAX")]
        r=self.engine.gruppieren(t,["Gruppe"],agg)
        zeile=next(z for z in self.zeilen(r) if z[0]=="a")
        self.assertEqual(zeile[1:5],(3,3,2,7))
        self.assertAlmostEqual(zeile[5],7/3)
        self.assertEqual(zeile[6:],(1,3))

    def test_qualitaetsregeln_sind_kleine_zaehlberichte(self):
        t=self.csv("ID;N\na;1\na;x\n;3\n")
        bericht=self.engine.pruefen(t,[{"art":"NICHT_LEER","spalte":"ID"},{"art":"EINDEUTIG","spalten":["ID"]},{"art":"TYP","spalte":"N","typ":"GANZZAHL"},{"art":"ERLAUBTE_WERTE","spalte":"ID","werte":[wert("b")]},{"art":"ZEILENZAHL","minimum":4}])
        self.assertFalse(bericht["gueltig"])
        self.assertEqual([r["fehler"] for r in bericht["regeln"]],[1,1,1,2,1])
        with self.assertRaises(DatenFehler): self.engine.pruefen(t,[{"art":"NICHT_LEER","spalte":"ID"}],"STOPP")
        t=self.typ(self.csv("N\n1\n5\n10\n",name="bereich.csv"),"N","GANZZAHL")
        self.assertEqual(self.engine.pruefen(t,[{"art":"BEREICH","spalte":"N","minimum":wert(2,"GANZZAHL"),"maximum":wert(9,"GANZZAHL")}])["regeln"][0]["fehler"],2)

    def test_vorschau_hat_zeilen_und_bytebudget(self):
        t=self.csv("N\n"+"\n".join(str(i) for i in range(300))+"\n")
        bericht=self.engine.vorschau(t)
        self.assertEqual(len(bericht["zeilen"]),200)
        self.assertTrue(bericht["gekuerzt"])
        self.assertIn("dauer_ms",bericht)
        bericht=self.engine.vorschau(t,byte_limit=10)
        self.assertLessEqual(bericht["bytes"],10)
        self.assertTrue(bericht["gekuerzt"])
        with self.assertRaises(DatenFehler): self.engine.vorschau(t,201)

    def test_quellaenderung_nach_import_stoppt_auswertung_und_export(self):
        t=self.csv("ID\n1\n")
        (self.basis/"quelle.csv").write_text("ID\n2\n",encoding="utf-8")
        with self.assertRaises(BusinessFehler): self.zeilen(t)
        with self.assertRaises(BusinessFehler): self.engine.exportieren(t,"CSV","ziel.csv")
        self.assertFalse((self.basis/"ziel.csv").exists())

    def test_fremde_tabellenreferenzen_werden_abgewiesen(self):
        t=self.csv("ID\n1\n")
        with self.assertRaises(DatenFehler): self.engine.tabelle({"tabelle":"SELECT * FROM x"})
        self.assertEqual(self.engine.tabelle(t.json()),t)

    def test_csv_json_parquet_roundtrip(self):
        t=self.csv("ID;Text\n001;Ähre 😀\n002;\n")
        for format,ende,opt in (("CSV","csv",{}),("JSON","json",{}),("JSON","jsonl",{"zeilenformat":True}),("PARQUET","parquet",{})):
            with self.subTest(format=format,ende=ende):
                name=f"ausgabe.{ende}"
                self.kontext.beginn_schritt(name,{}, {})
                bericht=self.engine.exportieren(t,format,name,optionen=opt)
                self.assertEqual(bericht["status"],"publiziert")
                rueck=self.engine.importieren(format,name,opt)
                self.assertEqual(self.zeilen(rueck),self.zeilen(t))

    def test_xlsx_formelinjektion_als_text_exportiert_und_streaming_import(self):
        from openpyxl import load_workbook
        t=self.csv('ID;Text\n001;"=HYPERLINK(\"\"https://example.invalid\"\")"\n002;Ähre 😀\n')
        self.kontext.beginn_schritt("xlsx",{}, {})
        self.engine.exportieren(t,"XLSX","ausgabe.xlsx")
        wb=load_workbook(self.basis/"ausgabe.xlsx",read_only=True,data_only=False)
        try:
            z=list(wb.active.iter_rows())[1][1]
            self.assertEqual(z.data_type,"s")
            self.assertTrue(z.value.startswith("=HYPERLINK"))
        finally: wb.close()
        self.assertEqual(self.zeilen(self.engine.importieren("XLSX","ausgabe.xlsx")),self.zeilen(t))

    def test_xlsx_formelquelle_wird_standardmaessig_abgewiesen(self):
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.append(["Wert"]);ws.append(["=1+1"]);wb.save(self.basis/"formel.xlsx");wb.close()
        with self.assertRaises(DatenFehler): self.engine.importieren("XLSX","formel.xlsx")

    def test_ausgabekonflikte_und_prueflauf_beruehren_keine_ziele(self):
        t=self.csv("ID\n1\n")
        ziel=self.basis/"ziel.csv";ziel.write_bytes(b"ALT")
        self.kontext.beginn_schritt("konflikt",{}, {})
        with self.assertRaises(BusinessFehler): self.engine.exportieren(t,"CSV",ziel.name)
        self.assertEqual(ziel.read_bytes(),b"ALT")
        self.assertEqual(self.engine.exportieren(t,"CSV",ziel.name,"UEBERSPRINGEN")["status"],"uebersprungen")
        self.kontext.beginn_schritt("neuername",{}, {})
        neuer=self.engine.exportieren(t,"CSV",ziel.name,"NEUER_NAME")
        self.assertNotEqual(Path(neuer["pfad"]),ziel)
        self.assertEqual(ziel.read_bytes(),b"ALT")
        self.kontext.modus="prueflauf"
        self.assertEqual(self.engine.exportieren(t,"CSV",ziel.name,"ERSETZEN")["status"],"geplant")
        self.assertEqual(ziel.read_bytes(),b"ALT")

    def test_fehlgeschlagener_export_erhaelt_vorhandenes_ziel(self):
        t=self.csv("ID\n😀\n")
        ziel=self.basis/"ziel.csv";ziel.write_bytes(b"ALT")
        self.kontext.beginn_schritt("fehler",{}, {})
        with self.assertRaises((UnicodeError,BusinessFehler)): self.engine.exportieren(t,"CSV",ziel.name,"ERSETZEN",{"kodierung":"CP1252"})
        self.assertEqual(ziel.read_bytes(),b"ALT")

    def test_format_und_dateigroesse_werden_begrenzt(self):
        with self.assertRaises(DatenFehler): self.engine.importieren("SQL","quelle.sql")
        self.kontext.grenzen=Grenzen(max_datei_bytes=4)
        with self.assertRaises((DatenFehler,BusinessFehler)): self.csv("ID\n123456\n")

    def test_json_verschachtelte_werte_werden_abgewiesen(self):
        (self.basis/"quelle.json").write_text('[{"ID":1,"Objekt":{"a":2}}]',encoding="utf-8")
        with self.assertRaises(DatenFehler): self.engine.importieren("JSON","quelle.json")

    def test_alle_handler_sind_fest_und_ergebnis_jsonfaehig(self):
        t=self.csv("ID\n001\n")
        self.assertEqual(len(HANDLER),12)
        bericht=HANDLER["daten.vorschau"](self.kontext,{"limit":1},{"tabelle":t.json()})
        self.assertEqual(json.loads(json.dumps(bericht))["zeilen"],[["001"]])
        with self.assertRaises(DatenFehler): HANDLER["daten.filtern"](self.kontext,{"bedingung":"1=1"},{"tabelle":t.json()})

    def test_abbruch_wird_vor_auswertung_beachtet(self):
        t=self.csv("ID\n1\n")
        self.kontext.abgebrochen=True
        with self.assertRaises(BusinessFehler): self.zeilen(t)

    def test_deutsche_zahlen_und_datumsformate_sind_ausdruecklich(self):
        t=self.csv("Betrag;Datum\n1.234,50;29.02.2024\n-7,25;01.03.2024\n")
        t=self.engine.spalten(t,{"aktion":"TYP","spalte":"Betrag","typ":"DEZIMAL","dezimaltrennzeichen":",","tausendertrennzeichen":"."})
        t=self.engine.spalten(t,{"aktion":"TYP","spalte":"Datum","typ":"DATUM","datumsformat":"DMY"})
        self.assertEqual(self.zeilen(t),[(Decimal("1234.50"),date(2024,2,29)),(Decimal("-7.25"),date(2024,3,1))])

    def test_falsche_gruppierung_und_typfehlerpolicy(self):
        t=self.csv("Betrag\n1.23,45\n1.234,50\n")
        k={"aktion":"TYP","spalte":"Betrag","typ":"DEZIMAL","dezimaltrennzeichen":",","tausendertrennzeichen":"."}
        with self.assertRaises(DatenFehler): self.zeilen(self.engine.spalten(t,k))
        self.assertEqual(self.zeilen(self.engine.spalten(t,{**k,"fehlwerte":"NULL"})),[(None,),(Decimal("1234.50"),)])
        self.assertEqual(self.meldungen[-1]["anzahl"],1)
        self.assertEqual(self.zeilen(self.engine.spalten(t,{**k,"fehlwerte":"AUSLASSEN_MIT_BERICHT"})),[(Decimal("1234.50"),)])

    def test_textteile_ersetzen_ist_keine_regex_oder_codeausfuehrung(self):
        t=self.csv("Text\nAA.foo.BB\n")
        t=self.engine.werte(t,{"aktion":"ERSETZEN","spalten":["Text"],"modus":"TEXTTEIL","suchen":wert(".foo."),"ersatz":wert("' OR 1=1 --")})
        self.assertEqual(self.zeilen(t),[("AA' OR 1=1 --BB",)])

    def test_parquet_bewahrt_dezimal_datum_boolean_und_unicode_dateinamen(self):
        t=self.csv("Betrag;Datum;Ja\n12345678901234.25;2024-02-29;true\n")
        for sp,typ in (("Betrag","DEZIMAL"),("Datum","DATUM"),("Ja","WAHRHEIT")): t=self.typ(t,sp,typ)
        name="Ä 😀 ' ; plan.parquet"
        self.kontext.beginn_schritt("parquet-typen",{}, {})
        self.engine.exportieren(t,"PARQUET",name)
        rueck=self.engine.importieren("PARQUET",name)
        self.assertEqual([s.typ for s in rueck.spalten],[s.typ for s in t.spalten])
        self.assertEqual(self.zeilen(rueck),self.zeilen(t))

    def test_xlsx_decimal_und_formelzeichen_bleiben_text(self):
        from openpyxl import load_workbook
        t=self.csv("Wert\n+cmd\n-cmd\n@cmd\n=cmd\n")
        self.kontext.beginn_schritt("textzellen",{}, {})
        self.engine.exportieren(t,"XLSX","text.xlsx")
        wb=load_workbook(self.basis/"text.xlsx",read_only=True,data_only=False)
        try:
            self.assertEqual([(z[0].value,z[0].data_type) for z in list(wb.active.iter_rows())[1:]],[('+cmd','s'),('-cmd','s'),('@cmd','s'),('=cmd','s')])
        finally: wb.close()

    def test_zeilengrenze_und_xlsx_zellgrenze_verhindern_abgeschnittene_ausgabe(self):
        t=self.csv("ID\n1\n2\n3\n")
        self.kontext.grenzen=Grenzen(max_zeilen=2)
        self.kontext.beginn_schritt("grenze",{}, {})
        with self.assertRaises(DatenFehler): self.engine.exportieren(t,"CSV","ziel.csv")
        self.assertFalse((self.basis/"ziel.csv").exists())
        self.kontext.grenzen=Grenzen()
        t=self.csv("Text\n"+"x"*32768+"\n",name="lang.csv")
        self.kontext.beginn_schritt("zellgrenze",{}, {})
        with self.assertRaises(DatenFehler): self.engine.exportieren(t,"XLSX","lang.xlsx")
        self.assertFalse((self.basis/"lang.xlsx").exists())

    def test_20001_zeilen_werden_mit_mehreren_paketen_verarbeitet(self):
        t=self.csv("ID;Betrag\n"+"".join(f"{i:06};1.25\n" for i in range(20001)))
        t=self.typ(t,"Betrag","DEZIMAL")
        r=self.engine.gruppieren(t,[],[{"funktion":"ANZAHL","name":"Zeilen"},{"funktion":"SUMME","spalte":"Betrag","name":"Summe"}])
        self.assertEqual(self.zeilen(r),[(20001,Decimal("25001.25"))])
        self.kontext.beginn_schritt("gross",{}, {})
        self.engine.exportieren(t,"JSON","gross.jsonl",optionen={"zeilenformat":True})
        with (self.basis/"gross.jsonl").open(encoding="utf-8") as datei:
            self.assertEqual(sum(1 for _ in datei),20001)
        self.assertEqual(len(self.engine.vorschau(t)["zeilen"]),200)

    def test_query_interrupt_beendet_reale_arbeit_und_verbindung_bleibt_nutzbar(self):
        # Fest implementierte technische Prüfquery, kein öffentlicher Ausdruck.
        t=self.engine._neu("SELECT sum(i::HUGEINT*i::HUGEINT) FROM range(1000000000) AS t(i)")
        start=time.monotonic()
        with self.assertRaisesRegex(DatenFehler,"Zeitlimit"):
            list(self.engine.zeilen(t,zeitlimit=0.02))
        self.assertLess(time.monotonic()-start,3)
        self.assertEqual(self.zeilen(self.csv("ID\n1\n")),[("1",)])

    def test_duckdb_tempbudget_und_offline_erweiterungen_sind_gesetzt(self):
        con=self.engine._verbindung
        self.assertFalse(con.execute("SELECT current_setting('autoinstall_known_extensions')").fetchone()[0])
        self.assertFalse(con.execute("SELECT current_setting('autoload_known_extensions')").fetchone()[0])
        self.assertEqual(con.execute("SELECT current_setting('max_temp_directory_size')").fetchone()[0],"4.0 GiB")

    def test_quellpfade_bleiben_ueber_mehrere_plaene_erhalten(self):
        l=self.csv("ID\n1\n",name="links.csv")
        r=self.csv("ID\n2\n",name="rechts.csv")
        t=self.engine.vereinigen(l,r)
        self.assertEqual(set(self.engine.quellpfade(t)),{self.basis/"links.csv",self.basis/"rechts.csv"})

    def test_grosse_vorschauzellen_werden_vor_fetch_offen_gekuerzt(self):
        t=self.csv("Text\n"+"Ä"*10000+"\n")
        bericht=self.engine.vorschau(t)
        self.assertTrue(bericht["gekuerzt"])
        self.assertTrue(bericht["zellen_gekuerzt"])
        self.assertTrue(bericht["zeilen"][0][0].endswith("…"))
        self.assertLessEqual(len(bericht["zeilen"][0][0]),4097)
        self.assertEqual(len(self.zeilen(t)[0][0]),10000)

    def test_64mb_csv_reader_verarbeitet_groessere_quelle_und_join(self):
        self.engine.schliessen()
        self.kontext.grenzen=Grenzen(duckdb_speicher="64MB")
        self.engine=DatenEngine(self.kontext);self.kontext.dienste["daten"]=self.engine
        t=self.csv("ID;Text\n"+"".join(f"{i%2};{'Textdaten'*20}\n" for i in range(100000)),name="groesser.csv")
        r=self.csv("ID;Name\n0;Nord\n1;Süd\n",name="lookup.csv")
        t=self.engine.verknuepfen(t,r,{"schluessel":[{"links":"ID","rechts":"ID"}],"kardinalitaet":"n:1"})
        g=self.engine.gruppieren(t,["Name"],[{"funktion":"ANZAHL","name":"N"}])
        self.assertEqual(set(self.zeilen(g)),{("Nord",50000),("Süd",50000)})

    def test_xlsx_import_rollt_bei_spaeter_formel_alle_pakete_zurueck(self):
        from openpyxl import Workbook
        wb=Workbook(write_only=True);ws=wb.create_sheet("Daten");ws.append(["ID"])
        for i in range(1001): ws.append([str(i)])
        ws.append(["=1+1"]);wb.save(self.basis/"spaet.xlsx");wb.close()
        with self.assertRaises(DatenFehler): self.engine.importieren("XLSX","spaet.xlsx")
        # Kein teilweise eingelesener Tabellenhandle darf sichtbar werden.
        self.assertEqual(len(self.engine._tabellen),0)
        con=self.engine._verbindung
        self.assertEqual(con.execute("SELECT sum(estimated_size) FROM duckdb_tables() WHERE table_name LIKE 'sl_xlsx_%'").fetchone()[0],0)
        self.assertEqual(self.zeilen(self.csv("ID\nweiter\n")),[("weiter",)])

    def test_leere_bereichsgrenzen_duerfen_nicht_falsch_gueltig_sein(self):
        t=self.typ(self.csv("N\n1\n100\n"),"N","GANZZAHL")
        for grenze in ("minimum","maximum"):
            with self.subTest(grenze=grenze):
                with self.assertRaisesRegex(DatenFehler,"Bereichsgrenze darf nicht leer"):
                    self.engine.pruefen(t,[{"art":"BEREICH","spalte":"N",grenze:wert(None,"GANZZAHL")}])

    def test_xlsx_bereich_verwendet_eigene_kopfzeile_und_nur_ausgewaehlte_zellen(self):
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active
        ws["A1"]="Nicht Teil der Tabelle";ws["E3"]="Auch nicht Teil"
        ws["B2"]="ID";ws["C2"]="Name";ws["D2"]="Betrag"
        ws["B3"]="001";ws["C3"]="Äpfel";ws["D3"]="12.50"
        ws["B4"]="002";ws["C4"]="Birnen";ws["D4"]="7.25"
        ws["B6"]="=1+1";wb.save(self.basis/"bereich.xlsx");wb.close()
        t=self.engine.importieren("XLSX","bereich.xlsx",{"bereich":"$B$2:$D$4"})
        self.assertEqual([s.name for s in t.spalten],["ID","Name","Betrag"])
        self.assertEqual(self.zeilen(t),[("001","Äpfel","12.50"),("002","Birnen","7.25")])
        ohne_kopf=self.engine.importieren("XLSX","bereich.xlsx",{"bereich":"B3:D3","kopfzeile":False})
        self.assertEqual(self.zeilen(ohne_kopf),[("001","Äpfel","12.50")])

    def test_xlsx_ungueltige_unvollstaendige_und_umgekehrte_bereiche_stoppen(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.append(["ID"]);wb.active.append(["1"]);wb.save(self.basis/"bereich.xlsx");wb.close()
        for bereich in ("A:A","1:3","B4:A1","A0:C3","XFE1:XFE2","A1:B1048577","Daten!A1:B2","A1:B2:C3"):
            with self.subTest(bereich=bereich):
                with self.assertRaisesRegex(DatenFehler,"vollständigen Zellbereich"):
                    self.engine.importieren("XLSX","bereich.xlsx",{"bereich":bereich})


    def test_lange_ganzzahlen_bleiben_in_vorschau_und_excel_exakt(self):
        from openpyxl import load_workbook
        (self.basis/"lang.csv").write_text("ID\n123456789012345678\n",encoding="utf-8")
        t=self.engine.importieren("CSV","lang.csv")
        t=self.engine.spalten(t,{"aktion":"TYP","spalte":"ID","typ":"GANZZAHL"})
        v=self.engine.vorschau(t)
        self.assertEqual(v["zeilen"],[["123456789012345678"]])
        self.assertTrue(v["grosse_ganzzahlen_als_text"])
        self.assertEqual(self.zeilen(t),[(123456789012345678,)])
        self.engine.exportieren(t,"XLSX","lang.xlsx")
        wb=load_workbook(self.basis/"lang.xlsx",read_only=True)
        try:
            self.assertEqual(wb.active["A2"].value,"123456789012345678")
            self.assertEqual(wb.active["A2"].data_type,"s")
        finally: wb.close()

if __name__ == "__main__":
    unittest.main()
