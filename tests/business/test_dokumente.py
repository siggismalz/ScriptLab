"""Echte Dokumentfixtures in eigenen Tempordnern; keine Benutzerdateien oder Viewer."""
from __future__ import annotations

from datetime import datetime
import hashlib
from contextlib import redirect_stderr
from io import StringIO
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "laufzeit"))

from scriptlab_business.kontext import BusinessFehler, Grenzen, Kontext, UnklarerZustand
from scriptlab_business.dateien import archivieren, beleg_ablegen, csv_stapel, datei_aktion, datei_fingerabdruck, dateistapel, publizieren, zip_pruefen
from scriptlab_business.ausfuehrung import abschliessen, ausfuehren, kontext_laden, schritt
from scriptlab_business.dokumente import (
    docx_erstellen, docx_serie, docx_vorlage, pdf_bericht, pdf_seiten, pdf_teilen,
    pdf_text, pdf_verbinden, seiten_lesen, xlsx_vorlage,
)
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from PIL import Image


class TabellenFixture:
    def __init__(self, namen, zeilen, quellen=()):
        self.t = SimpleNamespace(spalten=[SimpleNamespace(name=n) for n in namen])
        self.daten = zeilen
        self.quellen = quellen
    def tabelle(self, ref):
        return self.t
    def zeilen(self, t, paket=1000):
        yield from self.daten
    def quellpfade(self, t):
        return self.quellen


class DokumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ScriptLab-Business-Dokumente-")
        self.basis = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def kontext(self, modus="ausfuehren", **optionen):
        k = Kontext(self.basis, modus, workflow_hash="a" * 64, **optionen)
        self.addCleanup(k.schliessen)
        return k

    def pdf(self, name, texte):
        pfad = self.basis / name
        c = canvas.Canvas(str(pfad))
        for text in texte:
            c.drawString(50, 780, text)
            c.showPage()
        c.save()
        return pfad

    def test_xlsx_erhaelt_format_formeln_und_textbleibt_daten(self):
        quelle = self.basis / "Vorlage.xlsx"
        buch = Workbook()
        buch.active["A1"] = "Alt"
        buch.active["A1"].font = Font(bold=True)
        buch.active["B1"] = "=A2+1"
        buch.active["A2"] = 3
        buch.save(quelle)
        original = quelle.read_bytes()
        r = xlsx_vorlage(self.kontext(), {"vorlage": str(quelle), "ziel": "Ergebnis.xlsx", "zellen": {"A1": "=2+2", "A2": 9}, "formeln": {"C1": "=SUM(A2:A2)"}}, {})
        self.assertEqual(r["status"], "publiziert")
        ausgabe = load_workbook(self.basis / "Ergebnis.xlsx", data_only=False)
        self.addCleanup(ausgabe.close)
        self.assertEqual(ausgabe.active["A1"].data_type, "s")
        self.assertTrue(ausgabe.active["A1"].font.bold)
        self.assertEqual(ausgabe.active["B1"].value, "=A2+1")
        self.assertEqual(ausgabe.active["C1"].value, "=SUM(A2:A2)")
        self.assertEqual(ausgabe.calculation.calcMode, "manual")
        self.assertEqual(quelle.read_bytes(), original)

    def test_docx_runuebergreifend_umlaute_kopf_und_tabelle(self):
        quelle = self.basis / "Vorlage.docx"
        d = Document()
        p = d.add_paragraph("Vor ")
        p.add_run("{{K").bold = True
        p.add_run("und").italic = True
        p.add_run("e}} nach")
        d.add_paragraph("{{Straße}}")
        d.sections[0].header.paragraphs[0].text = "Kopf {{Kunde}}"
        d.add_table(rows=1, cols=1).cell(0, 0).text = "{{Kunde}}"
        d.save(quelle)
        original = quelle.read_bytes()
        docx_vorlage(self.kontext(), {"vorlage": str(quelle), "ziel": "Brief.docx", "werte": {"Kunde": "Müller & Sohn", "Straße": "Hauptstraße 2"}}, {})
        r = Document(self.basis / "Brief.docx")
        self.assertEqual(r.paragraphs[0].text, "Vor Müller & Sohn nach")
        self.assertTrue(r.paragraphs[0].runs[1].bold)
        self.assertEqual(r.paragraphs[1].text, "Hauptstraße 2")
        self.assertEqual(r.tables[0].cell(0, 0).text, "Müller & Sohn")
        self.assertIn("Müller", r.sections[0].header.paragraphs[0].text)
        self.assertEqual(quelle.read_bytes(), original)

    def test_docx_positionen_wiederholen_formatierte_musterzeile(self):
        quelle = self.basis / "Positionen.docx"
        d = Document()
        d.add_paragraph("Kunde {{Kunde}}")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Artikel", "Menge"
        t.cell(1, 0).paragraphs[0].add_run("{{Position.Name}}").bold = True
        t.cell(1, 1).text = "{{Position.Menge}}"
        d.save(quelle)
        k = self.kontext()
        k.dienste["daten"] = TabellenFixture(["Name", "Menge"], [("Papier", 2), ("Stift", 3)])
        docx_vorlage(k, {"vorlage": str(quelle), "ziel": "Positionen-fertig.docx", "werte": {"Kunde": "Lea"}, "positionen": {"tabelle": 0, "zeile": 1}}, {"tabelle": {"tabelle": "test"}})
        r = Document(self.basis / "Positionen-fertig.docx")
        self.assertEqual([[c.text for c in row.cells] for row in r.tables[0].rows], [["Artikel", "Menge"], ["Papier", "2"], ["Stift", "3"]])
        self.assertTrue(r.tables[0].cell(1, 0).paragraphs[0].runs[-1].bold)

    def test_fehlender_wordwert_publiziert_nichts(self):
        quelle = self.basis / "Fehlt.docx"
        d = Document()
        d.add_paragraph("{{Unbekannt}}")
        d.save(quelle)
        with self.assertRaisesRegex(BusinessFehler, "keinen Wert"):
            docx_vorlage(self.kontext(), {"vorlage": str(quelle), "ziel": "Fehlt-out.docx", "werte": {}}, {})
        self.assertFalse((self.basis / "Fehlt-out.docx").exists())

    def test_wordserie_hat_sichere_namen_und_eigene_dokumente(self):
        quelle = self.basis / "Serie.docx"
        d = Document()
        d.add_paragraph("Hallo {{Name}}")
        d.save(quelle)
        k = self.kontext()
        k.dienste["daten"] = TabellenFixture(["Name"], [("Anna",), ("../Berta",)])
        r = docx_serie(k, {"vorlage": str(quelle), "zielordner": ".", "dateiname": "Brief-{Name}.docx"}, {"tabelle": {"tabelle": "serie"}})
        self.assertEqual(r["anzahl"], 2)
        self.assertTrue(all(Path(e["pfad"]).parent == self.basis for e in r["dokumente"]))
        self.assertEqual(Document(r["dokumente"][1]["pfad"]).paragraphs[0].text, "Hallo ../Berta")

    def test_neues_word_dokument_mit_bild(self):
        bild = self.basis / "Bild.png"
        Image.new("RGB", (20, 20), "blue").save(bild)
        docx_erstellen(self.kontext(), {"ziel": "Neu.docx", "abschnitte": [{"art": "UEBERSCHRIFT", "text": "Bericht ä"}, {"art": "ABSATZ", "text": "Klarer Text"}, {"art": "BILD", "pfad": str(bild), "breite_cm": 3}]}, {})
        d = Document(self.basis / "Neu.docx")
        self.assertEqual(len(d.inline_shapes), 1)
        self.assertIn("Bericht ä", [p.text for p in d.paragraphs])

    def test_pdf_verbinden_auswaehlen_und_teilen(self):
        a, b = self.pdf("A.pdf", ["Erste", "Zweite"]), self.pdf("B.pdf", ["Dritte"])
        k = self.kontext()
        pdf_verbinden(k, {"quellen": [str(a), str(b)], "ziel": "Gesamt.pdf"}, {})
        gesamt = PdfReader(self.basis / "Gesamt.pdf")
        self.assertEqual(len(gesamt.pages), 3)
        self.assertIn("Dritte", gesamt.pages[2].extract_text())
        pdf_seiten(k, {"quelle": "Gesamt.pdf", "ziel": "Auswahl.pdf", "seiten": "3,1"}, {})
        self.assertIn("Dritte", PdfReader(self.basis / "Auswahl.pdf").pages[0].extract_text())
        r = pdf_teilen(k, {"quelle": "Gesamt.pdf", "zielordner": ".", "bereiche": ["1-2", "3"]}, {})
        self.assertEqual([len(PdfReader(o["pfad"]).pages) for o in r["dokumente"]], [2, 1])

    def test_pdf_text_unterscheidet_scan_und_text(self):
        text = self.pdf("Text.pdf", ["Gruessung aus der Testdatei"])
        self.assertFalse(pdf_text(self.kontext("vorschau"), {"quelle": str(text)}, {})["scan_verdacht"])
        bild = self.basis / "Scan.png"
        Image.new("RGB", (30, 30), "white").save(bild)
        scan = self.basis / "Scan.pdf"
        c = canvas.Canvas(str(scan))
        c.drawImage(str(bild), 10, 10)
        c.save()
        r = pdf_text(self.kontext("vorschau"), {"quelle": str(scan)}, {})
        self.assertTrue(r["scan_verdacht"])
        self.assertFalse(r["ocr"])

    def test_pdf_seiten_grenzen_und_duplikate(self):
        for text in ("0", "3-1", "1,1", "1-9", "__import__('os')"):
            with self.subTest(text=text), self.assertRaises(BusinessFehler):
                seiten_lesen(text, 3)

    def test_pdf_bericht_unicode_kopf_und_seitennummer(self):
        font = next((p for p in [Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/segoeui.ttf")] if p.is_file()), None)
        self.assertIsNotNone(font, "Windows-Dokumenttest benötigt eine lokale TTF-Schrift.")
        k = self.kontext()
        k.dienste["daten"] = TabellenFixture(["Region", "Betrag"], [(f"München {i}", f"{i}.50") for i in range(130)])
        pdf_bericht(k, {"ziel": "Bericht.pdf", "titel": "Monatsbericht ä", "schrift": str(font)}, {"tabelle": {"tabelle": "bericht"}})
        r = PdfReader(self.basis / "Bericht.pdf")
        self.assertGreater(len(r.pages), 1)
        for seite in r.pages:
            text = seite.extract_text()
            self.assertIn("Region", text)
            self.assertIn("Seite", text)
            self.assertIn("Monatsbericht ä", text)
        self.assertIn("München", r.pages[0].extract_text())

    def test_quellidentitaet_ist_stabil_und_aenderung_verhindert_publikation(self):
        quelle = self.basis / "Quelle.txt"
        quelle.write_text("vorher", encoding="utf-8")
        k = self.kontext()
        self.assertEqual(datei_fingerabdruck(k, quelle), datei_fingerabdruck(k, quelle))
        def schreiben(temp):
            temp.write_text("Ausgabe", encoding="utf-8")
            quelle.write_text("nachher", encoding="utf-8")
        with self.assertRaisesRegex(BusinessFehler, "verändert"):
            publizieren(k, "Ziel.txt", schreiben, quellen=[quelle])
        self.assertFalse((self.basis / "Ziel.txt").exists())

    def test_konfliktregeln_und_prueflauf(self):
        ziel = self.basis / "Ziel.txt"
        ziel.write_text("Alt", encoding="utf-8")
        k = self.kontext()
        schreiben = lambda temp: temp.write_text("Neu", encoding="utf-8")
        with self.assertRaises(BusinessFehler):
            publizieren(k, ziel, schreiben)
        self.assertEqual(publizieren(k, ziel, schreiben, "UEBERSPRINGEN")["status"], "uebersprungen")
        self.assertEqual(ziel.read_text(), "Alt")
        self.assertEqual(Path(publizieren(k, ziel, schreiben, "NEUER_NAME")["pfad"]).name, "Ziel (2).txt")
        publizieren(k, ziel, schreiben, "ERSETZEN")
        self.assertEqual(ziel.read_text(), "Neu")
        p = self.kontext("prueflauf")
        r = publizieren(p, "Geplant.txt", schreiben)
        self.assertEqual(r["status"], "geplant")
        self.assertFalse((self.basis / "Geplant.txt").exists())
        self.assertFalse(list(self.basis.glob(".scriptlab-*")))

    def test_zip_traversal_dtd_und_groesse_werden_abgewiesen(self):
        for name, text, limit in [("../fremd.xml", "<a/>", Grenzen()), ("word/a.xml", '<!DOCTYPE a [<!ENTITY x "boom">]><a>&x;</a>', Grenzen()), ("word/a.xml", "<a>zu lang</a>", Grenzen(max_entpackt_bytes=4))]:
            with self.subTest(name=name, text=text):
                pfad = self.basis / "Boese.docx"
                with zipfile.ZipFile(pfad, "w") as z:
                    z.writestr(name, text)
                with self.assertRaises(BusinessFehler):
                    zip_pruefen(self.kontext(grenzen=limit), pfad)
        self.assertFalse((self.basis.parent / "fremd.xml").exists())

    def test_zip_lehnt_windows_mehrdeutigkeiten_und_kaputtes_xml_ab(self):
        for name, text in [("word/./a.xml", "<a/>"), ("word//a.xml", "<a/>"), ("word/a.xml ", "<a/>"), ("NUL.txt", "Daten"), ("word/a.xml", "<a>")]:
            with self.subTest(name=name):
                pfad = self.basis / "Ungueltig.zip"
                with zipfile.ZipFile(pfad, "w") as z:
                    z.writestr(name, text)
                with self.assertRaises(BusinessFehler):
                    zip_pruefen(self.kontext(), pfad)

    def test_stapel_hat_stabile_reihenfolge_und_hashes(self):
        for name in ("z.csv", "A.csv", "b.csv", "ignore.txt"):
            (self.basis / name).write_text(name, encoding="utf-8")
        r = dateistapel(self.kontext("vorschau"), {"ordner": ".", "muster": "*.csv"}, {})
        self.assertEqual([Path(e["pfad"]).name for e in r["dateien"]], ["A.csv", "b.csv", "z.csv"])
        self.assertTrue(all(len(e["sha256"]) == 64 for e in r["dateien"]))

    def test_archivierung_erst_nach_wirklich_geprueftem_output(self):
        k = self.kontext()
        quelle = self.basis / "Original.csv"
        quelle.write_text("Original", encoding="utf-8")
        output = publizieren(k, "Verarbeitet.csv", lambda p: p.write_text("Geprüft", encoding="utf-8"))
        r = archivieren(k, {"quelle": str(quelle), "ziel": "Archiv.csv"}, {"ausgaben": [output]})
        self.assertEqual(r["status"], "archiviert")
        self.assertFalse(quelle.exists())
        self.assertEqual((self.basis / "Archiv.csv").read_text(), "Original")

    def test_resume_verifiziert_output_und_unklarer_zustand_blockiert(self):
        journal = self.basis / "Journal.jsonl"
        k = self.kontext(journal_pfad=journal)
        k.beginn_schritt("block", {}, {})
        publizieren(k, "Ergebnis.txt", lambda p: p.write_text("Gut", encoding="utf-8"), ausgabe_id="eins")
        k.schliessen()
        neu = self.kontext(journal_pfad=journal, fortsetzen=True)
        neu.beginn_schritt("block", {}, {})
        self.assertEqual(publizieren(neu, "Ergebnis.txt", lambda p: p.write_text("Gut", encoding="utf-8"), ausgabe_id="eins")["status"], "wiederverwendet")
        neu.journal.schreiben("vorbereitet", "unklar", pfad=str(self.basis / "Unklar.txt"))
        neu.schliessen()
        offen = self.kontext(journal_pfad=journal, fortsetzen=True)
        with self.assertRaises(UnklarerZustand):
            _ = offen.journal

    def test_dateiaktionen_kopieren_verschieben_umbenennen_und_ordner(self):
        quelle = self.basis / "Quelle.txt"
        quelle.write_text("ä Dateiinhalt", encoding="utf-8")
        k = self.kontext()
        schritt(k, "ordner", "dateien.aktion", {"aktion": "ORDNER", "ziel": "Unterordner"}, {})
        schritt(k, "kopieren", "dateien.aktion", {"aktion": "KOPIEREN", "quelle": "Quelle.txt", "ziel": "Kopie.txt"}, {})
        self.assertEqual(quelle.read_bytes(), (self.basis / "Kopie.txt").read_bytes())
        schritt(k, "verschieben", "dateien.aktion", {"aktion": "VERSCHIEBEN", "quelle": "Kopie.txt", "ziel": "Unterordner/Datei.txt"}, {})
        self.assertFalse((self.basis / "Kopie.txt").exists())
        schritt(k, "umbenennen", "dateien.aktion", {"aktion": "UMBENENNEN", "quelle": "Unterordner/Datei.txt", "ziel": "Neu.txt"}, {})
        self.assertTrue((self.basis / "Unterordner/Neu.txt").is_file())
        self.assertTrue(quelle.is_file())

    def test_zip_aktion_rundlauf_und_prueflauf_hat_keine_zielordner(self):
        (self.basis / "Eingang/Tief").mkdir(parents=True)
        (self.basis / "Eingang/Tief/ä.txt").write_text("Sicher", encoding="utf-8")
        k = self.kontext()
        datei_aktion(k, {"aktion": "ZIP_ERSTELLEN", "quelle": "Eingang", "ziel": "Archiv.zip"}, {})
        plan = datei_aktion(self.kontext("prueflauf"), {"aktion": "ZIP_ENTPACKEN", "quelle": "Archiv.zip", "ziel": "NurGeplant"}, {})
        self.assertEqual(plan["anzahl"], 1)
        self.assertFalse((self.basis / "NurGeplant").exists())
        datei_aktion(k, {"aktion": "ZIP_ENTPACKEN", "quelle": "Archiv.zip", "ziel": "Ausgepackt"}, {})
        self.assertEqual((self.basis / "Ausgepackt/Tief/ä.txt").read_text(encoding="utf-8"), "Sicher")

    def test_word_gruppenserie_hat_zwei_positionen_und_summe_27(self):
        quelle = self.basis / "Bestellung.docx"
        d = Document()
        d.add_paragraph("Lieferant {{Lieferant}} Gesamt {{Gesamtbestellmenge}}")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Artikel", "Bestellung"
        t.cell(1, 0).text, t.cell(1, 1).text = "{{Position.Artikel}}", "{{Position.Bestellmenge}}"
        d.save(quelle)
        k = self.kontext()
        k.dienste["daten"] = TabellenFixture(["Lieferant", "Artikel", "Bestellmenge"], [("S1", "A", 15), ("S1", "C", 12)])
        r = docx_serie(k, {"vorlage": str(quelle), "zielordner": ".", "dateiname": "Bestellung-{Lieferant}.docx", "gruppe_nach": "Lieferant", "positionen": {"tabelle": 0, "zeile": 1}, "summen": {"Gesamtbestellmenge": "Bestellmenge"}}, {"tabelle": {"tabelle": "bestellung"}})
        self.assertEqual(r["anzahl"], 1)
        dokument = Document(r["dokumente"][0]["pfad"])
        self.assertEqual(dokument.paragraphs[0].text, "Lieferant S1 Gesamt 27")
        self.assertEqual(len(dokument.tables[0].rows), 3)
        self.assertEqual([row.cells[1].text for row in dokument.tables[0].rows[1:]], ["15", "12"])

    def test_beleg_monatsablage_und_scan_hinweis(self):
        beleg = self.pdf("Rechnung.pdf", ["Rechnung R-2026-001", "Datum 2026-10-05"])
        original = beleg.read_bytes()
        k = self.kontext()
        r = beleg_ablegen(k, {"quelle": str(beleg), "zielordner": "Ablage"}, {})
        self.assertEqual(r["datum"], "2026-10-05")
        self.assertEqual(Path(r["pfad"]).relative_to(self.basis).as_posix(), "Ablage/2026-10/R-2026-001.pdf")
        self.assertEqual(beleg.read_bytes(), original)
        bild = self.basis / "Scan.png"
        Image.new("RGB", (10, 10), "white").save(bild)
        c = canvas.Canvas(str(self.basis / "Scan.pdf"))
        c.drawImage(str(bild), 10, 10)
        c.save()
        h = beleg_ablegen(k, {"quelle": "Scan.pdf", "zielordner": "Ablage", "bei_scan": "BERICHT"}, {})
        self.assertEqual(h["status"], "hinweis")
        self.assertTrue(h["scan_verdacht"])
        self.assertTrue((self.basis / "Scan.pdf").exists())

    def test_csv_stapel_fehlerquelle_bleibt_und_resume_prueft_erledigte_outputs(self):
        for name in ("Eingang", "Ausgabe", "Archiv"):
            (self.basis / name).mkdir()
        (self.basis / "Eingang/Gut.csv").write_text("Nummer;Wert\n001;ä\n", encoding="utf-8")
        (self.basis / "Eingang/Kaputt.csv").write_text("A;B\n1;2;3\n", encoding="utf-8")
        journal = self.basis / "Journal.jsonl"
        config = {"ordner": "Eingang", "zielordner": "Ausgabe", "archivordner": "Archiv", "bei_fehler": "BERICHT"}
        k = self.kontext(journal_pfad=journal)
        k.beginn_schritt("stapel", config, {})
        r = csv_stapel(k, config, {})
        self.assertEqual(r["status"], "mit_fehlern")
        self.assertTrue((self.basis / "Archiv/Gut.csv").exists())
        self.assertFalse((self.basis / "Eingang/Gut.csv").exists())
        self.assertTrue((self.basis / "Eingang/Kaputt.csv").exists())
        self.assertFalse((self.basis / "Ausgabe/Kaputt.csv").exists())
        self.assertIn("001", (self.basis / "Ausgabe/Gut.csv").read_text(encoding="utf-8-sig"))
        k.schliessen()
        neu = self.kontext(journal_pfad=journal, fortsetzen=True)
        neu.beginn_schritt("stapel", config, {})
        fortgesetzt = csv_stapel(neu, config, {})
        self.assertTrue(any(e["status"] == "wiederverwendet" for e in fortgesetzt["dateien"]))
        self.assertTrue((self.basis / "Eingang/Kaputt.csv").exists())

    def test_prueflauf_csv_stapel_veraendert_keine_fachlichen_dateien(self):
        for name in ("Eingang", "Ausgabe", "Archiv"):
            (self.basis / name).mkdir()
        quelle = self.basis / "Eingang/Test.csv"
        quelle.write_text("A;B\n001;2\n", encoding="utf-8")
        original = quelle.read_bytes()
        r = csv_stapel(self.kontext("prueflauf"), {"ordner": "Eingang", "zielordner": "Ausgabe", "archivordner": "Archiv"}, {})
        self.assertEqual(r["status"], "erfolgreich")
        self.assertEqual(quelle.read_bytes(), original)
        self.assertEqual(list((self.basis / "Ausgabe").iterdir()), [])
        self.assertEqual(list((self.basis / "Archiv").iterdir()), [])

    def test_runner_parameter_verweise_und_persistenter_bericht(self):
        from scriptlab_business.kontext import daten_hash
        (self.basis / "Eingang.csv").write_text("Region;Wert\nNord;100\nSued;80\n", encoding="utf-8")
        meldungen = []
        k = self.kontext("prueflauf", parameter={"eingang": "Eingang.csv", "api_token": "nicht-in-bericht"}, journal_pfad=self.basis / "Berichte/Lauf.jsonl", ereignis_callback=meldungen.append)
        plan = {"version": 1, "schritte": [
            {"id": "lesen", "operation": "daten.importieren", "konfiguration": {"format": "CSV", "pfad": {"parameter": "eingang"}}, "eingaben": {}},
            {"id": "schreiben", "operation": "daten.exportieren", "konfiguration": {"format": "CSV", "pfad": "Ausgabe.csv"}, "eingaben": {"tabelle": {"referenz": "lesen"}}},
        ]}
        r = ausfuehren(plan, k)
        self.assertEqual(r["status"], "erfolgreich")
        self.assertFalse((self.basis / "Ausgabe.csv").exists())
        self.assertEqual(abschliessen(k), 0)
        bericht = json.loads(next((self.basis / "Berichte").glob("bericht-*.json")).read_text())
        self.assertEqual(bericht["parameter"]["api_token"], "[nicht protokolliert]")
        self.assertEqual(bericht["modus"], "prueflauf")
        self.assertTrue(bericht["runtime"]["pakete"]["duckdb"])
        self.assertEqual(meldungen[-1]["art"], "abschluss")

    def test_technischer_ordner_und_kooperatives_abbruchsignal(self):
        signal = self.basis / "Technisch/abbruch.signal"
        k = self.kontext(technischer_ordner=signal.parent, abbruch_datei=signal)
        self.assertTrue(k.arbeitsordner.is_relative_to(signal.parent))
        signal.write_text("abbrechen", encoding="utf-8")
        with self.assertRaisesRegex(BusinessFehler, "abgebrochen"):
            k.pruefen()
        arbeitsordner = k.arbeitsordner
        k.schliessen()
        self.assertFalse(arbeitsordner.exists())

    def test_exporthash_verwendet_scriptbytes_statt_metadaten(self):
        script = self.basis / "workflow.py"
        script.write_text("# Erste Fassung\n", encoding="utf-8")
        metadaten = {"workflow_hash": "veraltete-metadaten", "parameter": [], "standard_parameter": {}}
        with patch.object(sys, "argv", [str(script)]), redirect_stderr(StringIO()):
            k = kontext_laden(script, metadaten)
            self.addCleanup(k.schliessen)
            self.assertEqual(k.workflow_hash, hashlib.sha256(script.read_bytes()).hexdigest())
            alter_hash = k.workflow_hash
            script.write_text("# Zweite Fassung\n", encoding="utf-8")
            neu = kontext_laden(script, metadaten)
            self.addCleanup(neu.schliessen)
            self.assertNotEqual(neu.workflow_hash, alter_hash)

    def test_abschluss_meldet_vorbereitete_publikation_als_unklar(self):
        k = self.kontext(journal_pfad=self.basis / "Laufdaten/journal.jsonl")
        k.journal.schreiben("vorbereitet", "block:1:datei", pfad=str(self.basis / "Unklar.txt"))
        self.assertEqual(abschliessen(k), 1)
        bericht = json.loads(next((self.basis / "Laufdaten").glob("bericht-*.json")).read_text())
        self.assertEqual(bericht["status"], "unklar")
        self.assertEqual(bericht["unklare_schritte"], ["block:1:datei"])

    def test_vorschau_abschluss_zeigt_letzte_tabelle_ohne_ausgabe(self):
        (self.basis / "Daten.csv").write_text("Region;Wert\nNord;100\nSued;80\n", encoding="utf-8")
        meldungen = []
        k = self.kontext("vorschau", journal_pfad=self.basis / "Laufdaten/journal.jsonl", ereignis_callback=meldungen.append)
        schritt(k, "lesen", "daten.importieren", {"format": "CSV", "pfad": "Daten.csv"}, {})
        schritt(k, "ausgabe", "daten.exportieren", {"format": "CSV", "pfad": "NichtErzeugen.csv"}, {"tabelle": {"referenz": "lesen"}})
        self.assertEqual(abschliessen(k), 0)
        vorschau = next(e for e in meldungen if e["art"] == "vorschau")
        self.assertEqual(len(vorschau["zeilen"]), 2)
        self.assertFalse(vorschau["abgeschnitten"])
        self.assertFalse((self.basis / "NichtErzeugen.csv").exists())

    def test_tabellenexport_resume_nutzt_fachliche_herkunft_statt_uuid(self):
        quelle = self.basis / "Daten.csv"
        quelle.write_text("Region;Wert\nNord;100\nSued;80\n", encoding="utf-8")
        journal = self.basis / "Laufdaten/journal.jsonl"
        plan = {"version": 1, "schritte": [
            {"id": "lesen", "operation": "daten.importieren", "konfiguration": {"format": "CSV", "pfad": "Daten.csv"}, "eingaben": {}},
            {"id": "filtern", "operation": "daten.filtern", "konfiguration": {"bedingung": {"art": "vergleich", "spalte": "Region", "operator": "EQ", "wert": {"art": "wert", "typ": "TEXT", "wert": "Nord"}}}, "eingaben": {"tabelle": {"referenz": "lesen"}}},
            {"id": "schreiben", "operation": "daten.exportieren", "konfiguration": {"format": "CSV", "pfad": "Ergebnis.csv"}, "eingaben": {"tabelle": {"referenz": "filtern"}}},
        ]}
        k = self.kontext(journal_pfad=journal)
        ausfuehren(plan, k)
        alte_kennung = k.ergebnisse["lesen"]["tabelle"]
        original = (self.basis / "Ergebnis.csv").read_bytes()
        k.schliessen()
        neu = self.kontext(journal_pfad=journal, fortsetzen=True)
        ausfuehren(plan, neu)
        self.assertNotEqual(neu.ergebnisse["lesen"]["tabelle"], alte_kennung)
        self.assertEqual(neu.ergebnisse["schreiben"]["status"], "wiederverwendet")
        self.assertEqual((self.basis / "Ergebnis.csv").read_bytes(), original)
        neu.schliessen()
        veraendert = self.kontext(journal_pfad=journal, fortsetzen=True)
        plan["schritte"][1]["konfiguration"]["bedingung"]["wert"]["wert"] = "Sued"
        with self.assertRaises(UnklarerZustand):
            ausfuehren(plan, veraendert)
        veraendert.schliessen()
        plan["schritte"][1]["konfiguration"]["bedingung"]["wert"]["wert"] = "Nord"
        quelle.write_text("Region;Wert\nNord;999\n", encoding="utf-8")
        quelle_geaendert = self.kontext(journal_pfad=journal, fortsetzen=True)
        with self.assertRaises(UnklarerZustand):
            ausfuehren(plan, quelle_geaendert)
        self.assertEqual((self.basis / "Ergebnis.csv").read_bytes(), original)

    def test_ungueltiger_plan_wird_vor_erster_dateiaktion_abgelehnt(self):
        eintrag = {"id": "doppelt", "operation": "dateien.aktion", "konfiguration": {"aktion": "ORDNER", "ziel": "NichtAnlegen"}, "eingaben": {}}
        with self.assertRaises(BusinessFehler):
            ausfuehren({"version": 1, "schritte": [eintrag, eintrag]}, self.kontext())
        self.assertFalse((self.basis / "NichtAnlegen").exists())


if __name__ == "__main__":
    unittest.main()
