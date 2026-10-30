"""Die gestaltete Vorlage erhält den vorhandenen Vertrag des Beispiels C."""
from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
import re
import shutil
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document

WURZEL = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(WURZEL / "laufzeit"))
VORLAGE = WURZEL / "beispiele/business/C-Bestellvorschlag/vorlagen/Bestellung.docx"


def modul_laden(name, pfad):
    spezifikation = importlib.util.spec_from_file_location(name, pfad)
    modul = importlib.util.module_from_spec(spezifikation)
    spezifikation.loader.exec_module(modul)
    return modul


class BestellvorlageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ScriptLab-Bestellvorlage-")
        self.addCleanup(self.temp.cleanup)
        self.basis = Path(self.temp.name).resolve()

    def test_generator_erhaelt_vorlagenvertrag_und_eingecheckte_teile(self):
        generator = modul_laden("bestellvorlage", WURZEL / "scripts/vorlage-bestellung-erstellen.py")
        ziel = generator.erstellen(self.basis / "Vorlage.docx")
        # Logische Paketbestandteile vergleichen; Kompressionsbytes dürfen sich
        # zwischen Python-/zlib-Builds unterscheiden.
        with ZipFile(VORLAGE) as original, ZipFile(ziel) as erstellt:
            self.assertEqual(set(original.namelist()), set(erstellt.namelist()))
            for name in original.namelist():
                self.assertEqual(original.read(name), erstellt.read(name), name)
        dokument = Document(ziel)
        self.assertEqual(len(dokument.tables), 1)
        zeilen = [[z.text for z in zeile.cells] for zeile in dokument.tables[0].rows]
        self.assertEqual(zeilen, [
            ["Artikel", "Bestand", "Zielbestand", "Bestellmenge"],
            ["{{Position.Artikel}}", "{{Position.Bestand}}", "{{Position.Zielbestand}}", "{{Position.Bestellmenge}}"],
        ])
        text = "\n".join([p.text for p in dokument.paragraphs] + [z for zeile in zeilen for z in zeile])
        self.assertEqual(set(re.findall(r"\{\{([^{}]+)\}\}", text)), {
            "Lieferant", "Datum", "Gesamtbestellmenge", "Position.Artikel",
            "Position.Bestand", "Position.Zielbestand", "Position.Bestellmenge",
        })
        self.assertFalse(dokument.tables[0]._tbl.xpath(".//w:trHeight[@w:hRule='exact']"))

    def test_vollstaendiger_plan_schreibt_zwei_positionen_und_27_stueck(self):
        beispiele = modul_laden("beispiele_bestellvorlage", WURZEL / "scripts/business-beispiele-pruefen.py")
        basis = beispiele.beispiel_kopieren("C", self.basis)
        bericht = beispiele.pruefe_c(basis)
        self.assertEqual((bericht["dokumente"], bericht["positionen"], bericht["gesamtbestellmenge"]), (1, [15, 12], 27))
        self.assertTrue(bericht["quellen_unveraendert"])
        self.assertTrue(bericht["vorhandenes_ziel_geschuetzt"])
        dokument = Document(basis / "ausgaben/Bestellung-S1.docx")
        text = "\n".join(p.text for p in dokument.paragraphs)
        self.assertIn("Lieferant: S1", text)
        self.assertIn("Bestelldatum: 2026-10-26", text)
        self.assertIn("Gesamtbestellmenge: 27 Stück", text)
        self.assertEqual([[z.text for z in zeile.cells] for zeile in dokument.tables[0].rows[1:]], [
            ["Artikel A", "5", "20", "15"], ["Artikel C", "0", "12", "12"],
        ])

    def test_fixture_neuaufbau_setzt_die_vorlagengestaltung_nicht_zurueck(self):
        beispiele = modul_laden("beispiele_vorlagenneuaufbau", WURZEL / "scripts/business-beispiele-pruefen.py")
        beispiele.fixtures_erstellen(self.basis)
        ziel = self.basis / "C-Bestellvorschlag/vorlagen/Bestellung.docx"
        with ZipFile(VORLAGE) as original, ZipFile(ziel) as erstellt:
            self.assertEqual(set(original.namelist()), set(erstellt.namelist()))
            for name in original.namelist():
                self.assertEqual(original.read(name), erstellt.read(name), name)

    def test_eingesetzter_platzhaltertext_bleibt_wortgetreu_und_quelle_erhalten(self):
        from scriptlab_business.dokumente import docx_vorlage
        from scriptlab_business.kontext import Kontext
        quelle = self.basis / "Vorlage.docx"
        shutil.copy2(VORLAGE, quelle)
        original = hashlib.sha256(quelle.read_bytes()).hexdigest()
        artikel = 'Artikel {{Lieferant}} mit Umlauten äöü und Zeichen < & >'
        with Kontext(self.basis, "ausfuehren", workflow_hash="bestellvorlage") as kontext:
            bericht = docx_vorlage(kontext, {
                "vorlage": quelle.name, "ziel": "Ergebnis.docx",
                "werte": {"Lieferant": "S1 {{Datum}}", "Datum": "2026-10-26", "Gesamtbestellmenge": 15},
                "positionen": {"tabelle": 0, "zeile": 1, "praefix": "Position"},
            }, {"positions_zeilen": [{"Artikel": artikel, "Bestand": 5, "Zielbestand": 20, "Bestellmenge": 15}]})
            self.assertEqual(bericht["status"], "publiziert")
        dokument = Document(self.basis / "Ergebnis.docx")
        self.assertIn("Lieferant: S1 {{Datum}}", [p.text for p in dokument.paragraphs])
        self.assertEqual(dokument.tables[0].cell(1, 0).text, artikel)
        self.assertEqual(hashlib.sha256(quelle.read_bytes()).hexdigest(), original)


if __name__ == "__main__":
    unittest.main()
